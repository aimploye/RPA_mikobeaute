from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
import calendar
from pathlib import Path
import re
from typing import Any

import xlrd  # type: ignore[import-untyped]
from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.cell.cell import MergedCell  # type: ignore[import-untyped]
from openpyxl.formula.translate import Translator  # type: ignore[import-untyped]
from openpyxl.styles import PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries  # type: ignore[import-untyped]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]


R14_SHEETS = (
    "Summary",
    "站前4樓",
    "站前11樓",
    "忠孝國際醫學3樓",
    "忠孝7樓",
    "忠孝健康7樓",
    "忠孝預防醫學3樓",
    "領用表",
)
R14_BRANCH_SHEETS = (
    "站前4樓",
    "站前11樓",
    "忠孝國際醫學3樓",
    "忠孝7樓",
    "忠孝健康7樓",
    "忠孝預防醫學3樓",
)
SUMMARY_BRANCH_ORDER = ("站前4樓", "站前11樓", "忠孝7樓", "忠孝國際醫學3樓", "忠孝健康7樓", "忠孝預防醫學3樓")
LEGACY_SUMMARY_BRANCH_ORDER = ("站前4樓", "站前11樓", "忠孝7樓", "忠孝國際醫學3樓", "忠孝健康7樓")
R14_BASE_SHEETS = ("Summary", "站前4樓", "站前11樓", "忠孝國際醫學3樓", "忠孝7樓", "忠孝健康7樓", "領用表")
R14_BRANCH_TEMPLATE_SHEET = "忠孝健康7樓"
R13_SUMMARY_MARKER = "商品數量合計"
DATE_RE = re.compile(r"\d{4}/\d{2}/\d{2}")
R14_OUTPUT_REPORT_DATE_RE = re.compile(r"(?P<year>\d{4})\s+demand planning-(?P<mmdd>\d{4})(?:_\d+)?\.xlsx$", re.IGNORECASE)
SUMMARY_COLUMN_RANGE_RE = re.compile(r"(Summary!\$?)([A-Z]+)(:\$?)([A-Z]+)")
R14_ORDER_DAYS = 21
R14_SAFETY_STOCK_DAYS = 14
SUMMARY_HEADER_TEXT_COLOR = "FF4A86E8"
SUMMARY_HEADER_INVENTORY_FILL = "FFDBE5F1"
SUMMARY_HEADER_ORDER_TITLE_FILL = "FFC04F15"
SUMMARY_HEADER_ORDER_FILL = "FFFCE2D5"
SUMMARY_HEADER_TURNOVER_TITLE_FILL = "FF205C98"


@dataclass(frozen=True)
class R13UsageRow:
    sequence: int
    branch: str
    item_code: str
    item_name: str
    capacity: str
    quantity: float


@dataclass(frozen=True)
class R13UsageData:
    start_date: datetime
    end_date: datetime
    report_month: str
    rows: tuple[R13UsageRow, ...]


@dataclass(frozen=True)
class R14TransformResult:
    output_path: Path
    report_month: str
    report_date: datetime
    imported_rows: int
    added_summary_items: int
    added_branch_items: int
    inventory_updated: int = 0
    inventory_unmatched: int = 0


@dataclass(frozen=True)
class R14TemplateInventorySyncResult:
    template_path: Path
    inventory_updated: int
    inventory_unmatched: int
    added_summary_items: int = 0
    added_branch_items: int = 0


@dataclass(frozen=True)
class R14TemplateMonthStateSyncResult:
    template_path: Path
    source_path: Path
    month_label: str
    updated_columns: int


@dataclass(frozen=True)
class ColumnDimensionSnapshot:
    min_col: int
    max_col: int
    index: str
    dimension: Any


@dataclass(frozen=True)
class HeaderTailSnapshot:
    start_col: int
    max_col: int
    max_row: int
    values: dict[tuple[int, int], object]
    merged_ranges: tuple[tuple[int, int, int, int], ...]
    column_dimensions: tuple[ColumnDimensionSnapshot, ...]


@dataclass(frozen=True)
class SummaryMonthGroup:
    columns: dict[str, int]
    inserted_at: int | None = None
    inserted_width: int = 0


@dataclass(frozen=True)
class SummaryMetricGroup:
    total_col: int
    branch_cols: dict[str, int]


@dataclass(frozen=True)
class BranchPlanningColumns:
    actual_col: int
    forecast_col: int
    order_col: int
    safety_col: int
    stock_col: int


@dataclass(frozen=True)
class R14BranchItemUsage:
    branch: str
    item_code: str
    item_name: str
    actual: float


@dataclass(frozen=True)
class R14WorkbookSnapshot:
    path: Path
    report_date: datetime
    report_month: str
    items: tuple[R14BranchItemUsage, ...]
    missing_branches: tuple[str, ...] = ()


class R14TransformError(RuntimeError):
    def __init__(self, error_code: str, message: str) -> None:
        super().__init__(message)
        self.error_code = error_code
        self.message = message


def validate_r14_template_file(template_path: Path) -> None:
    if not template_path.is_file() or template_path.stat().st_size <= 0:
        raise R14TransformError("R14_TEMPLATE_FILE_INVALID", f"R14 模板檔不存在或為空：{template_path}")
    try:
        workbook = load_workbook(template_path, read_only=True)
    except Exception as exc:
        raise R14TransformError("R14_TEMPLATE_WORKBOOK_INVALID", f"R14 模板檔無法讀取為 Excel workbook：{template_path}") from exc
    try:
        _validate_template(workbook.sheetnames)
    finally:
        workbook.close()


def sync_r14_template_inventory(
    template_path: Path,
    branch_inventory: dict[str, dict[str, float]],
    *,
    item_names: dict[str, str] | None = None,
    inventory_date: date | datetime | None = None,
) -> R14TemplateInventorySyncResult:
    workbook = load_workbook(template_path)
    _validate_template(workbook.sheetnames)
    _ensure_r14_branch_sheets(workbook)
    _ensure_summary_inventory_branch_column(workbook["Summary"])
    _ensure_summary_operational_metric_columns(workbook["Summary"])
    _ensure_summary_actual_branch_columns(workbook["Summary"])
    if inventory_date is not None:
        workbook["Summary"]["F2"] = _coerce_report_date(inventory_date)
    _style_summary_header_layout(workbook["Summary"])
    added_summary_items, added_branch_items = _ensure_inventory_items(workbook, branch_inventory, item_names or {})
    inventory_updated, inventory_unmatched = _write_branch_inventory_values(workbook, branch_inventory)
    if inventory_updated <= 0:
        raise R14TransformError(
            "R14_TEMPLATE_INVENTORY_NO_MATCH",
            "R14 模板庫存同步沒有找到任何可更新品項；請確認 Google Sheet Summary 的凱惠料號與模板各分館頁籤一致。",
        )
    try:
        workbook.save(template_path)
    except PermissionError as exc:
        raise R14TransformError(
            "W01_TEMPLATE_LOCKED",
            "R14 模板檔目前無法寫入，可能正被 Excel/同步程式鎖定，"
            f"或 Windows 權限不允許寫入該位置：{template_path}",
        ) from exc
    return R14TemplateInventorySyncResult(
        template_path=template_path,
        inventory_updated=inventory_updated,
        inventory_unmatched=inventory_unmatched,
        added_summary_items=added_summary_items,
        added_branch_items=added_branch_items,
    )


def r14_template_has_actual_month_state(
    template_path: Path,
    month_label: str,
    *,
    allow_legacy_missing_n006: bool = False,
) -> bool:
    return not r14_template_actual_month_state_issues(
        template_path,
        month_label,
        allow_legacy_missing_n006=allow_legacy_missing_n006,
    )


def r14_template_actual_month_state_issues(
    template_path: Path,
    month_label: str,
    *,
    allow_legacy_missing_n006: bool = False,
) -> tuple[str, ...]:
    issues: list[str] = []
    try:
        workbook = load_workbook(template_path, data_only=False)
        _validate_template(workbook.sheetnames)
    except Exception as exc:
        return (f"無法讀取 workbook：{type(exc).__name__}",)
    has_n006 = "忠孝預防醫學3樓" in workbook.sheetnames
    if not has_n006 and not allow_legacy_missing_n006:
        issues.append("缺少分館頁籤：忠孝預防醫學3樓")
    summary_group = _find_summary_month_group(workbook["Summary"], month_label)
    if summary_group is None:
        summary_group = _find_legacy_summary_month_group(workbook["Summary"], month_label)
    if summary_group is None:
        issues.append(f"Summary 缺少 {month_label} Actual 欄位群組")
    missing_branch_actual = [
        sheet_name
        for sheet_name in R14_BRANCH_SHEETS
        if sheet_name in workbook.sheetnames
        and _find_month_actual_column(workbook[sheet_name], month_label) is None
    ]
    if missing_branch_actual:
        issues.append(f"分館缺少 {month_label} Actual：{'、'.join(missing_branch_actual)}")
    return tuple(issues)


def sync_r14_template_actual_month_state(
    template_path: Path,
    source_path: Path,
    month_label: str,
    *,
    overwrite_values: bool = False,
) -> R14TemplateMonthStateSyncResult:
    template_workbook = load_workbook(template_path)
    source_workbook = load_workbook(source_path, data_only=False)
    _validate_template(template_workbook.sheetnames)
    _validate_template(source_workbook.sheetnames)
    _ensure_r14_branch_sheets(template_workbook)
    _ensure_summary_inventory_branch_column(template_workbook["Summary"])
    _ensure_summary_operational_metric_columns(template_workbook["Summary"])
    _ensure_summary_actual_branch_columns(template_workbook["Summary"])
    _style_summary_header_layout(template_workbook["Summary"])
    source_has_n006 = "忠孝預防醫學3樓" in source_workbook.sheetnames

    updated_columns = 0
    source_summary_group = _find_summary_month_group(source_workbook["Summary"], month_label)
    if source_summary_group is None:
        source_summary_group = _find_legacy_summary_month_group(source_workbook["Summary"], month_label)
    if source_summary_group is None:
        raise R14TransformError(
            "R14_TEMPLATE_MONTH_STATE_SOURCE_MISSING",
            f"R14 前月狀態來源檔找不到 Summary {month_label} Actual 欄位群組：{source_path}",
        )
    target_summary_group = _find_summary_month_group(template_workbook["Summary"], month_label)
    summary_group_inserted = target_summary_group is None
    if summary_group_inserted:
        target_summary_group = _ensure_summary_month_group(template_workbook["Summary"], month_label).columns
    if overwrite_values or summary_group_inserted:
        assert target_summary_group is not None
        for key, source_col in source_summary_group.items():
            _copy_actual_values_by_item_code(
                source_workbook["Summary"],
                source_col,
                template_workbook["Summary"],
                target_summary_group[key],
            )
            updated_columns += 1
        if not source_has_n006:
            _set_all_item_values(template_workbook["Summary"], target_summary_group["忠孝預防醫學3樓"], 0)
            updated_columns += 1

    for sheet_name in R14_BRANCH_SHEETS:
        if sheet_name == "忠孝預防醫學3樓" and not source_has_n006:
            target_col = _find_month_actual_column(template_workbook[sheet_name], month_label)
            if target_col is None:
                target_col = _ensure_branch_month_column(template_workbook[sheet_name], month_label)
            _set_all_item_values(template_workbook[sheet_name], target_col, 0)
            updated_columns += 1
            continue
        branch_source_col = _find_month_actual_column(source_workbook[sheet_name], month_label)
        if branch_source_col is None:
            raise R14TransformError(
                "R14_TEMPLATE_MONTH_STATE_SOURCE_MISSING",
                f"R14 前月狀態來源檔頁籤「{sheet_name}」找不到 {month_label} Actual 欄：{source_path}",
            )
        target_col = _find_month_actual_column(template_workbook[sheet_name], month_label)
        if target_col is not None and not overwrite_values:
            continue
        if target_col is None:
            target_col = _ensure_branch_month_column(template_workbook[sheet_name], month_label)
        _copy_actual_values_by_item_code(
            source_workbook[sheet_name], branch_source_col, template_workbook[sheet_name], target_col
        )
        updated_columns += 1

    if updated_columns > 0:
        try:
            template_workbook.save(template_path)
        except PermissionError as exc:
            raise R14TransformError(
                "R14_TEMPLATE_STATE_UPDATE_FAILED",
                f"R14 無法補齊 runtime 模板的前月 Actual 狀態：{template_path}",
            ) from exc
    return R14TemplateMonthStateSyncResult(
        template_path=template_path,
        source_path=source_path,
        month_label=month_label,
        updated_columns=updated_columns,
    )


def transform_r13_to_r14(
    raw_path: Path,
    template_path: Path,
    output_path: Path,
    *,
    expected_end_date: date | None = None,
    branch_inventory: dict[str, dict[str, float]] | None = None,
    branch_inventory_item_names: dict[str, str] | None = None,
    inventory_date: date | datetime | None = None,
    update_template_path: Path | None = None,
    planning_baseline_available: bool = True,
) -> R14TransformResult:
    usage = parse_r13_usage_summary(raw_path)
    return transform_r13_usage_to_r14(
        usage,
        template_path,
        output_path,
        expected_end_date=expected_end_date,
        branch_inventory=branch_inventory,
        branch_inventory_item_names=branch_inventory_item_names,
        inventory_date=inventory_date,
        update_template_path=update_template_path,
        planning_baseline_available=planning_baseline_available,
    )


def transform_r13_usage_to_r14(
    usage: R13UsageData,
    template_path: Path,
    output_path: Path,
    *,
    expected_end_date: date | None = None,
    branch_inventory: dict[str, dict[str, float]] | None = None,
    branch_inventory_item_names: dict[str, str] | None = None,
    inventory_date: date | datetime | None = None,
    update_template_path: Path | None = None,
    planning_baseline_available: bool = True,
) -> R14TransformResult:
    if expected_end_date is not None and usage.end_date.date() != expected_end_date:
        raise R14TransformError(
            "R14_SOURCE_DATE_MISMATCH",
            f"R13 raw data 查詢迄日是 {usage.end_date:%Y/%m/%d}，但 R14 預期迄日是 {expected_end_date:%Y/%m/%d}。",
        )
    workbook = load_workbook(template_path)
    _validate_template(workbook.sheetnames)
    _ensure_r14_branch_sheets(workbook)

    month_label = usage.report_month
    summary_inventory_columns = _ensure_summary_inventory_branch_column(workbook["Summary"])
    summary_operational_columns = _ensure_summary_operational_metric_columns(workbook["Summary"])
    _ensure_summary_actual_branch_columns(workbook["Summary"])
    if not planning_baseline_available:
        _ensure_blank_actual_month_state(workbook, _previous_month_label(month_label))
    summary_group = _ensure_summary_month_group(workbook["Summary"], month_label)
    summary_columns = summary_group.columns
    branch_columns: dict[str, int] = {}
    for sheet_name in R14_BRANCH_SHEETS:
        branch_columns[sheet_name] = _ensure_branch_month_column(workbook[sheet_name], month_label)

    report_date = _coerce_report_date(inventory_date) if inventory_date is not None else _template_report_date(workbook["Summary"])
    workbook["Summary"]["F2"] = report_date
    _style_summary_header_layout(workbook["Summary"])

    usage_sheet = workbook["領用表"]
    _replace_usage_month_rows(usage_sheet, usage)
    added_summary_items = _ensure_summary_items(workbook["Summary"], usage.rows)
    added_branch_items = 0
    for sheet_name in R14_BRANCH_SHEETS:
        rows = tuple(row for row in usage.rows if row.branch == sheet_name)
        added_branch_items += _ensure_branch_items(workbook[sheet_name], rows)
    inventory_added_summary, inventory_added_branch = _ensure_inventory_items(
        workbook,
        branch_inventory,
        branch_inventory_item_names or {},
    )
    added_summary_items += inventory_added_summary
    added_branch_items += inventory_added_branch
    _write_summary_actuals(workbook["Summary"], usage.rows, summary_columns)
    for sheet_name in R14_BRANCH_SHEETS:
        rows = tuple(row for row in usage.rows if row.branch == sheet_name)
        _write_branch_actuals(workbook[sheet_name], rows, branch_columns[sheet_name])
    inventory_updated, inventory_unmatched = _write_branch_inventory_values(workbook, branch_inventory)
    _refresh_year_total_formulas(workbook["Summary"], usage.end_date)
    for sheet_name in R14_BRANCH_SHEETS:
        _refresh_year_total_formulas(workbook[sheet_name], usage.end_date)
    _update_forecast_month_headers(workbook, month_label)
    _shift_branch_summary_formula_references(workbook, summary_group.inserted_at, summary_group.inserted_width)
    branch_planning = {
        sheet_name: _refresh_branch_planning_columns(
            workbook[sheet_name],
            month_label,
            planning_baseline_available=planning_baseline_available,
        )
        for sheet_name in R14_BRANCH_SHEETS
    }
    _refresh_summary_planning_columns(
        workbook["Summary"],
        month_label,
        branch_planning,
        planning_baseline_available=planning_baseline_available,
    )
    _write_summary_inventory_formulas(workbook["Summary"], summary_inventory_columns, workbook)
    _write_summary_operational_metric_formulas(
        workbook["Summary"],
        summary_inventory_columns,
        summary_operational_columns,
        branch_planning,
        usage.end_date,
        planning_baseline_available=planning_baseline_available,
    )
    _apply_r14_view_state(workbook, usage.end_date)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(output_path)
    except PermissionError as exc:
        raise R14TransformError(
            "R14_OUTPUT_LOCKED",
            f"R14 輸出檔目前無法寫入，可能正被 Excel 或同步程式開啟鎖定：{output_path}",
        ) from exc
    if update_template_path is not None and output_path.resolve() != update_template_path.resolve():
        try:
            update_template_path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(update_template_path)
        except PermissionError as exc:
            raise R14TransformError(
                "R14_TEMPLATE_STATE_UPDATE_FAILED",
                f"R14 已產出報表，但無法更新 runtime 模板狀態；下次跨月可能找不到前月 Actual：{update_template_path}",
            ) from exc
    return R14TransformResult(
        output_path=output_path,
        report_month=usage.report_month,
        report_date=report_date,
        imported_rows=len(usage.rows),
        added_summary_items=added_summary_items,
        added_branch_items=added_branch_items,
        inventory_updated=inventory_updated,
        inventory_unmatched=inventory_unmatched,
    )


def parse_r13_usage_summary(raw_path: Path) -> R13UsageData:
    workbook = xlrd.open_workbook(str(raw_path))
    if not workbook.sheet_names():
        raise R14TransformError("R13_RAW_EMPTY", f"R13 raw data 沒有任何工作表：{raw_path}")
    sheet = workbook.sheet_by_index(0)
    start_date, end_date = _parse_r13_period(sheet)
    marker_row = _find_marker_row(sheet, R13_SUMMARY_MARKER)
    header_row = marker_row + 1
    _validate_r13_summary_header(sheet, header_row)

    rows: list[R13UsageRow] = []
    for row_index in range(header_row + 1, sheet.nrows):
        first = _cell_text(sheet, row_index, 0)
        if first == "合計:":
            break
        branch = _cell_text(sheet, row_index, 1)
        item_code = _cell_text(sheet, row_index, 2)
        item_name = _cell_text(sheet, row_index, 4)
        if not branch and not item_code and not item_name:
            continue
        if branch not in R14_BRANCH_SHEETS:
            continue
        rows.append(
            R13UsageRow(
                sequence=len(rows) + 1,
                branch=branch,
                item_code=item_code,
                item_name=item_name,
                capacity=_cell_text(sheet, row_index, 8),
                quantity=_cell_number(sheet, row_index, 9),
            )
        )

    if not rows:
        raise R14TransformError("R13_USAGE_SUMMARY_EMPTY", "R13 raw data 找到商品數量合計，但沒有可匯入的分店彙總資料。")

    return R13UsageData(
        start_date=start_date,
        end_date=end_date,
        report_month=end_date.strftime("%Y/%m"),
        rows=tuple(rows),
    )


def load_r14_workbook_snapshot(
    path: Path,
    *,
    allow_legacy_missing_n006: bool = False,
) -> R14WorkbookSnapshot:
    workbook = load_workbook(path, data_only=False)
    _validate_template(workbook.sheetnames)
    missing_branches = tuple(sheet_name for sheet_name in R14_BRANCH_SHEETS if sheet_name not in workbook.sheetnames)
    if missing_branches and (
        not allow_legacy_missing_n006 or missing_branches != ("忠孝預防醫學3樓",)
    ):
        raise R14TransformError("R14_TEMPLATE_SHEET_MISSING", f"R14 輸出檔缺少分館頁籤「忠孝預防醫學3樓」：{path}")
    report_date = _r14_report_date_from_filename(path.name) or _coerce_excel_datetime(workbook["Summary"]["F2"].value)
    report_month = report_date.strftime("%Y/%m")
    items: list[R14BranchItemUsage] = []
    for sheet_name in (sheet_name for sheet_name in R14_BRANCH_SHEETS if sheet_name in workbook.sheetnames):
        sheet = workbook[sheet_name]
        actual_col = _require_month_actual_column(sheet, report_month)
        for row_index in range(4, _last_item_row(sheet) + 1):
            item_code = str(sheet.cell(row_index, 2).value or "").strip()
            item_name = str(sheet.cell(row_index, 3).value or "").strip()
            if not item_code and not item_name:
                continue
            items.append(
                R14BranchItemUsage(
                    branch=sheet_name,
                    item_code=item_code,
                    item_name=item_name,
                    actual=_numeric_cell_value(sheet.cell(row_index, actual_col).value),
                )
            )
    return R14WorkbookSnapshot(
        path=path,
        report_date=report_date,
        report_month=report_month,
        items=tuple(items),
        missing_branches=missing_branches,
    )


def _r14_report_date_from_filename(filename: str) -> datetime | None:
    match = R14_OUTPUT_REPORT_DATE_RE.search(filename)
    if match is None:
        return None
    year = int(match.group("year"))
    mmdd = match.group("mmdd")
    try:
        return datetime(year, int(mmdd[:2]), int(mmdd[2:]))
    except ValueError:
        return None


def _parse_r13_period(sheet: Any) -> tuple[datetime, datetime]:
    for row_index in range(min(sheet.nrows, 20)):
        row_text = " ".join(_cell_text(sheet, row_index, col) for col in range(sheet.ncols))
        dates = DATE_RE.findall(row_text)
        if len(dates) >= 2:
            return (
                datetime.strptime(dates[0], "%Y/%m/%d"),
                datetime.strptime(dates[1], "%Y/%m/%d"),
            )
    raise R14TransformError("R13_PERIOD_NOT_FOUND", "R13 raw data 找不到查詢範圍日期。")


def _find_marker_row(sheet: Any, marker: str) -> int:
    for row_index in range(sheet.nrows):
        for col_index in range(sheet.ncols):
            if marker in _cell_text(sheet, row_index, col_index):
                return row_index
    raise R14TransformError("R13_SUMMARY_MARKER_NOT_FOUND", f"R13 raw data 找不到標記：{marker}")


def _validate_r13_summary_header(sheet: Any, row_index: int) -> None:
    expected = {
        0: "序",
        1: "分店",
        2: "商品碼",
        4: "商品名稱",
        8: "容量",
        9: "領用數量",
    }
    missing = [
        label
        for col_index, label in expected.items()
        if label not in _cell_text(sheet, row_index, col_index)
    ]
    if missing:
        raise R14TransformError("R13_SUMMARY_HEADER_INVALID", f"R13 商品數量合計表頭不完整：{', '.join(missing)}")


def _cell_text(sheet: Any, row_index: int, col_index: int) -> str:
    if row_index >= sheet.nrows or col_index >= sheet.ncols:
        return ""
    value = sheet.cell_value(row_index, col_index)
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _cell_number(sheet: Any, row_index: int, col_index: int) -> float:
    value = sheet.cell_value(row_index, col_index)
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").strip()
    return float(text) if text else 0.0


def _validate_template(sheetnames: list[str]) -> None:
    missing = [sheet_name for sheet_name in R14_BASE_SHEETS if sheet_name not in sheetnames]
    if missing:
        raise R14TransformError("R14_TEMPLATE_SHEET_MISSING", f"R14 模板缺少頁籤：{', '.join(missing)}")


def _ensure_r14_branch_sheets(workbook: Any) -> None:
    if "忠孝預防醫學3樓" in workbook.sheetnames:
        return
    source = workbook[R14_BRANCH_TEMPLATE_SHEET]
    target = workbook.copy_worksheet(source)
    target.title = "忠孝預防醫學3樓"
    source_name = _normalized_header_text(R14_BRANCH_TEMPLATE_SHEET)
    target_name = "忠孝預防醫學3樓"
    for row_index in range(1, min(target.max_row, 3) + 1):
        for col_index in range(1, target.max_column + 1):
            cell = target.cell(row_index, col_index)
            if isinstance(cell.value, str) and source_name in _normalized_header_text(cell.value):
                cell.value = cell.value.replace(R14_BRANCH_TEMPLATE_SHEET, target_name)
    usage_index = workbook.sheetnames.index("領用表")
    workbook._sheets.remove(target)
    workbook._sheets.insert(usage_index, target)


def _validate_month_supported(sheet: Worksheet, month_label: str) -> None:
    for col_index in range(1, sheet.max_column + 1):
        if _cell_month_value(sheet.cell(2, col_index).value) == month_label and sheet.cell(3, col_index).value == "Actual":
            return
    raise R14TransformError(
        "R14_TEMPLATE_MONTH_NOT_FOUND",
        f"R14 模板頁籤「{sheet.title}」找不到 {month_label} 的 Actual 欄位；為避免破壞公式，本版不自動插入新月份欄。",
    )


def _ensure_branch_month_column(sheet: Worksheet, month_label: str) -> int:
    existing = _find_month_actual_column(sheet, month_label)
    if existing is not None:
        return existing
    previous_month = _previous_month_label(month_label)
    previous_col = _find_month_actual_column(sheet, previous_month)
    if previous_col is None:
        raise R14TransformError(
            "R14_TEMPLATE_MONTH_NOT_FOUND",
            f"R14 模板頁籤「{sheet.title}」找不到 {month_label} 或前一月 {previous_month} 的 Actual 欄位。",
        )
    target_col = previous_col + 1
    header_tail = _capture_header_tail(sheet, target_col)
    sheet.insert_cols(target_col)
    _restore_shifted_header_tail(sheet, header_tail, shift=1)
    _unmerge_cells_in_columns(sheet, target_col, target_col, max_row=3)
    _copy_column(sheet, previous_col, target_col)
    sheet.cell(2, target_col).value = month_label
    sheet.cell(3, target_col).value = "Actual"
    _clear_item_column_values(sheet, target_col)
    return target_col


def _ensure_blank_actual_month_state(workbook: Any, month_label: str) -> None:
    """Create only the columns needed for a degraded report, never guessed values."""
    summary = workbook["Summary"]
    target = _parse_month_label(month_label)
    if target is None:
        raise R14TransformError("R14_TEMPLATE_MONTH_NOT_FOUND", f"無法解析 R14 月份：{month_label}")

    while _find_summary_month_group(summary, month_label) is None:
        existing_months = _complete_summary_months_at_or_before(summary, target)
        if not existing_months or existing_months[-1] >= target:
            raise R14TransformError(
                "R14_TEMPLATE_MONTH_NOT_FOUND",
                f"R14 模板沒有可延伸至 {month_label} 的既有 Actual 欄位群組。",
            )
        _ensure_summary_month_group(summary, _next_month_label(existing_months[-1]))

    for sheet_name in R14_BRANCH_SHEETS:
        sheet = workbook[sheet_name]
        while _find_month_actual_column(sheet, month_label) is None:
            existing_months = _branch_months_at_or_before(sheet, target)
            if not existing_months or existing_months[-1] >= target:
                raise R14TransformError(
                    "R14_TEMPLATE_MONTH_NOT_FOUND",
                    f"R14 模板頁籤「{sheet_name}」沒有可延伸至 {month_label} 的 Actual 欄。",
                )
            _ensure_branch_month_column(sheet, _next_month_label(existing_months[-1]))


def _complete_summary_months_at_or_before(sheet: Worksheet, target: tuple[int, int]) -> list[tuple[int, int]]:
    months: set[tuple[int, int]] = set()
    for col_index in range(1, sheet.max_column + 1):
        if sheet.cell(3, col_index).value != "Actual":
            continue
        parsed = _parse_month_label(sheet.cell(2, col_index).value)
        if parsed is None or parsed > target:
            continue
        label = f"{parsed[0]:04d}/{parsed[1]:02d}"
        if _find_summary_month_group(sheet, label) is not None:
            months.add(parsed)
    return sorted(months)


def _branch_months_at_or_before(sheet: Worksheet, target: tuple[int, int]) -> list[tuple[int, int]]:
    months = {
        parsed
        for col_index in range(1, sheet.max_column + 1)
        if sheet.cell(3, col_index).value == "Actual"
        and (parsed := _parse_month_label(sheet.cell(2, col_index).value)) is not None
        and parsed <= target
    }
    return sorted(months)


def _next_month_label(month: tuple[int, int]) -> str:
    year, month_number = month
    if month_number == 12:
        return f"{year + 1:04d}/01"
    return f"{year:04d}/{month_number + 1:02d}"


def _summary_branch_inventory_columns(sheet: Worksheet) -> dict[str, int]:
    columns: dict[str, int] = {}
    for offset, branch in enumerate(SUMMARY_BRANCH_ORDER):
        col_index = 6 + offset
        if _summary_inventory_header_matches(sheet.cell(3, col_index).value, branch):
            columns[branch] = col_index
            continue
        break
    return columns


def _summary_inventory_header_matches(value: object, branch: str) -> bool:
    normalized = _normalized_header_text(value)
    return "庫存" in normalized and _normalized_header_text(branch) in normalized


def _ensure_summary_inventory_branch_column(sheet: Worksheet) -> dict[str, int]:
    columns = _summary_branch_inventory_columns(sheet)
    if "忠孝預防醫學3樓" in columns:
        return columns
    expected_existing = SUMMARY_BRANCH_ORDER[: len(columns)]
    if tuple(columns) != expected_existing:
        raise R14TransformError(
            "R14_SUMMARY_BRANCH_LAYOUT_INVALID",
            "R14 Summary 左側分館庫存欄位順序不符合既有契約，拒絕插入新館欄位以避免公式錯位。",
        )
    target_col = 6 + len(columns)
    source_col = target_col - 1
    header_tail = _capture_header_tail(sheet, target_col)
    sheet.insert_cols(target_col)
    _restore_shifted_header_tail(sheet, header_tail, shift=1)
    _unmerge_cells_in_columns(sheet, target_col, target_col, max_row=3)
    _copy_column(sheet, source_col, target_col)
    sheet.cell(2, target_col).value = None
    sheet.cell(3, target_col).value = "忠孝預防醫學3樓\n庫存"
    _clear_item_column_values(sheet, target_col)
    return {**columns, "忠孝預防醫學3樓": target_col}


def _ensure_summary_operational_metric_columns(sheet: Worksheet) -> dict[str, list[int]]:
    inventory_end_col = 6 + len(SUMMARY_BRANCH_ORDER) - 1
    order_start = inventory_end_col + 1
    order_branches = SUMMARY_BRANCH_ORDER
    order_matches = [
        str(sheet.cell(3, order_start + offset).value or "").strip() == branch
        for offset, branch in enumerate(order_branches)
    ]
    turnover_start = order_start + len(order_branches)
    turnover_matches = [
        str(sheet.cell(3, turnover_start + offset).value or "").strip() == branch
        for offset, branch in enumerate(order_branches)
    ]
    if all(order_matches) and all(turnover_matches):
        return {
            "order": list(range(order_start, turnover_start)),
            "turnover": list(range(turnover_start, turnover_start + len(order_branches))),
        }
    if any(order_matches) or any(turnover_matches):
        raise R14TransformError(
            "R14_SUMMARY_METRIC_LAYOUT_INVALID",
            "R14 Summary 已存在不完整的下單數/週轉天數欄位，拒絕再次插入以避免破壞右側公式。",
        )

    width = len(order_branches) * 2
    header_tail = _capture_header_tail(sheet, order_start)
    sheet.insert_cols(order_start, amount=width)
    _restore_shifted_header_tail(sheet, header_tail, shift=width)
    _unmerge_cells_in_columns(sheet, order_start, order_start + width - 1, max_row=3)
    style_source = order_start - 1
    for col_index in range(order_start, order_start + width):
        _copy_column(sheet, style_source, col_index)
        _clear_item_column_values(sheet, col_index)
    order_end = order_start + len(order_branches) - 1
    turnover_start = order_end + 1
    turnover_end = turnover_start + len(order_branches) - 1
    sheet.merge_cells(start_row=2, start_column=order_start, end_row=2, end_column=order_end)
    sheet.merge_cells(start_row=2, start_column=turnover_start, end_row=2, end_column=turnover_end)
    sheet.cell(2, order_start).value = "下單數"
    sheet.cell(2, turnover_start).value = "週轉天數"
    for offset, branch in enumerate(order_branches):
        sheet.cell(3, order_start + offset).value = branch
        sheet.cell(3, turnover_start + offset).value = branch
    return {
        "order": list(range(order_start, order_end + 1)),
        "turnover": list(range(turnover_start, turnover_end + 1)),
    }


def _style_summary_header_layout(sheet: Worksheet) -> None:
    groups = (
        (6, 11, SUMMARY_HEADER_INVENTORY_FILL, SUMMARY_HEADER_TEXT_COLOR, False),
        (12, 17, SUMMARY_HEADER_ORDER_TITLE_FILL, "FFFFFFFF", True),
        (18, 23, SUMMARY_HEADER_TURNOVER_TITLE_FILL, "FFFFFFFF", True),
    )
    for start_col, end_col, _fill_color, _font_color, _bold in groups:
        _unmerge_summary_header_range(sheet, start_col, end_col)

    sheet.cell(2, 6).number_format = 'yyyy/m/d"庫存"'
    sheet.cell(2, 12).value = "下單數"
    sheet.cell(2, 18).value = "週轉天數"
    for start_col, end_col, fill_color, font_color, bold in groups:
        for row_index in (2, 3):
            row_fill = fill_color
            if row_index == 3 and start_col == 12:
                row_fill = SUMMARY_HEADER_ORDER_FILL
            elif row_index == 3 and start_col == 18:
                row_fill = SUMMARY_HEADER_INVENTORY_FILL
            for col_index in range(start_col, end_col + 1):
                cell = sheet.cell(row_index, col_index)
                cell.fill = PatternFill(fill_type="solid", fgColor=row_fill)
                font = copy(cell.font)
                font.color = font_color if row_index == 2 else SUMMARY_HEADER_TEXT_COLOR
                font.bold = bold if row_index == 2 else False
                cell.font = font
                alignment = copy(cell.alignment)
                alignment.horizontal = "center"
                alignment.vertical = "center"
                alignment.wrap_text = True
                cell.alignment = alignment

    sheet.merge_cells(start_row=2, start_column=6, end_row=2, end_column=11)
    sheet.merge_cells(start_row=2, start_column=12, end_row=2, end_column=17)
    sheet.merge_cells(start_row=2, start_column=18, end_row=2, end_column=23)


def _unmerge_summary_header_range(sheet: Worksheet, start_col: int, end_col: int) -> None:
    ranges_to_remove = []
    for merged_range in sheet.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
        if min_row <= 2 <= max_row and min_col <= end_col and max_col >= start_col:
            ranges_to_remove.append(merged_range)
    for merged_range in ranges_to_remove:
        sheet.unmerge_cells(str(merged_range))


def _ensure_summary_actual_branch_columns(sheet: Worksheet) -> None:
    legacy_groups: list[int] = []
    for col_index in range(1, sheet.max_column + 1):
        if _cell_month_value(sheet.cell(2, col_index).value) is None or sheet.cell(3, col_index).value != "Actual":
            continue
        if all(
            str(sheet.cell(2, col_index + offset).value or "").strip() == branch
            and sheet.cell(3, col_index + offset).value == "Actual"
            for offset, branch in enumerate(LEGACY_SUMMARY_BRANCH_ORDER, start=1)
        ):
            next_col = col_index + 1 + len(LEGACY_SUMMARY_BRANCH_ORDER)
            if not (
                str(sheet.cell(2, next_col).value or "").strip() == "忠孝預防醫學3樓"
                and sheet.cell(3, next_col).value == "Actual"
            ):
                legacy_groups.append(col_index)

    for total_col in reversed(legacy_groups):
        target_col = total_col + 1 + len(LEGACY_SUMMARY_BRANCH_ORDER)
        header_tail = _capture_header_tail(sheet, target_col)
        sheet.insert_cols(target_col)
        _restore_shifted_header_tail(sheet, header_tail, shift=1)
        _unmerge_cells_in_columns(sheet, target_col, target_col, max_row=3)
        _copy_column(sheet, target_col - 1, target_col)
        sheet.cell(2, target_col).value = "忠孝預防醫學3樓"
        sheet.cell(3, target_col).value = "Actual"
        _clear_item_column_values(sheet, target_col)


def _ensure_summary_month_group(sheet: Worksheet, month_label: str) -> SummaryMonthGroup:
    existing = _find_summary_month_group(sheet, month_label)
    if existing is not None:
        return SummaryMonthGroup(columns=existing)
    previous_month = _previous_month_label(month_label)
    previous_group = _find_summary_month_group(sheet, previous_month)
    if previous_group is None:
        raise R14TransformError(
            "R14_TEMPLATE_MONTH_NOT_FOUND",
            f"R14 模板頁籤「Summary」找不到 {month_label} 或前一月 {previous_month} 的 Actual 欄位群組。",
        )
    previous_total_col = previous_group["total"]
    width = 1 + len(SUMMARY_BRANCH_ORDER)
    target_col = previous_total_col + width
    header_tail = _capture_header_tail(sheet, target_col)
    sheet.insert_cols(target_col, amount=width)
    _restore_shifted_header_tail(sheet, header_tail, shift=width)
    _unmerge_cells_in_columns(sheet, target_col, target_col + width - 1, max_row=3)
    for offset in range(width):
        _copy_column(sheet, previous_total_col + offset, target_col + offset)
    _unmerge_cells_in_columns(sheet, target_col, target_col + width - 1, max_row=3)
    sheet.cell(2, target_col).value = month_label
    sheet.cell(3, target_col).value = "Actual"
    for offset, branch in enumerate(SUMMARY_BRANCH_ORDER, start=1):
        sheet.cell(2, target_col + offset).value = branch
        sheet.cell(3, target_col + offset).value = "Actual"
    for col_index in range(target_col, target_col + width):
        _clear_item_column_values(sheet, col_index)
    columns = {"total": target_col, **{branch: target_col + index + 1 for index, branch in enumerate(SUMMARY_BRANCH_ORDER)}}
    return SummaryMonthGroup(columns=columns, inserted_at=target_col, inserted_width=width)


def _find_month_actual_column(sheet: Worksheet, month_label: str) -> int | None:
    for col_index in range(1, sheet.max_column + 1):
        if _cell_month_value(sheet.cell(2, col_index).value) == month_label and sheet.cell(3, col_index).value == "Actual":
            return col_index
    return None


def _find_summary_month_group(sheet: Worksheet, month_label: str) -> dict[str, int] | None:
    total_col = _find_month_actual_column(sheet, month_label)
    if total_col is None:
        return None
    group = {"total": total_col}
    for offset, branch in enumerate(SUMMARY_BRANCH_ORDER, start=1):
        col_index = total_col + offset
        if str(sheet.cell(2, col_index).value or "").strip() != branch or sheet.cell(3, col_index).value != "Actual":
            return None
        group[branch] = col_index
    return group


def _find_legacy_summary_month_group(sheet: Worksheet, month_label: str) -> dict[str, int] | None:
    """Find a five-branch Actual group from a snapshot made before N006 existed."""
    total_col = _find_month_actual_column(sheet, month_label)
    if total_col is None:
        return None
    group = {"total": total_col}
    for offset, branch in enumerate(LEGACY_SUMMARY_BRANCH_ORDER, start=1):
        col_index = total_col + offset
        if (
            str(sheet.cell(2, col_index).value or "").strip() != branch
            or sheet.cell(3, col_index).value != "Actual"
        ):
            return None
        group[branch] = col_index
    return group


def _previous_month_label(month_label: str) -> str:
    year, month = (int(part) for part in month_label.split("/"))
    if month == 1:
        return f"{year - 1}/12"
    return f"{year}/{month - 1:02d}"


def _days_in_month_label(month_label: str) -> int:
    year, month = (int(part) for part in month_label.split("/"))
    return calendar.monthrange(year, month)[1]


def _capture_header_tail(sheet: Worksheet, start_col: int, *, max_row: int = 3) -> HeaderTailSnapshot:
    values: dict[tuple[int, int], object] = {}
    for col_index in range(start_col, sheet.max_column + 1):
        for row_index in range(1, max_row + 1):
            cell = sheet.cell(row_index, col_index)
            if isinstance(cell, MergedCell):
                continue
            values[(row_index, col_index)] = cell.value

    merged_ranges: list[tuple[int, int, int, int]] = []
    for merged_range in sheet.merged_cells.ranges:
        min_col, min_row, max_col, range_max_row = range_boundaries(str(merged_range))
        if max_col < start_col:
            continue
        if min_col < start_col:
            continue
        merged_ranges.append((min_col, min_row, max_col, range_max_row))

    column_dimensions: list[ColumnDimensionSnapshot] = []
    for key, dimension in sheet.column_dimensions.items():
        min_col = dimension.min or column_index_from_string(key)
        max_col = dimension.max or min_col
        if max_col < start_col:
            continue
        if min_col < start_col:
            continue
        column_dimensions.append(
            ColumnDimensionSnapshot(
                min_col=min_col,
                max_col=max_col,
                index=key,
                dimension=copy(dimension),
            )
        )

    return HeaderTailSnapshot(
        start_col=start_col,
        max_col=sheet.max_column,
        max_row=max(max_row, sheet.max_row),
        values=values,
        merged_ranges=tuple(merged_ranges),
        column_dimensions=tuple(column_dimensions),
    )


def _restore_shifted_header_tail(sheet: Worksheet, snapshot: HeaderTailSnapshot, *, shift: int) -> None:
    if snapshot.max_col < snapshot.start_col:
        return
    restore_end_col = snapshot.max_col + shift
    _unmerge_cells_from_column(sheet, snapshot.start_col, max_row=snapshot.max_row)
    _delete_column_dimensions_from_column(sheet, snapshot.start_col, restore_end_col)

    for (row_index, col_index), value in snapshot.values.items():
        target = sheet.cell(row_index, col_index + shift)
        if not isinstance(target, MergedCell):
            target.value = value

    for dimension_snapshot in snapshot.column_dimensions:
        dimension = copy(dimension_snapshot.dimension)
        dimension.min = dimension_snapshot.min_col + shift
        dimension.max = dimension_snapshot.max_col + shift
        dimension.index = get_column_letter(dimension.min)
        sheet.column_dimensions[dimension.index] = dimension

    for min_col, min_row, max_col, max_row in snapshot.merged_ranges:
        sheet.merge_cells(
            start_row=min_row,
            start_column=min_col + shift,
            end_row=max_row,
            end_column=max_col + shift,
        )


def _unmerge_cells_from_column(sheet: Worksheet, start_col: int, *, max_row: int) -> None:
    ranges_to_remove = []
    for merged_range in sheet.merged_cells.ranges:
        min_col, min_row, max_col, _max_row = range_boundaries(str(merged_range))
        if min_row > max_row or max_col < start_col:
            continue
        if min_col < start_col:
            continue
        ranges_to_remove.append(merged_range)
    for merged_range in ranges_to_remove:
        try:
            sheet.unmerge_cells(str(merged_range))
        except KeyError:
            try:
                sheet.merged_cells.ranges.remove(merged_range)
            except KeyError:
                pass


def _delete_column_dimensions_from_column(sheet: Worksheet, start_col: int, end_col: int) -> None:
    for key, dimension in list(sheet.column_dimensions.items()):
        min_col = dimension.min or column_index_from_string(key)
        if start_col <= min_col <= end_col:
            del sheet.column_dimensions[key]


def _replace_usage_month_rows(sheet: Worksheet, usage: R13UsageData) -> None:
    target_month = usage.report_month
    existing_rows = _existing_usage_month_rows(sheet, target_month)
    if existing_rows:
        insert_at = existing_rows[0]
        style_template = _capture_row_styles(sheet, insert_at, max_col=13)
        _delete_rows(sheet, existing_rows)
        sheet.insert_rows(insert_at, amount=len(usage.rows))
        target_rows = range(insert_at, insert_at + len(usage.rows))
    else:
        last_used_row = _last_used_row(sheet, max_col=13)
        style_template = _capture_row_styles(sheet, max(last_used_row, 2), max_col=13)
        target_rows = range(last_used_row + 1, last_used_row + len(usage.rows) + 1)

    for row, target_row in zip(usage.rows, target_rows, strict=True):
        _apply_row_styles(sheet, style_template, target_row, max_col=13)
        _set_cell_value(sheet, target_row, 1, target_month)
        _set_cell_value(sheet, target_row, 2, f"{target_month}{row.branch}")
        _set_cell_value(sheet, target_row, 3, row.sequence)
        _set_cell_value(sheet, target_row, 4, row.branch)
        _set_cell_value(sheet, target_row, 5, row.item_code)
        _set_cell_value(sheet, target_row, 6, None)
        _set_cell_value(sheet, target_row, 7, row.item_name)
        _set_cell_value(sheet, target_row, 8, None)
        _set_cell_value(sheet, target_row, 9, None)
        _set_cell_value(sheet, target_row, 10, None)
        _set_cell_value(sheet, target_row, 11, row.capacity)
        _set_cell_value(sheet, target_row, 12, row.quantity)
        _set_cell_value(sheet, target_row, 13, None)


def _write_summary_actuals(sheet: Worksheet, rows: tuple[R13UsageRow, ...], columns: dict[str, int]) -> None:
    quantities = _quantity_by_item_and_branch(rows)
    item_rows = _item_code_rows(sheet)
    item_codes = set(item_rows)
    for item_code in item_codes:
        row_index = item_rows[item_code]
        total = 0.0
        any_value = False
        for branch in SUMMARY_BRANCH_ORDER:
            value = quantities.get((item_code, branch), 0.0)
            if value:
                any_value = True
                total += value
                _set_cell_value(sheet, row_index, columns[branch], value)
            else:
                _set_cell_value(sheet, row_index, columns[branch], None)
        _set_cell_value(sheet, row_index, columns["total"], total if any_value else None)


def _write_branch_actuals(sheet: Worksheet, rows: tuple[R13UsageRow, ...], column_index: int) -> None:
    quantities = _quantity_by_item(rows)
    item_rows = _item_code_rows(sheet)
    for item_code, row_index in item_rows.items():
        value = quantities.get(item_code, 0.0)
        _set_cell_value(sheet, row_index, column_index, value if value else None)


def _copy_actual_values_by_item_code(
    source_sheet: Worksheet,
    source_col: int,
    target_sheet: Worksheet,
    target_col: int,
) -> None:
    source_rows = _item_code_rows(source_sheet)
    target_rows = _item_code_rows(target_sheet)
    for item_code, source_row in source_rows.items():
        target_row = target_rows.get(item_code)
        if target_row is None:
            continue
        _set_cell_value(target_sheet, target_row, target_col, source_sheet.cell(source_row, source_col).value)
        target_sheet.cell(target_row, target_col).number_format = "0"


def _set_all_item_values(sheet: Worksheet, column_index: int, value: float) -> None:
    for row_index in _item_code_rows(sheet).values():
        _set_cell_value(sheet, row_index, column_index, value)
        sheet.cell(row_index, column_index).number_format = "0"


def _write_branch_inventory_values(
    workbook: Any,
    branch_inventory: dict[str, dict[str, float]] | None,
) -> tuple[int, int]:
    if not branch_inventory:
        return 0, 0

    updated = 0
    unmatched = 0
    for sheet_name in R14_BRANCH_SHEETS:
        sheet = workbook[sheet_name]
        values = branch_inventory.get(sheet_name, {})
        if not values:
            unmatched += len(_item_code_rows(sheet))
            continue
        stock_col = _require_branch_stock_column(sheet)
        for item_code, row_index in _item_code_rows(sheet).items():
            if item_code not in values:
                unmatched += 1
                continue
            _set_cell_value(sheet, row_index, stock_col, values[item_code])
            sheet.cell(row_index, stock_col).number_format = "0"
            updated += 1
    return updated, unmatched


def _apply_r14_view_state(workbook: Any, report_date: datetime) -> None:
    summary = workbook["Summary"]
    _apply_summary_actual_outline(summary)
    _apply_summary_planning_outline(summary, report_date.strftime("%Y/%m"))
    for sheet_name in R14_BRANCH_SHEETS:
        _apply_branch_actual_month_visibility(workbook[sheet_name], report_date)


def _apply_summary_actual_outline(sheet: Worksheet) -> None:
    sheet.sheet_properties.outlinePr.summaryRight = False
    for group in _iter_summary_month_groups(sheet):
        total_col = group["total"]
        detail_cols = [group[branch] for branch in SUMMARY_BRANCH_ORDER]
        _set_column_dimension(sheet, total_col, hidden=False, outline_level=0, collapsed=True)
        _group_columns(sheet, detail_cols[0], detail_cols[-1], hidden=True)


def _apply_summary_planning_outline(sheet: Worksheet, month_label: str) -> None:
    forecast_total = _find_summary_forecast_total_column(sheet, month_label)
    if forecast_total is not None:
        forecast_cols = [forecast_total + offset for offset in range(1, len(SUMMARY_BRANCH_ORDER) + 1)]
        _set_column_dimension(sheet, forecast_total, hidden=False, outline_level=0, collapsed=True)
        _group_columns(sheet, forecast_cols[0], forecast_cols[-1], hidden=True)

    safety_total = _find_header_column(sheet, "安庫")
    if safety_total is not None:
        safety_cols = [safety_total + offset for offset in range(1, len(SUMMARY_BRANCH_ORDER) + 1)]
        if safety_cols[-1] <= sheet.max_column:
            _set_column_dimension(sheet, safety_total, hidden=False, outline_level=0, collapsed=True)
            _group_columns(sheet, safety_cols[0], safety_cols[-1], hidden=True)


def _iter_summary_month_groups(sheet: Worksheet) -> list[dict[str, int]]:
    groups: list[dict[str, int]] = []
    for col_index in range(1, sheet.max_column + 1):
        month_label = _cell_month_value(sheet.cell(2, col_index).value)
        if not month_label or sheet.cell(3, col_index).value != "Actual":
            continue
        group = _find_summary_month_group(sheet, month_label)
        if group is not None and group["total"] == col_index:
            groups.append(group)
    return groups


def _apply_branch_actual_month_visibility(sheet: Worksheet, report_date: datetime) -> None:
    sheet.sheet_properties.outlinePr.summaryRight = False
    target = (report_date.year, report_date.month)
    month_columns = _actual_month_columns(sheet)
    future_columns: list[int] = []
    for col_index, month in month_columns:
        if month[0] != target[0]:
            continue
        if month <= target:
            _set_column_dimension(sheet, col_index, hidden=False, outline_level=0, collapsed=month == target)
        else:
            future_columns.append(col_index)

    for start_col, end_col in _contiguous_ranges(future_columns):
        _group_columns(sheet, start_col, end_col, hidden=True)


def _actual_month_columns(sheet: Worksheet) -> list[tuple[int, tuple[int, int]]]:
    columns: list[tuple[int, tuple[int, int]]] = []
    for col_index in range(1, sheet.max_column + 1):
        if sheet.cell(3, col_index).value != "Actual":
            continue
        month = _parse_month_label(sheet.cell(2, col_index).value)
        if month is not None:
            columns.append((col_index, month))
    return columns


def _parse_month_label(value: object) -> tuple[int, int] | None:
    month_label = _cell_month_value(value)
    if not month_label:
        return None
    year, month = (int(part) for part in month_label.split("/"))
    return year, month


def _coerce_excel_datetime(value: object) -> datetime:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, str):
        for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%Y/%m/%d %H:%M:%S"):
            try:
                return datetime.strptime(value.strip(), fmt)
            except ValueError:
                continue
    raise R14TransformError("R14_REPORT_DATE_NOT_FOUND", "R14 報表找不到可解析的 F2 報表日期。")


def _template_report_date(sheet: Worksheet) -> datetime:
    try:
        return _coerce_excel_datetime(sheet["F2"].value)
    except R14TransformError as exc:
        raise R14TransformError(
            "R14_TEMPLATE_INVENTORY_DATE_MISSING",
            "R14 模板 Summary!F2 沒有有效庫存日期；請先執行 W01 從 Google Sheet Summary!G2 同步庫存日期。",
        ) from exc


def _coerce_report_date(value: date | datetime) -> datetime:
    if isinstance(value, datetime):
        return value
    return datetime(value.year, value.month, value.day)


def _numeric_cell_value(value: object) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, bool):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return 0.0


def _set_column_dimension(
    sheet: Worksheet,
    col_index: int,
    *,
    hidden: bool,
    outline_level: int,
    collapsed: bool,
) -> None:
    letter = get_column_letter(col_index)
    dimension = sheet.column_dimensions[letter]
    dimension.min = col_index
    dimension.max = col_index
    dimension.hidden = hidden
    dimension.outlineLevel = outline_level
    dimension.collapsed = collapsed


def _group_columns(sheet: Worksheet, start_col: int, end_col: int, *, hidden: bool) -> None:
    if start_col > end_col:
        return
    start_letter = get_column_letter(start_col)
    end_letter = get_column_letter(end_col)
    sheet.column_dimensions.group(start_letter, end_letter, outline_level=1, hidden=hidden)
    dimension = sheet.column_dimensions[start_letter]
    dimension.min = start_col
    dimension.max = end_col
    dimension.hidden = hidden
    dimension.outlineLevel = 1
    dimension.collapsed = False


def _quantity_by_item_and_branch(rows: tuple[R13UsageRow, ...]) -> dict[tuple[str, str], float]:
    result: dict[tuple[str, str], float] = {}
    for row in rows:
        key = (row.item_code, row.branch)
        result[key] = result.get(key, 0.0) + row.quantity
    return result


def _quantity_by_item(rows: tuple[R13UsageRow, ...]) -> dict[str, float]:
    result: dict[str, float] = {}
    for row in rows:
        result[row.item_code] = result.get(row.item_code, 0.0) + row.quantity
    return result


def _existing_usage_month_rows(sheet: Worksheet, month_label: str) -> list[int]:
    rows: list[int] = []
    for row_index in range(2, _last_used_row(sheet, max_col=13) + 1):
        month = _usage_month_from_row(sheet, row_index)
        if month == month_label:
            rows.append(row_index)
    return rows


def _delete_rows(sheet: Worksheet, row_indexes: list[int]) -> None:
    for start, end in reversed(_contiguous_ranges(row_indexes)):
        sheet.delete_rows(start, end - start + 1)


def _usage_month_from_row(sheet: Worksheet, row_index: int) -> str:
    raw_month = sheet.cell(row_index, 1).value
    month = _cell_month_value(raw_month)
    if month:
        return month
    fx = sheet.cell(row_index, 2).value
    if isinstance(fx, str):
        formula_match = re.search(r'"(\d{4}/\d{2})"', fx)
        if formula_match:
            return formula_match.group(1)
        text_match = re.match(r"(\d{4}/\d{2})", fx)
        if text_match:
            return text_match.group(1)
    return ""


def _set_cell_value(sheet: Worksheet, row_index: int, col_index: int, value: object) -> None:
    cell = sheet.cell(row_index, col_index)
    if isinstance(cell, MergedCell):
        return
    cell.value = value


def _cell_month_value(value: object) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y/%m")
    if isinstance(value, str):
        match = re.search(r"\d{4}/\d{2}", value)
        return match.group(0) if match else ""
    return ""


def _ensure_summary_items(sheet: Worksheet, rows: tuple[R13UsageRow, ...]) -> int:
    existing_codes = _item_code_rows(sheet)
    unique_rows = _unique_rows_by_item_code(rows)
    added = 0
    for item_code, row in unique_rows.items():
        if item_code in existing_codes:
            continue
        target_row = _append_item_row(sheet, row)
        existing_codes[item_code] = target_row
        added += 1
    return added


def _ensure_branch_items(sheet: Worksheet, rows: tuple[R13UsageRow, ...]) -> int:
    existing_codes = _item_code_rows(sheet)
    unique_rows = _unique_rows_by_item_code(rows)
    added = 0
    for item_code, row in unique_rows.items():
        if item_code in existing_codes:
            continue
        target_row = _append_item_row(sheet, row)
        existing_codes[item_code] = target_row
        added += 1
    return added


def _ensure_inventory_items(
    workbook: Any,
    branch_inventory: dict[str, dict[str, float]] | None,
    item_names: dict[str, str],
) -> tuple[int, int]:
    if not branch_inventory:
        return 0, 0

    summary_added = _ensure_sheet_item_codes(
        workbook["Summary"],
        _inventory_item_codes(branch_inventory),
        item_names,
    )
    branch_added = 0
    for sheet_name in R14_BRANCH_SHEETS:
        branch_added += _ensure_sheet_item_codes(
            workbook[sheet_name],
            _inventory_item_codes({sheet_name: branch_inventory.get(sheet_name, {})}),
            item_names,
            branch=sheet_name,
        )
    return summary_added, branch_added


def _ensure_sheet_item_codes(
    sheet: Worksheet,
    item_codes: list[str],
    item_names: dict[str, str],
    *,
    branch: str = "",
) -> int:
    existing_codes = _item_code_rows(sheet)
    added = 0
    for item_code in item_codes:
        if item_code in existing_codes:
            continue
        target_row = _append_item_row(
            sheet,
            R13UsageRow(
                sequence=0,
                branch=branch,
                item_code=item_code,
                item_name=item_names.get(item_code, ""),
                capacity="",
                quantity=0.0,
            ),
        )
        existing_codes[item_code] = target_row
        added += 1
    return added


def _inventory_item_codes(branch_inventory: dict[str, dict[str, float]]) -> list[str]:
    item_codes: list[str] = []
    seen: set[str] = set()
    for sheet_name in R14_BRANCH_SHEETS:
        for item_code in branch_inventory.get(sheet_name, {}):
            normalized = str(item_code).strip()
            if not normalized or normalized in seen:
                continue
            seen.add(normalized)
            item_codes.append(normalized)
    return item_codes


def _unique_rows_by_item_code(rows: tuple[R13UsageRow, ...]) -> dict[str, R13UsageRow]:
    result: dict[str, R13UsageRow] = {}
    for row in rows:
        result.setdefault(row.item_code, row)
    return result


def _item_code_rows(sheet: Worksheet) -> dict[str, int]:
    rows: dict[str, int] = {}
    for row_index in range(4, _last_item_row(sheet) + 1):
        value = sheet.cell(row_index, 2).value
        if value in (None, ""):
            continue
        rows[str(value).strip()] = row_index
    return rows


def _append_item_row(sheet: Worksheet, row: R13UsageRow) -> int:
    source_row = _last_item_row(sheet)
    target_row = source_row + 1
    _copy_row_with_translated_formulas(sheet, source_row, target_row)
    sheet.cell(target_row, 1).value = None
    sheet.cell(target_row, 2).value = row.item_code
    sheet.cell(target_row, 3).value = row.item_name
    sheet.cell(target_row, 4).value = None
    sheet.cell(target_row, 5).value = None
    return target_row


def _refresh_year_total_formulas(sheet: Worksheet, report_date: datetime) -> None:
    year = report_date.year
    target = (report_date.year, report_date.month)
    actual_columns = [
        col_index
        for col_index in range(1, sheet.max_column + 1)
        if (
            (month := _parse_month_label(sheet.cell(2, col_index).value)) is not None
            and month[0] == year
            and month <= target
            and sheet.cell(3, col_index).value == "Actual"
        )
    ]
    if not actual_columns:
        return
    total_col = _find_year_summary_column(sheet, year, "總銷量")
    average_col = _find_year_summary_column(sheet, year, "月均銷")
    if total_col is None and average_col is None:
        return
    for row_index in range(4, _last_item_row(sheet) + 1):
        refs = ",".join(f"{get_column_letter(col_index)}{row_index}" for col_index in actual_columns)
        if total_col is not None:
            _set_cell_value(sheet, row_index, total_col, f"=SUM({refs})")
        if average_col is not None:
            average_function = _existing_average_function(sheet.cell(row_index, average_col).value)
            _set_cell_value(sheet, row_index, average_col, f"={average_function}({refs})")


def _existing_average_function(value: object) -> str:
    if isinstance(value, str) and value.upper().startswith("=AVERAGEA("):
        return "AVERAGEA"
    return "AVERAGE"


def _update_forecast_month_headers(workbook: Any, month_label: str) -> None:
    for sheet_name in ("Summary", *R14_BRANCH_SHEETS):
        sheet = workbook[sheet_name]
        for col_index in range(1, sheet.max_column + 1):
            if sheet.cell(3, col_index).value == "Forecast" and _cell_month_value(sheet.cell(2, col_index).value):
                _set_cell_value(sheet, 2, col_index, month_label)


def _shift_branch_summary_formula_references(workbook: Any, inserted_at: int | None, inserted_width: int) -> None:
    if inserted_at is None or inserted_width <= 0:
        return
    for sheet_name in R14_BRANCH_SHEETS:
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if not (isinstance(value, str) and value.startswith("=") and "Summary!" in value):
                    continue
                cell.value = _shift_summary_column_ranges(value, inserted_at=inserted_at, inserted_width=inserted_width)


def _shift_summary_column_ranges(formula: str, *, inserted_at: int, inserted_width: int) -> str:
    def replace(match: re.Match[str]) -> str:
        prefix, start_col, separator, end_col = match.groups()
        start_index = column_index_from_string(start_col)
        end_index = column_index_from_string(end_col)
        if start_index != end_index or start_index < inserted_at:
            return match.group(0)
        shifted_col = get_column_letter(start_index + inserted_width)
        return f"{prefix}{shifted_col}{separator}{shifted_col}"

    return SUMMARY_COLUMN_RANGE_RE.sub(replace, formula)


def _refresh_branch_planning_columns(
    sheet: Worksheet,
    month_label: str,
    *,
    planning_baseline_available: bool = True,
) -> BranchPlanningColumns:
    actual_col = _require_month_actual_column(sheet, month_label)
    previous_month_label = _previous_month_label(month_label)
    previous_actual_col = _require_month_actual_column(sheet, previous_month_label)
    previous_month_days = _days_in_month_label(previous_month_label)
    forecast_col = _require_branch_forecast_column(sheet, month_label)
    order_col = _ensure_branch_order_column(sheet, forecast_col)
    safety_col = _require_branch_safety_column(sheet, after_col=order_col)
    stock_col = _require_branch_stock_column(sheet)

    for row_index in range(4, _last_item_row(sheet) + 1):
        if not planning_baseline_available:
            _set_cell_value(sheet, row_index, forecast_col, None)
            _set_cell_value(sheet, row_index, order_col, None)
            _set_cell_value(sheet, row_index, safety_col, None)
            continue
        previous_actual_ref = f"{get_column_letter(previous_actual_col)}{row_index}"
        safety_ref = f"{get_column_letter(safety_col)}{row_index}"
        stock_ref = f"{get_column_letter(stock_col)}{row_index}"
        box_capacity_ref = f"D{row_index}"
        daily_average_ref = f"({previous_actual_ref}/{previous_month_days})"
        order_quantity_ref = f"({daily_average_ref}*{R14_ORDER_DAYS}+{safety_ref}-{stock_ref})"
        _set_cell_value(sheet, row_index, forecast_col, f"=CEILING({previous_actual_ref}*1.2,1)")
        _set_cell_value(
            sheet,
            row_index,
            order_col,
            (
                f"=MAX(0,IFERROR(IF({box_capacity_ref}>0,"
                f"CEILING({order_quantity_ref},{box_capacity_ref}),"
                f"CEILING({order_quantity_ref},1)),0))"
            ),
        )
        _set_cell_value(sheet, row_index, safety_col, f"=CEILING({daily_average_ref}*{R14_SAFETY_STOCK_DAYS},1)")
        sheet.cell(row_index, forecast_col).number_format = "0"
        sheet.cell(row_index, order_col).number_format = "0"
        sheet.cell(row_index, safety_col).number_format = "0"

    return BranchPlanningColumns(
        actual_col=actual_col,
        forecast_col=forecast_col,
        order_col=order_col,
        safety_col=safety_col,
        stock_col=stock_col,
    )


def _refresh_summary_planning_columns(
    sheet: Worksheet,
    month_label: str,
    branch_planning: dict[str, BranchPlanningColumns],
    *,
    planning_baseline_available: bool = True,
) -> None:
    forecast_group = _ensure_summary_forecast_group(sheet, month_label)
    safety_group = _ensure_summary_safety_group(sheet)
    if not planning_baseline_available:
        for col_index in (
            forecast_group.total_col,
            *forecast_group.branch_cols.values(),
            safety_group.total_col,
            *safety_group.branch_cols.values(),
        ):
            _clear_item_column_values(sheet, col_index)
        return
    for row_index in range(4, _last_item_row(sheet) + 1):
        _write_summary_metric_formulas(
            sheet,
            row_index=row_index,
            group=forecast_group,
            branch_value_cols={branch: columns.forecast_col for branch, columns in branch_planning.items()},
        )
        _write_summary_metric_formulas(
            sheet,
            row_index=row_index,
            group=safety_group,
            branch_value_cols={branch: columns.safety_col for branch, columns in branch_planning.items()},
        )


def _write_summary_metric_formulas(
    sheet: Worksheet,
    *,
    row_index: int,
    group: SummaryMetricGroup,
    branch_value_cols: dict[str, int],
) -> None:
    for branch in SUMMARY_BRANCH_ORDER:
        source_col = branch_value_cols[branch]
        source_letter = get_column_letter(source_col)
        branch_col = group.branch_cols[branch]
        sheet_ref = _quote_sheet_name(branch)
        _set_cell_value(
            sheet,
            row_index,
            branch_col,
            f"=SUMIF({sheet_ref}!B:B,B{row_index},{sheet_ref}!{source_letter}:{source_letter})",
        )
        sheet.cell(row_index, branch_col).number_format = "0"
    first_col = get_column_letter(group.branch_cols[SUMMARY_BRANCH_ORDER[0]])
    last_col = get_column_letter(group.branch_cols[SUMMARY_BRANCH_ORDER[-1]])
    _set_cell_value(sheet, row_index, group.total_col, f"=SUM({first_col}{row_index}:{last_col}{row_index})")
    sheet.cell(row_index, group.total_col).number_format = "0"


def _write_summary_inventory_formulas(sheet: Worksheet, inventory_columns: dict[str, int], workbook: Any) -> None:
    for branch, summary_col in inventory_columns.items():
        branch_sheet = workbook[branch]
        stock_col = _require_branch_stock_column(branch_sheet)
        stock_letter = get_column_letter(stock_col)
        branch_ref = _quote_sheet_name(branch)
        for row_index in range(4, _last_item_row(sheet) + 1):
            _set_cell_value(
                sheet,
                row_index,
                summary_col,
                f"=SUMIF({branch_ref}!B:B,B{row_index},{branch_ref}!{stock_letter}:{stock_letter})",
            )
            sheet.cell(row_index, summary_col).number_format = "0"


def _write_summary_operational_metric_formulas(
    sheet: Worksheet,
    inventory_columns: dict[str, int],
    metric_columns: dict[str, list[int]],
    branch_planning: dict[str, BranchPlanningColumns],
    report_date: datetime,
    *,
    planning_baseline_available: bool = True,
) -> None:
    report_day = report_date.day
    for offset, branch in enumerate(SUMMARY_BRANCH_ORDER):
        branch_ref = _quote_sheet_name(branch)
        order_col = get_column_letter(branch_planning[branch].order_col)
        actual_col = get_column_letter(branch_planning[branch].actual_col)
        inventory_col = get_column_letter(inventory_columns[branch])
        for row_index in range(4, _last_item_row(sheet) + 1):
            order_summary_col = metric_columns["order"][offset]
            if planning_baseline_available:
                _set_cell_value(
                    sheet,
                    row_index,
                    order_summary_col,
                    f"=SUMIF({branch_ref}!B:B,B{row_index},{branch_ref}!{order_col}:{order_col})",
                )
            else:
                _set_cell_value(sheet, row_index, order_summary_col, None)
            sheet.cell(row_index, order_summary_col).number_format = "0"

            usage_sum = f"SUMIF({branch_ref}!B:B,B{row_index},{branch_ref}!{actual_col}:{actual_col})"
            turnover_summary_col = metric_columns["turnover"][offset]
            _set_cell_value(
                sheet,
                row_index,
                turnover_summary_col,
                f'=IFERROR(IF({usage_sum}=0,"",{inventory_col}{row_index}/({usage_sum}/{report_day})),"")',
            )
            sheet.cell(row_index, turnover_summary_col).number_format = "0.0"


def _ensure_summary_forecast_group(sheet: Worksheet, month_label: str) -> SummaryMetricGroup:
    total_col = _find_summary_forecast_total_column(sheet, month_label)
    if total_col is None:
        raise R14TransformError("R14_SUMMARY_FORECAST_MISSING", f"Summary 找不到 {month_label} Forecast 總欄。")
    branch_cols = _ensure_summary_metric_branch_columns(
        sheet,
        total_col=total_col,
        metric_label="Forecast",
        style_col=total_col + 1 if total_col + 1 <= sheet.max_column else total_col,
    )
    return SummaryMetricGroup(total_col=total_col, branch_cols=branch_cols)


def _ensure_summary_safety_group(sheet: Worksheet) -> SummaryMetricGroup:
    total_col = _find_header_column(sheet, "安庫")
    if total_col is None:
        raise R14TransformError("R14_SUMMARY_SAFETY_STOCK_MISSING", "Summary 找不到安庫欄。")
    _unmerge_cells_in_columns(sheet, total_col, total_col, max_row=3)
    sheet.cell(2, total_col).value = "安庫"
    sheet.cell(3, total_col).value = "安庫"
    branch_cols = _ensure_summary_metric_branch_columns(sheet, total_col=total_col, metric_label="安庫", style_col=total_col)
    return SummaryMetricGroup(total_col=total_col, branch_cols=branch_cols)


def _find_summary_forecast_total_column(sheet: Worksheet, month_label: str) -> int | None:
    for col_index in range(1, sheet.max_column + 1):
        if sheet.cell(3, col_index).value == "Forecast" and _cell_month_value(sheet.cell(2, col_index).value) == month_label:
            return col_index
    return None


def _ensure_summary_metric_branch_columns(
    sheet: Worksheet,
    *,
    total_col: int,
    metric_label: str,
    style_col: int,
) -> dict[str, int]:
    branch_cols: dict[str, int] = {}
    for offset, branch in enumerate(SUMMARY_BRANCH_ORDER, start=1):
        expected_col = total_col + offset
        if not _is_summary_metric_branch_column(sheet, expected_col, branch=branch, metric_label=metric_label):
            header_tail = _capture_header_tail(sheet, expected_col)
            sheet.insert_cols(expected_col)
            _restore_shifted_header_tail(sheet, header_tail, shift=1)
            _unmerge_cells_in_columns(sheet, expected_col, expected_col, max_row=3)
            copy_source_col = min(max(style_col, 1), sheet.max_column)
            _copy_column(sheet, copy_source_col, expected_col)
        _unmerge_cells_in_columns(sheet, expected_col, expected_col, max_row=3)
        sheet.cell(2, expected_col).value = branch
        sheet.cell(3, expected_col).value = metric_label
        _clear_item_column_values(sheet, expected_col)
        branch_cols[branch] = expected_col
    return branch_cols


def _is_summary_metric_branch_column(sheet: Worksheet, col_index: int, *, branch: str, metric_label: str) -> bool:
    if col_index > sheet.max_column:
        return False
    return str(sheet.cell(2, col_index).value or "").strip() == branch and sheet.cell(3, col_index).value == metric_label


def _require_month_actual_column(sheet: Worksheet, month_label: str) -> int:
    col_index = _find_month_actual_column(sheet, month_label)
    if col_index is None:
        raise R14TransformError("R14_TEMPLATE_MONTH_NOT_FOUND", f"頁籤「{sheet.title}」找不到 {month_label} Actual 欄。")
    return col_index


def _require_branch_forecast_column(sheet: Worksheet, month_label: str) -> int:
    for col_index in range(1, sheet.max_column + 1):
        if sheet.cell(3, col_index).value == "Forecast" and _cell_month_value(sheet.cell(2, col_index).value) == month_label:
            return col_index
    raise R14TransformError("R14_BRANCH_FORECAST_MISSING", f"頁籤「{sheet.title}」找不到 {month_label} Forecast 欄。")


def _ensure_branch_order_column(sheet: Worksheet, forecast_col: int) -> int:
    order_col = forecast_col + 1
    if _is_header_column(sheet, order_col, "下單數"):
        return order_col

    header_tail = _capture_header_tail(sheet, order_col)
    sheet.insert_cols(order_col)
    _restore_shifted_header_tail(sheet, header_tail, shift=1)
    style_source_col = _find_header_column(sheet, "安庫") or order_col + 1
    _unmerge_cells_in_columns(sheet, order_col, order_col, max_row=3)
    _copy_column(sheet, style_source_col, order_col)
    _unmerge_cells_in_columns(sheet, order_col, order_col, max_row=3)
    sheet.cell(2, order_col).value = "下單數"
    sheet.cell(3, order_col).value = None
    sheet.merge_cells(start_row=2, start_column=order_col, end_row=3, end_column=order_col)
    _clear_item_column_values(sheet, order_col)
    return order_col


def _require_branch_safety_column(sheet: Worksheet, *, after_col: int) -> int:
    for col_index in range(after_col + 1, sheet.max_column + 1):
        if _is_header_column(sheet, col_index, "安庫"):
            return col_index
    raise R14TransformError("R14_BRANCH_SAFETY_STOCK_MISSING", f"頁籤「{sheet.title}」找不到安庫欄。")


def _require_branch_stock_column(sheet: Worksheet) -> int:
    sheet_name = _normalized_header_text(sheet.title)
    fallback: int | None = None
    for col_index in range(1, sheet.max_column + 1):
        text = _normalized_header_text(sheet.cell(3, col_index).value)
        if "庫存" not in text or text == "庫存單位":
            continue
        if sheet_name in text:
            return col_index
        fallback = fallback or col_index
    if fallback is not None:
        return fallback
    raise R14TransformError("R14_BRANCH_STOCK_COLUMN_MISSING", f"頁籤「{sheet.title}」找不到該分館庫存欄。")


def _require_header_column(sheet: Worksheet, label: str) -> int:
    col_index = _find_header_column(sheet, label)
    if col_index is None:
        raise R14TransformError("R14_HEADER_COLUMN_MISSING", f"頁籤「{sheet.title}」找不到表頭：{label}")
    return col_index


def _find_header_column(sheet: Worksheet, label: str) -> int | None:
    for col_index in range(1, sheet.max_column + 1):
        if _is_header_column(sheet, col_index, label):
            return col_index
    return None


def _is_header_column(sheet: Worksheet, col_index: int, label: str) -> bool:
    normalized_label = _normalized_header_text(label)
    return any(_normalized_header_text(sheet.cell(row_index, col_index).value) == normalized_label for row_index in (1, 2, 3))


def _normalized_header_text(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip()


def _quote_sheet_name(sheet_name: str) -> str:
    return "'" + sheet_name.replace("'", "''") + "'"


def _find_year_summary_column(sheet: Worksheet, year: int, label: str) -> int | None:
    for col_index in range(1, sheet.max_column + 1):
        text = str(sheet.cell(2, col_index).value or "")
        if str(year) in text and label in text:
            return col_index
    return None


def _cell_year_value(value: object) -> int | None:
    if isinstance(value, datetime):
        return value.year
    if isinstance(value, str):
        match = re.search(r"\d{4}", value)
        return int(match.group(0)) if match else None
    return None


def _last_item_row(sheet: Worksheet) -> int:
    for row_index in range(sheet.max_row, 3, -1):
        if sheet.cell(row_index, 2).value not in (None, "") or sheet.cell(row_index, 3).value not in (None, ""):
            return row_index
    return 4


def _last_used_row(sheet: Worksheet, *, max_col: int) -> int:
    for row_index in range(sheet.max_row, 0, -1):
        if any(sheet.cell(row_index, col_index).value not in (None, "") for col_index in range(1, max_col + 1)):
            return row_index
    return 1


def _contiguous_ranges(indexes: list[int]) -> list[tuple[int, int]]:
    if not indexes:
        return []
    ranges: list[tuple[int, int]] = []
    start = previous = sorted(indexes)[0]
    for index in sorted(indexes)[1:]:
        if index == previous + 1:
            previous = index
            continue
        ranges.append((start, previous))
        start = previous = index
    ranges.append((start, previous))
    return ranges


def _capture_row_styles(sheet: Worksheet, source_row: int, *, max_col: int) -> list[dict[str, object]]:
    styles: list[dict[str, object]] = []
    for col_index in range(1, max_col + 1):
        source = sheet.cell(source_row, col_index)
        styles.append(
            {
                "style": copy(source._style) if source.has_style else None,
                "number_format": source.number_format,
                "alignment": copy(source.alignment) if source.alignment else None,
            }
        )
    return styles


def _apply_row_styles(sheet: Worksheet, styles: list[dict[str, object]], target_row: int, *, max_col: int) -> None:
    for col_index in range(1, max_col + 1):
        target = sheet.cell(target_row, col_index)
        if isinstance(target, MergedCell):
            continue
        style = styles[col_index - 1]
        if style["style"] is not None:
            target._style = copy(style["style"])
        if style["number_format"]:
            target.number_format = str(style["number_format"])
        if style["alignment"] is not None:
            target.alignment = copy(style["alignment"])


def _copy_row_styles(sheet: Worksheet, source_row: int, target_row: int, *, max_col: int) -> None:
    if source_row == target_row:
        return
    for col_index in range(1, max_col + 1):
        source = sheet.cell(source_row, col_index)
        target = sheet.cell(target_row, col_index)
        if isinstance(target, MergedCell):
            continue
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)


def _copy_column(sheet: Worksheet, source_col: int, target_col: int) -> None:
    source_letter = get_column_letter(source_col)
    target_letter = get_column_letter(target_col)
    sheet.column_dimensions[target_letter].width = sheet.column_dimensions[source_letter].width
    for row_index in range(1, sheet.max_row + 1):
        source = sheet.cell(row_index, source_col)
        target = sheet.cell(row_index, target_col)
        if isinstance(target, MergedCell):
            continue
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.fill:
            target.fill = copy(source.fill)
        if source.border:
            target.border = copy(source.border)


def _clear_item_column_values(sheet: Worksheet, col_index: int) -> None:
    for row_index in range(4, _last_item_row(sheet) + 1):
        _set_cell_value(sheet, row_index, col_index, None)


def _unmerge_cells_in_columns(sheet: Worksheet, start_col: int, end_col: int, *, max_row: int) -> None:
    ranges_to_remove = []
    for merged_range in sheet.merged_cells.ranges:
        min_col, min_row, max_col, _max_row = range_boundaries(str(merged_range))
        if min_row > max_row:
            continue
        if max_col < start_col or min_col > end_col:
            continue
        ranges_to_remove.append(merged_range)
    for merged_range in ranges_to_remove:
        try:
            sheet.unmerge_cells(str(merged_range))
        except KeyError:
            try:
                sheet.merged_cells.ranges.remove(merged_range)
            except KeyError:
                pass


def _copy_row_with_translated_formulas(sheet: Worksheet, source_row: int, target_row: int) -> None:
    _copy_row_styles(sheet, source_row, target_row, max_col=sheet.max_column)
    for col_index in range(1, sheet.max_column + 1):
        source = sheet.cell(source_row, col_index)
        target = sheet.cell(target_row, col_index)
        if isinstance(target, MergedCell):
            continue
        value = source.value
        if isinstance(value, str) and value.startswith("="):
            origin = f"{get_column_letter(col_index)}{source_row}"
            destination = f"{get_column_letter(col_index)}{target_row}"
            target.value = Translator(value, origin=origin).translate_formula(destination)
        else:
            target.value = None
