from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from pos_report_bot.reports.w02_order_builder import build_w02_order_plan


def test_build_w02_order_plan_uses_r14_order_quantity_column_and_groups_by_department(tmp_path: Path) -> None:
    r14_path = _write_r14_fixture(tmp_path)

    plan = build_w02_order_plan(
        r14_path,
        item_departments={"6150001": "護理部"},
        branch_sheet_names=["站前4樓"],
    )

    assert plan.report_date == datetime(2026, 7, 14)
    assert len(plan.forms) == 1
    form = plan.forms[0]
    assert form.branch == "站前4樓"
    assert form.department == "護理部"
    assert len(form.items) == 1
    assert form.items[0].item_code == "6150001"
    assert form.items[0].item_name == "測試商品"
    assert form.items[0].quantity == 42
    assert plan.issues == ()


def test_build_w02_order_plan_evaluates_r14_order_formula_when_cached_value_is_missing(tmp_path: Path) -> None:
    r14_path = _write_r14_fixture(tmp_path, order_formula=True)

    plan = build_w02_order_plan(
        r14_path,
        item_departments={"6150001": "護理部"},
        branch_sheet_names=["站前4樓"],
    )

    assert len(plan.forms) == 1
    assert plan.forms[0].items[0].quantity == 40


def test_build_w02_order_plan_preserves_legacy_r14_day_formula_support(tmp_path: Path) -> None:
    r14_path = _write_r14_fixture(tmp_path, order_formula=True)
    workbook = load_workbook(r14_path)
    sheet = workbook["站前4樓"]
    sheet["I4"] = 14
    sheet["H4"] = "=MAX(0,IFERROR(IF(D4>0,CEILING(((G4/DAY($F$2))*14+I4-F4),D4),CEILING(((G4/DAY($F$2))*14+I4-F4),1)),0))"
    workbook.save(r14_path)

    plan = build_w02_order_plan(
        r14_path,
        item_departments={"6150001": "護理部"},
        branch_sheet_names=["站前4樓"],
    )

    assert len(plan.forms) == 1
    assert plan.forms[0].items[0].quantity == 60


def test_build_w02_order_plan_uses_filename_report_date_when_inventory_date_is_different_month(tmp_path: Path) -> None:
    r14_path = _write_r14_fixture(
        tmp_path,
        filename="診所stock status - 2026 demand planning-0714.xlsx",
        summary_date=datetime(2026, 6, 25),
    )

    plan = build_w02_order_plan(
        r14_path,
        item_departments={"6150001": "護理部"},
        branch_sheet_names=["站前4樓"],
    )

    assert plan.report_date == datetime(2026, 7, 14)
    assert len(plan.forms) == 1
    assert plan.forms[0].items[0].quantity == 42


def test_build_w02_order_plan_uses_previous_month_daily_average_not_inventory_date_day(tmp_path: Path) -> None:
    r14_path = _write_r14_fixture(
        tmp_path,
        filename="診所stock status - 2026 demand planning-0701.xlsx",
        summary_date=datetime(2026, 6, 25),
        order_formula=True,
        branch_f2_formula=True,
        actual_value=1,
        stock_value=2,
        box_capacity=1,
    )

    plan = build_w02_order_plan(
        r14_path,
        item_departments={"6150001": "護理部"},
        branch_sheet_names=["站前4樓"],
    )

    assert plan.report_date == datetime(2026, 7, 1)
    assert plan.forms == ()
    assert plan.issues == ()


def test_build_w02_order_plan_uses_previous_month_days_for_nested_branch_formula(tmp_path: Path) -> None:
    r14_path = _write_r14_fixture(
        tmp_path,
        filename="診所stock status - 2026 demand planning-0701.xlsx",
        summary_date=datetime(2026, 6, 25),
        order_formula=True,
        branch_f2_formula=True,
        nested_branch_f2_formula=True,
        actual_value=85,
        stock_value=0,
        box_capacity=50,
    )

    plan = build_w02_order_plan(
        r14_path,
        item_departments={"6150099": "護理部"},
        branch_sheet_names=["站前11樓"],
    )

    assert len(plan.forms) == 1
    assert plan.forms[0].items[0].quantity == 100


def test_build_w02_order_plan_skips_positive_quantity_item_without_department(tmp_path: Path) -> None:
    r14_path = _write_r14_fixture(tmp_path)

    plan = build_w02_order_plan(
        r14_path,
        item_departments={},
        known_item_codes={"6150001"},
        branch_sheet_names=["站前4樓"],
    )

    assert plan.forms == ()
    assert len(plan.issues) == 1
    assert plan.issues[0].branch == "站前4樓"
    assert plan.issues[0].item_code == "6150001"
    assert plan.issues[0].quantity == 42
    assert "F 欄部門空白" in plan.issues[0].reason


def test_build_w02_order_plan_reports_missing_google_sheet_item_code(tmp_path: Path) -> None:
    r14_path = _write_r14_fixture(tmp_path)

    plan = build_w02_order_plan(
        r14_path,
        item_departments={},
        known_item_codes=set(),
        branch_sheet_names=["站前4樓"],
    )

    assert plan.forms == ()
    assert len(plan.issues) == 1
    assert plan.issues[0].item_code == "6150001"
    assert "B 欄找不到" in plan.issues[0].reason


def test_build_w02_order_plan_only_uses_configured_branch_sheets(tmp_path: Path) -> None:
    r14_path = _write_r14_fixture(tmp_path)

    plan = build_w02_order_plan(
        r14_path,
        item_departments={"6150001": "護理部", "6150099": "護理部"},
        known_item_codes={"6150001", "6150099"},
        branch_sheet_names=["站前4樓"],
    )

    assert [form.branch for form in plan.forms] == ["站前4樓"]
    assert all(item.item_code != "6150099" for form in plan.forms for item in form.items)


def test_build_w02_order_plan_preserves_configured_branch_order(tmp_path: Path) -> None:
    r14_path = _write_r14_fixture(tmp_path)

    plan = build_w02_order_plan(
        r14_path,
        item_departments={"6150001": "護理部", "6150099": "護理部"},
        known_item_codes={"6150001", "6150099"},
        branch_sheet_names=["站前11樓", "站前4樓"],
    )

    assert [form.branch for form in plan.forms] == ["站前11樓", "站前4樓"]


def _write_r14_fixture(
    tmp_path: Path,
    *,
    filename: str = "r14.xlsx",
    summary_date: datetime = datetime(2026, 7, 14),
    order_formula: bool = False,
    branch_f2_formula: bool = False,
    nested_branch_f2_formula: bool = False,
    actual_value: int = 28,
    stock_value: int = 5,
    box_capacity: int = 20,
) -> Path:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["F2"] = summary_date

    branch = workbook.create_sheet("站前4樓")
    branch["F2"] = "=Summary!F2" if branch_f2_formula else summary_date
    branch.cell(2, 7).value = "2026/07"
    branch.cell(3, 2).value = "凱惠料號"
    branch.cell(3, 3).value = "品名"
    branch.cell(3, 4).value = "盒入數"
    branch.cell(2, 5).value = "2026/06"
    branch.cell(3, 5).value = "Actual"
    branch.cell(3, 6).value = "站前4樓庫存"
    branch.cell(3, 7).value = "Actual"
    branch.cell(2, 8).value = "下單數"
    branch.cell(2, 9).value = "安庫"
    branch.cell(4, 2).value = "6150001"
    branch.cell(4, 3).value = "測試商品"
    branch.cell(4, 4).value = box_capacity
    branch.cell(4, 5).value = actual_value
    branch.cell(4, 6).value = stock_value
    branch.cell(4, 7).value = actual_value
    branch.cell(4, 8).value = (
        "=MAX(0,IFERROR(IF(D4>0,CEILING(((E4/30)*21+I4-F4),D4),CEILING(((E4/30)*21+I4-F4),1)),0))"
        if order_formula
        else 42
    )
    branch.cell(4, 9).value = "=CEILING((E4/30)*14,1)"
    branch.cell(5, 2).value = "6150002"
    branch.cell(5, 3).value = "不用下單商品"
    branch.cell(5, 4).value = 10
    branch.cell(5, 5).value = 1
    branch.cell(5, 6).value = 99
    branch.cell(5, 7).value = 1
    branch.cell(5, 8).value = (
        "=MAX(0,IFERROR(IF(D5>0,CEILING(((E5/30)*21+I5-F5),D5),CEILING(((E5/30)*21+I5-F5),1)),0))"
        if order_formula
        else 0
    )
    branch.cell(5, 9).value = "=CEILING((E5/30)*14,1)"

    other_branch = workbook.create_sheet("站前11樓")
    other_branch["F2"] = "=站前4樓!F2" if nested_branch_f2_formula else summary_date
    other_branch.cell(2, 7).value = "2026/07"
    other_branch.cell(3, 2).value = "凱惠料號"
    other_branch.cell(3, 3).value = "品名"
    other_branch.cell(3, 4).value = "盒入數"
    other_branch.cell(2, 5).value = "2026/06"
    other_branch.cell(3, 5).value = "Actual"
    other_branch.cell(3, 6).value = "站前11樓庫存"
    other_branch.cell(3, 7).value = "Actual"
    other_branch.cell(2, 8).value = "下單數"
    other_branch.cell(4, 2).value = "6150099"
    other_branch.cell(4, 3).value = "其他分館商品"
    other_branch.cell(4, 4).value = box_capacity
    other_branch.cell(4, 5).value = actual_value
    other_branch.cell(4, 6).value = stock_value
    other_branch.cell(4, 7).value = actual_value
    other_branch.cell(4, 8).value = (
        "=MAX(0,IFERROR(IF(D4>0,CEILING(((E4/30)*21+I4-F4),D4),CEILING(((E4/30)*21+I4-F4),1)),0))"
        if order_formula
        else 12
    )
    other_branch.cell(4, 9).value = "=CEILING((E4/30)*14,1)"

    r14_path = tmp_path / filename
    workbook.save(r14_path)
    return r14_path
