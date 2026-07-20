from __future__ import annotations

import calendar
from dataclasses import dataclass
from datetime import datetime
import math
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from pos_report_bot.reports.r14_transformer import R14_BRANCH_SHEETS, R14TransformError

R14_OUTPUT_REPORT_DATE_RE = re.compile(r"(?P<year>\d{4})\s+demand planning-(?P<mmdd>\d{4})(?:_\d+)?\.xlsx$", re.IGNORECASE)
R14_ORDER_DAYS = 21
R14_SAFETY_STOCK_DAYS = 14


@dataclass(frozen=True)
class W02OrderItem:
    branch: str
    department: str
    item_code: str
    item_name: str
    quantity: int


@dataclass(frozen=True)
class W02OrderIssue:
    branch: str
    item_code: str
    item_name: str
    quantity: int | None
    reason: str


@dataclass(frozen=True)
class W02OrderForm:
    branch: str
    department: str
    items: tuple[W02OrderItem, ...]


@dataclass(frozen=True)
class W02OrderPlan:
    r14_path: Path
    report_date: datetime
    forms: tuple[W02OrderForm, ...]
    issues: tuple[W02OrderIssue, ...]

    @property
    def order_item_count(self) -> int:
        return sum(len(form.items) for form in self.forms)


def build_w02_order_plan(
    r14_path: Path,
    *,
    item_departments: dict[str, str],
    known_item_codes: set[str] | None = None,
    branch_sheet_names: list[str] | tuple[str, ...] | None = None,
) -> W02OrderPlan:
    workbook = load_workbook(r14_path, data_only=False)
    if "Summary" not in workbook.sheetnames:
        raise R14TransformError("W02_R14_SUMMARY_MISSING", "W02 找不到 R14 報表 Summary 頁籤。")
    report_date = _r14_report_date_from_filename(r14_path.name) or _coerce_report_datetime(workbook["Summary"]["F2"].value)
    report_month = report_date.strftime("%Y/%m")
    branch_names = _candidate_branch_sheet_names(workbook.sheetnames, branch_sheet_names)
    forms: dict[tuple[str, str], list[W02OrderItem]] = {}
    issues: list[W02OrderIssue] = []

    for branch in branch_names:
        sheet = workbook[branch]
        columns = _branch_columns(sheet, report_month)
        for row_index in range(4, _last_item_row(sheet) + 1):
            item_code = _cell_text(sheet.cell(row_index, 2).value)
            item_name = _cell_text(sheet.cell(row_index, 3).value)
            if not item_code and not item_name:
                continue
            quantity = _calculate_order_quantity(sheet, row_index, columns, report_date)
            if quantity <= 0:
                continue
            department = item_departments.get(item_code, "").strip()
            if not department:
                issues.append(
                    W02OrderIssue(
                        branch=branch,
                        item_code=item_code,
                        item_name=item_name,
                        quantity=quantity,
                        reason=_missing_department_reason(item_code, known_item_codes),
                    )
                )
                continue
            forms.setdefault((branch, department), []).append(
                W02OrderItem(
                    branch=branch,
                    department=department,
                    item_code=item_code,
                    item_name=item_name,
                    quantity=quantity,
                )
            )

    branch_order = {branch: index for index, branch in enumerate(branch_names)}
    ordered_forms = tuple(
        W02OrderForm(branch=branch, department=department, items=tuple(items))
        for (branch, department), items in sorted(
            forms.items(),
            key=lambda item: (branch_order.get(item[0][0], 9999), item[0][1]),
        )
    )
    return W02OrderPlan(
        r14_path=r14_path,
        report_date=report_date,
        forms=ordered_forms,
        issues=tuple(issues),
    )


def _missing_department_reason(item_code: str, known_item_codes: set[str] | None) -> str:
    if known_item_codes is not None and item_code not in known_item_codes:
        return (
            "Google Sheet Summary B 欄找不到此凱惠料號，已跳過。"
            "請先在庫存 Google Sheet Summary 新增此料號，並在 F 欄填入 POS 部門後重跑 W02。"
        )
    return (
        "Google Sheet Summary 已有此凱惠料號，但 F 欄部門空白，已跳過。"
        "請在 F 欄填入與 POS 部門下拉選單完全一致的部門名稱後重跑 W02。"
    )


def _candidate_branch_sheet_names(
    workbook_sheet_names: list[str],
    configured_branch_names: list[str] | tuple[str, ...] | None,
) -> list[str]:
    candidates = list(configured_branch_names) if configured_branch_names is not None else list(R14_BRANCH_SHEETS)
    return [
        sheet_name
        for sheet_name in candidates
        if sheet_name in workbook_sheet_names and sheet_name not in {"Summary", "領用表"}
    ]


def _branch_columns(sheet: Worksheet, report_month: str) -> dict[str, int]:
    actual_col = _find_month_actual_column(sheet, report_month)
    if actual_col is None:
        raise R14TransformError("W02_R14_ACTUAL_COLUMN_MISSING", f"W02 找不到 {sheet.title} 的 {report_month} Actual 欄。")
    previous_actual_col = _find_month_actual_column(sheet, _previous_month_label(report_month))
    if previous_actual_col is None:
        raise R14TransformError(
            "W02_R14_PREVIOUS_ACTUAL_COLUMN_MISSING",
            f"W02 找不到 {sheet.title} 的 {_previous_month_label(report_month)} Actual 欄。",
        )
    order_col = _find_branch_order_column(sheet)
    if order_col is None:
        raise R14TransformError("W02_R14_ORDER_COLUMN_MISSING", f"W02 找不到 {sheet.title} 的下單數欄。")
    return {"actual": actual_col, "previous_actual": previous_actual_col, "order": order_col}


def _calculate_order_quantity(sheet: Worksheet, row_index: int, columns: dict[str, int], report_date: datetime) -> int:
    order_cell_value = sheet.cell(row_index, columns["order"]).value
    if not _is_formula(order_cell_value):
        return int(math.ceil(max(0.0, _numeric_cell_value(order_cell_value))))
    elapsed_days = _r14_formula_elapsed_days(sheet, fallback_days=report_date.day)
    evaluated = _evaluate_r14_order_formula(sheet, row_index, str(order_cell_value), elapsed_days)
    if evaluated is not None:
        return evaluated
    previous_actual_col = columns.get("previous_actual") or columns["actual"]
    previous_actual = _numeric_cell_value(sheet.cell(row_index, previous_actual_col).value)
    stock_col = _find_branch_stock_column(sheet)
    if stock_col is None:
        raise R14TransformError("W02_R14_STOCK_COLUMN_MISSING", f"W02 找不到 {sheet.title} 的庫存欄。")
    stock = _numeric_cell_value(sheet.cell(row_index, stock_col).value)
    box_capacity = _numeric_cell_value(sheet.cell(row_index, 4).value)
    return _r14_order_quantity_formula_result(
        actual=previous_actual,
        stock=stock,
        box_capacity=box_capacity,
        average_days=_days_in_previous_month(report_date),
        order_days=R14_ORDER_DAYS,
        safety_days=R14_SAFETY_STOCK_DAYS,
    )


def _evaluate_r14_order_formula(sheet: Worksheet, row_index: int, formula: str, elapsed_days: int) -> int | None:
    refs = _parse_r14_order_formula_refs(formula, row_index)
    if refs is None:
        return None
    actual = _numeric_cell_value(sheet[f"{refs['actual']}{row_index}"].value)
    stock = _numeric_cell_value(sheet[f"{refs['stock']}{row_index}"].value)
    box_capacity = _numeric_cell_value(sheet[f"{refs['box']}{row_index}"].value)
    return _r14_order_quantity_formula_result(
        actual=actual,
        stock=stock,
        box_capacity=box_capacity,
        average_days=int(refs.get("average_days") or elapsed_days),
        order_days=int(refs.get("order_days") or 14),
        safety_days=R14_SAFETY_STOCK_DAYS,
    )


def _r14_formula_elapsed_days(sheet: Worksheet, *, fallback_days: int) -> int:
    value = _resolve_workbook_cell_value(sheet, sheet["F2"].value)
    try:
        date_value = _coerce_report_datetime(value)
    except R14TransformError:
        return max(int(fallback_days), 1)
    return max(int(date_value.day), 1)


def _resolve_workbook_cell_value(sheet: Worksheet, value: object, *, depth: int = 0) -> object:
    if depth > 5 or not _is_formula(value):
        return value
    reference = str(value).strip()[1:]
    match = re.fullmatch(r"(?:'(?P<quoted>[^']+)'|(?P<sheet>[^!]+))!\$?(?P<cell>[A-Z]{1,3})\$?(?P<row>\d+)", reference)
    if match is None:
        return value
    sheet_name = match.group("quoted") or match.group("sheet")
    cell_ref = f"{match.group('cell')}{match.group('row')}"
    workbook = sheet.parent
    if sheet_name not in workbook.sheetnames:
        return value
    target_value = workbook[sheet_name][cell_ref].value
    return _resolve_workbook_cell_value(workbook[sheet_name], target_value, depth=depth + 1)


def _parse_r14_order_formula_refs(formula: str, row_index: int) -> dict[str, str] | None:
    row = str(row_index)
    normalized = re.sub(r"\s+", "", formula.upper())
    box_match = re.search(rf"IF\(\$?(?P<box>[A-Z]{{1,3}})\$?{row}>0", normalized)
    quantity_match = re.search(
        rf"\(\(\$?(?P<actual>[A-Z]{{1,3}})\$?{row}/(?P<average_days>\d+)\)\*(?P<order_days>\d+)"
        rf"\+\$?(?P<safety>[A-Z]{{1,3}})\$?{row}-\$?(?P<stock>[A-Z]{{1,3}})\$?{row}\)",
        normalized,
    )
    if quantity_match is None:
        quantity_match = re.search(
            rf"\(\(\$?(?P<actual>[A-Z]{{1,3}})\$?{row}/DAY\(\$F\$2\)\)\*14"
            rf"\+\$?(?P<safety>[A-Z]{{1,3}})\$?{row}-\$?(?P<stock>[A-Z]{{1,3}})\$?{row}\)",
            normalized,
        )
    if box_match is None or quantity_match is None:
        return None
    return {
        "actual": quantity_match.group("actual"),
        "safety": quantity_match.group("safety"),
        "stock": quantity_match.group("stock"),
        "box": box_match.group("box"),
        "average_days": quantity_match.groupdict().get("average_days") or "",
        "order_days": quantity_match.groupdict().get("order_days") or "",
    }


def _r14_order_quantity_formula_result(
    *,
    actual: float,
    stock: float,
    box_capacity: float,
    average_days: int,
    order_days: int,
    safety_days: int,
) -> int:
    average_days = max(int(average_days), 1)
    daily_average = actual / average_days
    safety_stock = math.ceil(daily_average * safety_days)
    raw_quantity = max(0.0, (daily_average * order_days) + safety_stock - stock)
    if raw_quantity <= 0:
        return 0
    if box_capacity > 0:
        return int(math.ceil(raw_quantity / box_capacity) * box_capacity)
    return int(math.ceil(raw_quantity))


def _find_month_actual_column(sheet: Worksheet, report_month: str) -> int | None:
    for col_index in range(1, sheet.max_column + 1):
        if _cell_month_value(sheet.cell(2, col_index).value) == report_month and _cell_text(sheet.cell(3, col_index).value) == "Actual":
            return col_index
    return None


def _previous_month_label(month_label: str) -> str:
    year, month = (int(part) for part in month_label.split("/"))
    if month == 1:
        return f"{year - 1}/12"
    return f"{year}/{month - 1:02d}"


def _days_in_previous_month(report_date: datetime) -> int:
    year = report_date.year
    month = report_date.month - 1
    if month == 0:
        year -= 1
        month = 12
    return calendar.monthrange(year, month)[1]


def _find_branch_stock_column(sheet: Worksheet) -> int | None:
    normalized_sheet = _normalized_text(sheet.title)
    fallback: int | None = None
    for col_index in range(1, sheet.max_column + 1):
        text = _normalized_text(sheet.cell(3, col_index).value)
        if "庫存" not in text or text == "庫存單位":
            continue
        if normalized_sheet in text:
            return col_index
        fallback = fallback or col_index
    return fallback


def _find_branch_order_column(sheet: Worksheet) -> int | None:
    for col_index in range(1, sheet.max_column + 1):
        if _normalized_text(sheet.cell(2, col_index).value) == "下單數" or _normalized_text(sheet.cell(3, col_index).value) == "下單數":
            return col_index
    return None


def _is_formula(value: object) -> bool:
    return isinstance(value, str) and value.strip().startswith("=")


def _last_item_row(sheet: Worksheet) -> int:
    last = 3
    for row_index in range(4, sheet.max_row + 1):
        if _cell_text(sheet.cell(row_index, 2).value) or _cell_text(sheet.cell(row_index, 3).value):
            last = row_index
    return last


def _coerce_report_datetime(value: object) -> datetime:
    if isinstance(value, datetime):
        return value
    if hasattr(value, "year") and hasattr(value, "month") and hasattr(value, "day"):
        return datetime(int(value.year), int(value.month), int(value.day))
    text = _cell_text(value)
    for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    raise R14TransformError("W02_R14_REPORT_DATE_MISSING", "W02 找不到 R14 Summary!F2 的有效日期。")


def _r14_report_date_from_filename(filename: str) -> datetime | None:
    match = R14_OUTPUT_REPORT_DATE_RE.search(filename)
    if match is None:
        return None
    year = int(match.group("year"))
    mmdd = match.group("mmdd")
    try:
        return datetime(year, int(mmdd[:2]), int(mmdd[2:]))
    except ValueError:
        return None


def _cell_month_value(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.strftime("%Y/%m")
    text = _cell_text(value).replace("\n", " ")
    match = re.search(r"(20\d{2})[/.-](\d{1,2})", text)
    if not match:
        return None
    return f"{int(match.group(1)):04d}/{int(match.group(2)):02d}"


def _numeric_cell_value(value: object) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    text = _cell_text(value).replace(",", "")
    if not text or text.startswith("="):
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalized_text(value: object) -> str:
    return re.sub(r"\s+", "", _cell_text(value))
