from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from pos_report_bot.reports.r14_transformer import (
    R14TransformError,
    _days_in_month_label,
    load_r14_workbook_snapshot,
    parse_r13_usage_summary,
    sync_r14_template_actual_month_state,
    sync_r14_template_inventory,
    transform_r13_to_r14,
)


ROOT = Path(__file__).resolve().parents[2]
FIXTURE_DIR = ROOT / "tests" / "R14_TEST"
RAW_PATH = FIXTURE_DIR / "診所stock status - 2026 demand planning-0609-rawdata.xls"
TEMPLATE_PATH = FIXTURE_DIR / "診所stock status - 2026 demand planning-0531.xlsx"


def test_parse_r13_usage_summary_reads_bottom_summary_block() -> None:
    usage = parse_r13_usage_summary(RAW_PATH)

    assert usage.start_date.date() == date(2026, 6, 1)
    assert usage.end_date.date() == date(2026, 6, 8)
    assert usage.report_month == "2026/06"
    assert len(usage.rows) == 208
    assert usage.rows[0].branch == "站前4樓"
    assert usage.rows[0].item_code == "6050010"
    assert usage.rows[0].item_name == "(V)ExoPower 泌力機泌安瓶 6mL"
    assert usage.rows[0].capacity == "瓶"
    assert usage.rows[0].quantity == 1
    assert all(row.branch != "合計:" for row in usage.rows)


def test_r14_previous_month_day_count_handles_month_boundaries() -> None:
    assert _days_in_month_label("2026/01") == 31
    assert _days_in_month_label("2026/02") == 28
    assert _days_in_month_label("2028/02") == 29
    assert _days_in_month_label("2026/12") == 31


def test_transform_r13_to_r14_rebuilds_usage_sheet_and_actual_columns(tmp_path: Path) -> None:
    output_path = tmp_path / "診所stock status - 2026 demand planning-0608.xlsx"

    result = transform_r13_to_r14(
        RAW_PATH,
        TEMPLATE_PATH,
        output_path,
        expected_end_date=date(2026, 6, 8),
    )

    assert result.output_path == output_path
    assert result.report_month == "2026/06"
    assert result.report_date.date() == date(2026, 5, 31)
    assert result.imported_rows == 208
    assert result.added_summary_items == 2
    assert result.added_branch_items == 3
    assert output_path.exists()
    assert output_path.stat().st_size > 0

    workbook = load_workbook(output_path, data_only=False)
    assert workbook.sheetnames == [
        "Summary",
        "站前4樓",
        "站前11樓",
        "忠孝國際醫學3樓",
        "忠孝7樓",
        "忠孝健康7樓",
        "領用表",
    ]

    usage_sheet = workbook["領用表"]
    usage_rows = [row for row in range(2, usage_sheet.max_row + 1) if usage_sheet.cell(row, 1).value == "2026/06"]
    assert len(usage_rows) == 208
    assert [usage_sheet.cell(usage_rows[0], col).value for col in range(1, 14)] == [
        "2026/06",
        "2026/06站前4樓",
        1,
        "站前4樓",
        "6050010",
        None,
        "(V)ExoPower 泌力機泌安瓶 6mL",
        None,
        None,
        None,
        "瓶",
        1,
        None,
    ]

    summary = workbook["Summary"]
    assert summary["F2"].value.date() == date(2026, 5, 31)
    assert [(summary.cell(2, col).value, summary.cell(3, col).value) for col in range(121, 127)] == [
        ("2026/06", "Actual"),
        ("站前4樓", "Actual"),
        ("站前11樓", "Actual"),
        ("忠孝7樓", "Actual"),
        ("忠孝國際醫學3樓", "Actual"),
        ("忠孝健康7樓", "Actual"),
    ]
    assert summary.column_dimensions["DQ"].hidden is False
    assert summary.column_dimensions["DQ"].collapsed is True
    assert summary.column_dimensions["DR"].hidden is True
    assert summary.column_dimensions["DR"].outlineLevel == 1
    assert summary.column_dimensions["DR"].min == 122
    assert summary.column_dimensions["DR"].max == 126
    assert "DW2:DW3" in _merged_ranges(summary, min_col=121)
    assert "DX2:DX3" in _merged_ranges(summary, min_col=121)
    assert summary["DW2"].value == "2026\n總銷量\n(單支/條/點)"
    assert summary["DX2"].value == "2026\n月均銷\n(單支/條/點)"

    forecast_group = _find_summary_metric_group(summary, total_label="2026/06", metric_label="Forecast")
    safety_group = _find_summary_metric_group(summary, total_label="安庫", metric_label="安庫")
    assert [(summary.cell(2, col).value, summary.cell(3, col).value) for col in forecast_group] == [
        ("2026/06", "Forecast"),
        ("站前4樓", "Forecast"),
        ("站前11樓", "Forecast"),
        ("忠孝7樓", "Forecast"),
        ("忠孝國際醫學3樓", "Forecast"),
        ("忠孝健康7樓", "Forecast"),
    ]
    assert [(summary.cell(2, col).value, summary.cell(3, col).value) for col in safety_group] == [
        ("安庫", "安庫"),
        ("站前4樓", "安庫"),
        ("站前11樓", "安庫"),
        ("忠孝7樓", "安庫"),
        ("忠孝國際醫學3樓", "安庫"),
        ("忠孝健康7樓", "安庫"),
    ]
    forecast_total_letter = _col_letter(forecast_group[0])
    forecast_first_detail = _col_letter(forecast_group[1])
    forecast_last_detail = _col_letter(forecast_group[-1])
    assert summary.column_dimensions[forecast_total_letter].hidden is False
    assert summary.column_dimensions[forecast_total_letter].collapsed is True
    assert summary.column_dimensions[forecast_first_detail].hidden is True
    assert summary.column_dimensions[forecast_first_detail].outlineLevel == 1
    assert summary.column_dimensions[forecast_first_detail].min == forecast_group[1]
    assert summary.column_dimensions[forecast_first_detail].max == forecast_group[-1]
    safety_total_letter = _col_letter(safety_group[0])
    safety_first_detail = _col_letter(safety_group[1])
    assert summary.column_dimensions[safety_total_letter].hidden is False
    assert summary.column_dimensions[safety_total_letter].collapsed is True
    assert summary.column_dimensions[safety_first_detail].hidden is True
    assert summary.column_dimensions[safety_first_detail].outlineLevel == 1
    assert summary.column_dimensions[safety_first_detail].min == safety_group[1]
    assert summary.column_dimensions[safety_first_detail].max == safety_group[-1]
    assert _merged_ranges_in_column(summary, forecast_group[0], min_row=4) == []
    assert _merged_ranges_in_column(summary, safety_group[0], min_row=4) == []

    summary_row = _find_item_row(summary, "6050010")
    assert [summary.cell(summary_row, col).value for col in range(121, 127)] == [1, 1, None, None, None, None]
    assert summary.cell(summary_row, 127).value == f"=SUM(CN{summary_row},CS{summary_row},CY{summary_row},DE{summary_row},DK{summary_row},DQ{summary_row})"
    assert summary.cell(summary_row, 128).value == (
        f"=AVERAGEA(CN{summary_row},CS{summary_row},CY{summary_row},DE{summary_row},DK{summary_row},DQ{summary_row})"
    )
    assert _find_item_row(summary, "6110181") > 0
    assert summary.cell(summary_row, forecast_group[0]).value == (
        f"=SUM({forecast_first_detail}{summary_row}:{forecast_last_detail}{summary_row})"
    )
    branch_planning_columns: dict[str, tuple[int, int]] = {}
    for sheet_name in ("站前4樓", "站前11樓", "忠孝7樓", "忠孝國際醫學3樓", "忠孝健康7樓"):
        branch = workbook[sheet_name]
        branch_forecast_col = _find_forecast_month_column(branch)
        branch_planning_columns[sheet_name] = (branch_forecast_col, branch_forecast_col + 2)
    for branch_index, sheet_name in enumerate(("站前4樓", "站前11樓", "忠孝7樓", "忠孝國際醫學3樓", "忠孝健康7樓"), start=1):
        branch_forecast_col, _branch_safety_col = branch_planning_columns[sheet_name]
        branch_forecast_letter = _col_letter(branch_forecast_col)
        assert summary.cell(summary_row, forecast_group[branch_index]).value == (
            f"=SUMIF('{sheet_name}'!B:B,B{summary_row},'{sheet_name}'!{branch_forecast_letter}:{branch_forecast_letter})"
        )
    safety_first_detail_letter = _col_letter(safety_group[1])
    safety_last_detail_letter = _col_letter(safety_group[-1])
    assert summary.cell(summary_row, safety_group[0]).value == (
        f"=SUM({safety_first_detail_letter}{summary_row}:{safety_last_detail_letter}{summary_row})"
    )
    for branch_index, sheet_name in enumerate(("站前4樓", "站前11樓", "忠孝7樓", "忠孝國際醫學3樓", "忠孝健康7樓"), start=1):
        _branch_forecast_col, branch_safety_col = branch_planning_columns[sheet_name]
        branch_safety_letter = _col_letter(branch_safety_col)
        assert summary.cell(summary_row, safety_group[branch_index]).value == (
            f"=SUMIF('{sheet_name}'!B:B,B{summary_row},'{sheet_name}'!{branch_safety_letter}:{branch_safety_letter})"
        )

    branch_sheet = workbook["站前4樓"]
    branch_row = _find_item_row(branch_sheet, "6050010")
    branch_col = _find_month_column(branch_sheet, "2026/06")
    assert branch_sheet.cell(branch_row, branch_col).value == 1
    assert branch_sheet["BW2"].value == "2026/06"
    assert branch_sheet.cell(branch_row, 73).value == (
        f"=SUM(BI{branch_row},BJ{branch_row},BK{branch_row},BL{branch_row},BM{branch_row},BN{branch_row})"
    )
    assert branch_sheet.cell(branch_row, 74).value == (
        f"=AVERAGE(BI{branch_row},BJ{branch_row},BK{branch_row},BL{branch_row},BM{branch_row},BN{branch_row})"
    )
    forecast_col = _find_forecast_month_column(branch_sheet)
    order_col = forecast_col + 1
    safety_col = order_col + 1
    previous_actual_col = _find_month_column(branch_sheet, "2026/05")
    assert branch_sheet.cell(2, order_col).value == "下單數"
    assert branch_sheet.cell(2, safety_col).value == "安庫"
    assert branch_sheet.cell(branch_row, forecast_col).value == (
        f"=CEILING({_col_letter(previous_actual_col)}{branch_row}*1.2,1)"
    )
    assert branch_sheet.cell(branch_row, order_col).value == (
        f"=MAX(0,IFERROR(IF(D{branch_row}>0,"
        f"CEILING((({_col_letter(previous_actual_col)}{branch_row}/31)*21+BY{branch_row}-F{branch_row}),D{branch_row}),"
        f"CEILING((({_col_letter(previous_actual_col)}{branch_row}/31)*21+BY{branch_row}-F{branch_row}),1)),0))"
    )
    assert branch_sheet.cell(branch_row, safety_col).value == (
        f"=CEILING(({_col_letter(previous_actual_col)}{branch_row}/31)*14,1)"
    )

    expected_visibility = {
        "站前4樓": ("BM", "BN", "BO", 67, 72),
        "站前11樓": ("BM", "BN", "BO", 67, 72),
        "忠孝國際醫學3樓": ("T", "U", "V", 22, 27),
        "忠孝7樓": ("V", "W", "X", 24, 29),
        "忠孝健康7樓": ("T", "U", "V", 22, 27),
    }
    for sheet_name, (previous_month, current_month, future_group, group_min, group_max) in expected_visibility.items():
        branch_sheet = workbook[sheet_name]
        assert branch_sheet.column_dimensions[previous_month].collapsed is False
        assert branch_sheet.column_dimensions[current_month].hidden is False
        assert branch_sheet.column_dimensions[current_month].collapsed is True
        assert branch_sheet.column_dimensions[future_group].hidden is True
        assert branch_sheet.column_dimensions[future_group].outlineLevel == 1
        assert branch_sheet.column_dimensions[future_group].min == group_min
        assert branch_sheet.column_dimensions[future_group].max == group_max
        forecast_col = _find_forecast_month_column(branch_sheet)
        assert branch_sheet.cell(2, forecast_col).value == "2026/06"

    for sheet_name in ("站前4樓", "站前11樓", "忠孝國際醫學3樓", "忠孝7樓", "忠孝健康7樓"):
        branch_sheet = workbook[sheet_name]
        item_row = _first_item_row(branch_sheet)
        previous_actual_col = _find_month_column(branch_sheet, "2026/05")
        forecast_col = _find_forecast_month_column(branch_sheet)
        order_col = forecast_col + 1
        safety_col = order_col + 1
        stock_col = _find_branch_stock_column(branch_sheet, sheet_name)
        assert branch_sheet.cell(2, order_col).value == "下單數"
        assert branch_sheet.cell(2, safety_col).value == "安庫"
        assert branch_sheet.cell(item_row, forecast_col).value == (
            f"=CEILING({_col_letter(previous_actual_col)}{item_row}*1.2,1)"
        )
        assert branch_sheet.cell(item_row, order_col).value == (
            f"=MAX(0,IFERROR(IF(D{item_row}>0,"
            f"CEILING((({_col_letter(previous_actual_col)}{item_row}/31)*21+"
            f"{_col_letter(safety_col)}{item_row}-{_col_letter(stock_col)}{item_row}),D{item_row}),"
            f"CEILING((({_col_letter(previous_actual_col)}{item_row}/31)*21+"
            f"{_col_letter(safety_col)}{item_row}-{_col_letter(stock_col)}{item_row}),1)),0))"
        )
        assert branch_sheet.cell(item_row, safety_col).value == (
            f"=CEILING(({_col_letter(previous_actual_col)}{item_row}/31)*14,1)"
        )

    assert workbook["站前4樓"].column_dimensions["H"].hidden is True
    assert workbook["站前4樓"].column_dimensions["H"].max == 44
    assert workbook["站前4樓"].column_dimensions["AU"].hidden is True
    assert workbook["站前4樓"].column_dimensions["AU"].max == 49
    assert workbook["站前4樓"].column_dimensions["AX"].hidden is True
    assert workbook["站前4樓"].column_dimensions["AY"].hidden is True
    assert workbook["站前11樓"].column_dimensions["H"].hidden is True
    assert workbook["站前11樓"].column_dimensions["H"].max == 44
    for hidden_col in ("H", "I", "J", "K"):
        assert workbook["忠孝健康7樓"].column_dimensions[hidden_col].hidden is True


def test_transform_r13_to_r14_writes_weekly_google_sheet_inventory(tmp_path: Path) -> None:
    output_path = tmp_path / "with-inventory.xlsx"

    result = transform_r13_to_r14(
        RAW_PATH,
        TEMPLATE_PATH,
        output_path,
        expected_end_date=date(2026, 6, 8),
        branch_inventory={
            "站前4樓": {"6050010": 123},
            "站前11樓": {"6120001": 45},
            "忠孝7樓": {"6110181": 67},
            "忠孝國際醫學3樓": {"6120001": 89},
            "忠孝健康7樓": {"6120004": 10},
        },
        inventory_date=date(2026, 6, 9),
    )

    workbook = load_workbook(output_path, data_only=False)
    assert workbook["Summary"]["F2"].value.date() == date(2026, 6, 9)
    expected = {
        "站前4樓": ("6050010", 123),
        "站前11樓": ("6120001", 45),
        "忠孝7樓": ("6110181", 67),
        "忠孝國際醫學3樓": ("6120001", 89),
        "忠孝健康7樓": ("6120004", 10),
    }
    for sheet_name, (item_code, inventory_value) in expected.items():
        branch_sheet = workbook[sheet_name]
        branch_row = _find_item_row(branch_sheet, item_code)
        stock_col = _find_branch_stock_column(branch_sheet, sheet_name)
        forecast_col = _find_forecast_month_column(branch_sheet)
        order_col = forecast_col + 1
        stock_letter = _col_letter(stock_col)

        assert branch_sheet.cell(branch_row, stock_col).value == inventory_value
        assert branch_sheet.cell(branch_row, order_col).value.endswith(
            f"-{stock_letter}{branch_row}),1)),0))"
        )

    assert _find_branch_stock_column(workbook["忠孝7樓"], "忠孝7樓") != 6
    assert result.inventory_updated == 5
    assert result.inventory_unmatched > 0


def test_transform_r13_to_r14_appends_google_sheet_only_inventory_items(tmp_path: Path) -> None:
    output_path = tmp_path / "with-new-inventory-item.xlsx"

    result = transform_r13_to_r14(
        RAW_PATH,
        TEMPLATE_PATH,
        output_path,
        expected_end_date=date(2026, 6, 8),
        branch_inventory={
            "站前4樓": {"NEWGOOGLE001": 12},
            "站前11樓": {},
            "忠孝7樓": {"NEWGOOGLE001": 0},
            "忠孝國際醫學3樓": {},
            "忠孝健康7樓": {},
        },
        branch_inventory_item_names={"NEWGOOGLE001": "Google Sheet 新增品項"},
        inventory_date=date(2026, 6, 9),
    )

    workbook = load_workbook(output_path, data_only=False)
    summary_row = _find_item_row(workbook["Summary"], "NEWGOOGLE001")
    assert workbook["Summary"].cell(summary_row, 3).value == "Google Sheet 新增品項"
    assert workbook["Summary"]["F2"].value.date() == date(2026, 6, 9)

    for sheet_name, expected_value in {"站前4樓": 12, "忠孝7樓": 0}.items():
        sheet = workbook[sheet_name]
        row = _find_item_row(sheet, "NEWGOOGLE001")
        stock_col = _find_branch_stock_column(sheet, sheet_name)
        forecast_col = _find_forecast_month_column(sheet)
        assert sheet.cell(row, 3).value == "Google Sheet 新增品項"
        assert sheet.cell(row, stock_col).value == expected_value
        assert str(sheet.cell(row, forecast_col).value).startswith("=CEILING(")

    for sheet_name in ("站前11樓", "忠孝國際醫學3樓", "忠孝健康7樓"):
        with pytest.raises(AssertionError):
            _find_item_row(workbook[sheet_name], "NEWGOOGLE001")

    assert result.added_summary_items >= 1
    assert result.added_branch_items >= 2
    assert result.inventory_updated == 2


def test_sync_r14_template_inventory_updates_existing_template_file(tmp_path: Path) -> None:
    template_path = tmp_path / "r14-template.xlsx"
    template_path.write_bytes(TEMPLATE_PATH.read_bytes())

    result = sync_r14_template_inventory(
        template_path,
        {
            "站前4樓": {"6050010": 321},
            "站前11樓": {"6120001": 654},
        },
    )

    workbook = load_workbook(template_path, data_only=False)
    expected = {
        "站前4樓": ("6050010", 321),
        "站前11樓": ("6120001", 654),
    }
    for sheet_name, (item_code, expected_value) in expected.items():
        sheet = workbook[sheet_name]
        row = _find_item_row(sheet, item_code)
        stock_col = _find_branch_stock_column(sheet, sheet_name)
        assert sheet.cell(row, stock_col).value == expected_value
    assert result.template_path == template_path
    assert result.inventory_updated == 2


def test_sync_r14_template_inventory_writes_date_and_appends_google_sheet_only_items(tmp_path: Path) -> None:
    template_path = tmp_path / "r14-template.xlsx"
    template_path.write_bytes(TEMPLATE_PATH.read_bytes())

    result = sync_r14_template_inventory(
        template_path,
        {
            "站前4樓": {"NEWGOOGLE002": 5},
            "站前11樓": {},
            "忠孝7樓": {},
            "忠孝國際醫學3樓": {},
            "忠孝健康7樓": {"NEWGOOGLE002": 0},
        },
        item_names={"NEWGOOGLE002": "庫存表新增品項"},
        inventory_date=date(2026, 6, 25),
    )

    workbook = load_workbook(template_path, data_only=False)
    summary_row = _find_item_row(workbook["Summary"], "NEWGOOGLE002")
    assert workbook["Summary"]["F2"].value.date() == date(2026, 6, 25)
    assert workbook["Summary"].cell(summary_row, 3).value == "庫存表新增品項"
    for sheet_name, expected_value in {"站前4樓": 5, "忠孝健康7樓": 0}.items():
        sheet = workbook[sheet_name]
        row = _find_item_row(sheet, "NEWGOOGLE002")
        stock_col = _find_branch_stock_column(sheet, sheet_name)
        assert sheet.cell(row, 3).value == "庫存表新增品項"
        assert sheet.cell(row, stock_col).value == expected_value
    assert result.added_summary_items == 1
    assert result.added_branch_items == 2
    assert result.inventory_updated == 2


def test_transform_preserves_template_inventory_date_when_w01_is_not_running(tmp_path: Path) -> None:
    refreshed_template = tmp_path / "refreshed-template.xlsx"
    refreshed_template.write_bytes(TEMPLATE_PATH.read_bytes())
    workbook = load_workbook(refreshed_template)
    workbook["Summary"]["F2"] = date(2026, 6, 25)
    workbook.save(refreshed_template)
    output_path = tmp_path / "daily-r14-after-w01.xlsx"

    result = transform_r13_to_r14(
        RAW_PATH,
        refreshed_template,
        output_path,
        expected_end_date=date(2026, 6, 8),
    )

    workbook = load_workbook(output_path, data_only=False)
    assert result.report_date.date() == date(2026, 6, 25)
    assert workbook["Summary"]["F2"].value.date() == date(2026, 6, 25)


def test_transform_can_persist_month_state_to_runtime_template(tmp_path: Path) -> None:
    runtime_template = tmp_path / "runtime-template.xlsx"
    runtime_template.write_bytes(TEMPLATE_PATH.read_bytes())
    output_path = tmp_path / "daily-r14.xlsx"

    transform_r13_to_r14(
        RAW_PATH,
        runtime_template,
        output_path,
        expected_end_date=date(2026, 6, 8),
        update_template_path=runtime_template,
    )

    output_workbook = load_workbook(output_path, data_only=False)
    runtime_workbook = load_workbook(runtime_template, data_only=False)
    output_summary = output_workbook["Summary"]
    runtime_summary = runtime_workbook["Summary"]
    output_group = _find_summary_metric_group(output_summary, total_label="2026/06", metric_label="Actual")
    runtime_group = _find_summary_metric_group(runtime_summary, total_label="2026/06", metric_label="Actual")
    item_row = _find_item_row(runtime_summary, "6050010")

    assert runtime_group == output_group
    assert runtime_summary.cell(item_row, runtime_group[0]).value == output_summary.cell(item_row, output_group[0]).value


def test_sync_r14_template_actual_month_state_overwrites_existing_previous_month_values(tmp_path: Path) -> None:
    runtime_template = tmp_path / "runtime-template.xlsx"
    runtime_template.write_bytes(TEMPLATE_PATH.read_bytes())
    source_report = tmp_path / "source-report.xlsx"
    transform_r13_to_r14(
        RAW_PATH,
        runtime_template,
        source_report,
        expected_end_date=date(2026, 6, 8),
        update_template_path=runtime_template,
    )
    workbook = load_workbook(runtime_template)
    summary_group = _find_summary_metric_group(workbook["Summary"], total_label="2026/06", metric_label="Actual")
    summary_row = _find_item_row(workbook["Summary"], "6050010")
    workbook["Summary"].cell(summary_row, summary_group[0]).value = 999
    branch = workbook["站前4樓"]
    branch_col = _find_month_column(branch, "2026/06")
    branch_row = _find_item_row(branch, "6050010")
    branch.cell(branch_row, branch_col).value = 999
    workbook.save(runtime_template)

    result = sync_r14_template_actual_month_state(
        runtime_template,
        source_report,
        "2026/06",
        overwrite_values=True,
    )

    workbook = load_workbook(runtime_template, data_only=False)
    summary = workbook["Summary"]
    summary_group = _find_summary_metric_group(summary, total_label="2026/06", metric_label="Actual")
    branch = workbook["站前4樓"]
    branch_col = _find_month_column(branch, "2026/06")
    assert result.updated_columns == 11
    assert summary.cell(summary_row, summary_group[0]).value == 1
    assert branch.cell(branch_row, branch_col).value == 1


def test_transform_requires_template_inventory_date_when_w01_is_not_running(tmp_path: Path) -> None:
    bad_template = tmp_path / "bad-template-date.xlsx"
    bad_template.write_bytes(TEMPLATE_PATH.read_bytes())
    workbook = load_workbook(bad_template)
    workbook["Summary"]["F2"] = None
    workbook.save(bad_template)

    with pytest.raises(R14TransformError) as error:
        transform_r13_to_r14(
            RAW_PATH,
            bad_template,
            tmp_path / "bad-template-date-output.xlsx",
            expected_end_date=date(2026, 6, 8),
        )

    assert error.value.error_code == "R14_TEMPLATE_INVENTORY_DATE_MISSING"


def test_transform_replaces_existing_usage_month_instead_of_appending(tmp_path: Path) -> None:
    first_output = tmp_path / "first.xlsx"
    transform_r13_to_r14(RAW_PATH, TEMPLATE_PATH, first_output, expected_end_date=date(2026, 6, 8))

    workbook = load_workbook(first_output)
    usage_sheet = workbook["領用表"]
    fake_row = usage_sheet.max_row + 1
    usage_sheet.cell(fake_row, 1).value = "2026/06"
    usage_sheet.cell(fake_row, 4).value = "站前4樓"
    usage_sheet.cell(fake_row, 5).value = "FAKE999"
    workbook.save(first_output)

    second_output = tmp_path / "second.xlsx"
    transform_r13_to_r14(RAW_PATH, first_output, second_output, expected_end_date=date(2026, 6, 8))

    workbook = load_workbook(second_output, data_only=False)
    usage_sheet = workbook["領用表"]
    usage_rows = [row for row in range(2, usage_sheet.max_row + 1) if usage_sheet.cell(row, 1).value == "2026/06"]
    item_codes = {str(usage_sheet.cell(row, 5).value) for row in usage_rows}
    assert len(usage_rows) == 208
    assert "FAKE999" not in item_codes

    summary = workbook["Summary"]
    assert len(_find_summary_metric_group(summary, total_label="2026/06", metric_label="Forecast")) == 6
    assert len(_find_summary_metric_group(summary, total_label="安庫", metric_label="安庫")) == 6
    for sheet_name in ("站前4樓", "站前11樓", "忠孝國際醫學3樓", "忠孝7樓", "忠孝健康7樓"):
        sheet = workbook[sheet_name]
        assert _count_header_columns(sheet, "下單數") == 1


def test_transform_rejects_raw_file_with_unexpected_end_date(tmp_path: Path) -> None:
    with pytest.raises(R14TransformError) as error:
        transform_r13_to_r14(
            RAW_PATH,
            TEMPLATE_PATH,
            tmp_path / "bad.xlsx",
            expected_end_date=date(2026, 6, 9),
        )

    assert error.value.error_code == "R14_SOURCE_DATE_MISMATCH"


def test_transform_allows_inventory_date_from_different_month(tmp_path: Path) -> None:
    output_path = tmp_path / "cross-month-inventory-date.xlsx"

    result = transform_r13_to_r14(
        RAW_PATH,
        TEMPLATE_PATH,
        output_path,
        expected_end_date=date(2026, 6, 8),
        inventory_date=date(2026, 5, 31),
    )

    workbook = load_workbook(output_path, data_only=False)
    assert result.report_month == "2026/06"
    assert result.report_date.date() == date(2026, 5, 31)
    assert workbook["Summary"]["F2"].value.date() == date(2026, 5, 31)


def test_r14_snapshot_uses_filename_report_date_when_inventory_date_is_different_month(tmp_path: Path) -> None:
    output_path = tmp_path / "診所stock status - 2026 demand planning-0608.xlsx"
    transform_r13_to_r14(
        RAW_PATH,
        TEMPLATE_PATH,
        output_path,
        expected_end_date=date(2026, 6, 8),
        inventory_date=date(2026, 5, 31),
    )

    snapshot = load_r14_workbook_snapshot(output_path)
    station_front_item = next(item for item in snapshot.items if item.branch == "站前4樓" and item.item_code == "6050010")

    assert snapshot.report_date.date() == date(2026, 6, 8)
    assert snapshot.report_month == "2026/06"
    assert station_front_item.actual == 1


def _find_item_row(sheet, item_code: str) -> int:  # type: ignore[no-untyped-def]
    for row in range(4, sheet.max_row + 1):
        if str(sheet.cell(row, 2).value).strip() == item_code:
            return row
    raise AssertionError(f"item not found: {item_code}")


def _first_item_row(sheet) -> int:  # type: ignore[no-untyped-def]
    for row in range(4, sheet.max_row + 1):
        if sheet.cell(row, 2).value not in (None, ""):
            return row
    raise AssertionError("first item row not found")


def _find_month_column(sheet, month_label: str) -> int:  # type: ignore[no-untyped-def]
    for col in range(1, sheet.max_column + 1):
        if sheet.cell(2, col).value == month_label and sheet.cell(3, col).value == "Actual":
            return col
    raise AssertionError(f"month column not found: {month_label}")


def _find_forecast_month_column(sheet):  # type: ignore[no-untyped-def]
    for col in range(1, sheet.max_column + 1):
        if sheet.cell(3, col).value == "Forecast" and str(sheet.cell(2, col).value).startswith("2026/"):
            return col
    raise AssertionError("forecast month column not found")


def _find_summary_metric_group(sheet, *, total_label: str, metric_label: str) -> list[int]:  # type: ignore[no-untyped-def]
    for col in range(1, sheet.max_column + 1):
        total_cell = sheet.cell(2, col).value
        if total_label == "安庫":
            total_matches = str(total_cell).strip() == total_label and sheet.cell(3, col).value == metric_label
        else:
            total_matches = total_cell == total_label and sheet.cell(3, col).value == metric_label
        if not total_matches:
            continue
        columns = [col]
        for branch in ("站前4樓", "站前11樓", "忠孝7樓", "忠孝國際醫學3樓", "忠孝健康7樓"):
            next_col = columns[-1] + 1
            assert sheet.cell(2, next_col).value == branch
            assert sheet.cell(3, next_col).value == metric_label
            columns.append(next_col)
        return columns
    raise AssertionError(f"summary metric group not found: {total_label} {metric_label}")


def _find_header_column(sheet, label: str) -> int:  # type: ignore[no-untyped-def]
    for col in range(1, sheet.max_column + 1):
        if any(_normal_header(sheet.cell(row, col).value) == _normal_header(label) for row in (1, 2, 3)):
            return col
    raise AssertionError(f"header not found: {label}")


def _find_branch_stock_column(sheet, sheet_name: str) -> int:  # type: ignore[no-untyped-def]
    normalized_sheet = _normal_header(sheet_name)
    fallback = None
    for col in range(1, sheet.max_column + 1):
        text = _normal_header(sheet.cell(3, col).value)
        if "庫存" not in text or text == "庫存單位":
            continue
        if normalized_sheet in text:
            return col
        fallback = fallback or col
    if fallback is not None:
        return fallback
    raise AssertionError(f"stock column not found: {sheet_name}")


def _count_header_columns(sheet, label: str) -> int:  # type: ignore[no-untyped-def]
    return sum(
        1
        for col in range(1, sheet.max_column + 1)
        if any(_normal_header(sheet.cell(row, col).value) == _normal_header(label) for row in (1, 2, 3))
    )


def _normal_header(value) -> str:  # type: ignore[no-untyped-def]
    return "".join(str(value or "").split())


def _col_letter(col: int) -> str:
    return get_column_letter(col)


def _merged_ranges(sheet, *, min_col: int) -> list[str]:  # type: ignore[no-untyped-def]
    return sorted(
        str(cell_range)
        for cell_range in sheet.merged_cells.ranges
        if cell_range.min_row <= 3 and cell_range.min_col >= min_col
    )


def _merged_ranges_in_column(sheet, col: int, *, min_row: int) -> list[str]:  # type: ignore[no-untyped-def]
    return sorted(
        str(cell_range)
        for cell_range in sheet.merged_cells.ranges
        if cell_range.max_row >= min_row and cell_range.min_col <= col <= cell_range.max_col
    )
