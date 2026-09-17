from datetime import date, datetime, timedelta
import json
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook
from openpyxl import load_workbook

from pos_report_bot.app.automation_runner import AutomationRunner
from pos_report_bot.config.loader import load_project_config
from pos_report_bot.reports.planner import build_dry_run_plan
from pos_report_bot.reports.w02_order_builder import W02OrderIssue


ROOT = Path(__file__).resolve().parents[2]


class FakeW02DepartmentClient:
    def __init__(self, departments: dict[str, str], known_item_codes: set[str] | None = None) -> None:
        self.departments = departments
        self.known_item_codes = known_item_codes if known_item_codes is not None else set(departments)

    def read_w02_departments(self, _settings):  # type: ignore[no-untyped-def]
        return SimpleNamespace(
            item_departments=self.departments,
            known_item_codes=self.known_item_codes,
            rows_read=3,
        )


class FakeGmailSender:
    sent: list[dict[str, object]] = []

    def __init__(self, _config) -> None:  # type: ignore[no-untyped-def]
        return

    def send(self, settings, *, subject: str, body: str, attachments=None):  # type: ignore[no-untyped-def]
        self.__class__.sent.append(
            {
                "recipients": list(settings.recipients),
                "subject": subject,
                "body": body,
                "attachments": list(attachments or []),
            }
        )
        return SimpleNamespace(ok=True, message="sent")


class FakeW02PosOrderAutomator:
    calls: list[dict[str, object]] = []
    result = SimpleNamespace(
        ok=True,
        message="fake W02 POS orders completed",
        actions=["fake_w02_pos_submit"],
        completed_forms=1,
        skipped_forms=0,
        error_code=None,
        diagnostic_path=None,
        skipped_issues=(),
        completed_form_counts_by_branch={},
    )

    def __init__(self, window, **kwargs) -> None:  # type: ignore[no-untyped-def]
        self.window = window
        self.kwargs = kwargs

    def submit_plan(self, plan):  # type: ignore[no-untyped-def]
        self.__class__.calls.append(
            {
                "window": self.window,
                "form_count": len(plan.forms),
                "item_count": plan.order_item_count,
                "item_codes": [
                    item.item_code
                    for form in plan.forms
                    for item in form.items
                ],
                "kwargs": self.kwargs,
            }
        )
        return self.__class__.result


class FakeReadyPosWindow:
    def __init__(self, name: str = "SPA-POS N001-站前4樓") -> None:
        self.name = name
        self.closed = False
        self._children = [
            SimpleNamespace(window_text=lambda: "常用表單", children=lambda: []),
            SimpleNamespace(window_text=lambda: "維護設定", children=lambda: []),
            SimpleNamespace(window_text=lambda: "統計報表", children=lambda: []),
            SimpleNamespace(window_text=lambda: "庫存管理", children=lambda: []),
        ]

    def window_text(self) -> str:
        return self.name

    def children(self) -> list[object]:
        return self._children

    def close(self) -> None:
        self.closed = True

    def is_visible(self) -> bool:
        return not self.closed


def test_w02_issues_are_skipped_and_advance_next_run_date_when_no_order_forms(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    _write_r14_for_w02(tmp_path, run_date=date(2026, 7, 3), positive_order=True)
    FakeGmailSender.sent = []
    runner = _w02_runner(config, tmp_path, departments={}, run_date=date(2026, 7, 3))
    output, planned_outputs = _w02_output(config, date(2026, 7, 3))

    result = runner._run_w02_order_workflow(output, planned_outputs, prior_failures=[])

    assert result.ok is True
    assert result.error_code is None
    assert config.w02_order.next_run_date == "2026/07/17"
    assert len(FakeGmailSender.sent) == 1
    assert "W02 下單異常品項" in str(FakeGmailSender.sent[0]["body"])
    assert "B 欄找不到" in str(FakeGmailSender.sent[0]["body"])
    assert "w02_skipped_items:1" in result.actions


def test_w02_issues_report_blank_department_and_advance_when_no_order_forms(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    _write_r14_for_w02(tmp_path, run_date=date(2026, 7, 3), positive_order=True)
    FakeGmailSender.sent = []
    runner = _w02_runner(
        config,
        tmp_path,
        departments={},
        known_item_codes={"6150001"},
        run_date=date(2026, 7, 3),
    )
    output, planned_outputs = _w02_output(config, date(2026, 7, 3))

    result = runner._run_w02_order_workflow(output, planned_outputs, prior_failures=[])

    assert result.ok is True
    assert result.error_code is None
    assert config.w02_order.next_run_date == "2026/07/17"
    assert len(FakeGmailSender.sent) == 1
    assert "F 欄部門空白" in str(FakeGmailSender.sent[0]["body"])


def test_w02_issues_do_not_block_normal_forms_but_pos_disabled_keeps_date(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    _write_r14_for_w02(
        tmp_path,
        run_date=date(2026, 7, 3),
        positive_order=True,
        extra_positive_order=True,
    )
    FakeGmailSender.sent = []
    runner = _w02_runner(
        config,
        tmp_path,
        departments={"6150002": "護理部"},
        known_item_codes={"6150001", "6150002"},
        run_date=date(2026, 7, 3),
    )
    output, planned_outputs = _w02_output(config, date(2026, 7, 3))

    result = runner._run_w02_order_workflow(output, planned_outputs, prior_failures=[])

    assert result.ok is False
    assert result.error_code == "W02_POS_SUBMISSION_DISABLED"
    assert config.w02_order.next_run_date == "2026/07/03"
    assert len(FakeGmailSender.sent) == 0
    assert "異常品項已記錄但尚未寄送通知" in result.message


def test_w02_email_disabled_fails_when_issues_exist(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    config.email.enabled = False
    _write_r14_for_w02(tmp_path, run_date=date(2026, 7, 3), positive_order=True)
    runner = _w02_runner(config, tmp_path, departments={}, run_date=date(2026, 7, 3))
    output, planned_outputs = _w02_output(config, date(2026, 7, 3))

    result = runner._run_w02_order_workflow(output, planned_outputs, prior_failures=[])

    assert result.ok is False
    assert result.error_code == "W02_EMAIL_DISABLED"
    assert config.w02_order.next_run_date == "2026/07/03"


def test_w02_locked_canonical_plan_uses_recovery_plan_path(
    tmp_path: Path,
    monkeypatch,
) -> None:  # type: ignore[no-untyped-def]
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    _write_r14_for_w02(tmp_path, run_date=date(2026, 7, 3), positive_order=False)
    runner = _w02_runner(config, tmp_path, departments={"6150001": "護理部"}, run_date=date(2026, 7, 3))
    output, planned_outputs = _w02_output(config, date(2026, 7, 3))
    canonical_path = runner._w02_order_plan_path()
    canonical_path.parent.mkdir(parents=True, exist_ok=True)
    canonical_path.write_text('{"historical": true}', encoding="utf-8")

    original_write_text = Path.write_text
    original_replace = Path.replace

    def deny_canonical_write(path: Path, *args, **kwargs):  # type: ignore[no-untyped-def]
        if path == canonical_path:
            raise PermissionError(13, "Permission denied", str(path))
        return original_write_text(path, *args, **kwargs)

    def deny_canonical_replace(path: Path, target: Path):
        if Path(target) == canonical_path:
            raise PermissionError(13, "Permission denied", str(target))
        return original_replace(path, target)

    monkeypatch.setattr(Path, "write_text", deny_canonical_write)
    monkeypatch.setattr(Path, "replace", deny_canonical_replace)
    monkeypatch.setattr("pos_report_bot.storage.run_state.sleep", lambda _seconds: None)

    result = runner._run_w02_order_workflow(output, planned_outputs, prior_failures=[])

    assert result.ok is True
    assert result.output_path != canonical_path
    assert result.output_path.parent == canonical_path.parent
    assert result.output_path.name.startswith("w02_order_plan_20260703_recovery_")
    assert result.output_path.is_file()
    assert canonical_path.read_text(encoding="utf-8") == '{"historical": true}'
    assert f"w02_plan:{result.output_path}" in result.actions


def test_w02_no_order_forms_advances_next_run_date(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    _write_r14_for_w02(tmp_path, run_date=date(2026, 7, 3), positive_order=False)
    runner = _w02_runner(config, tmp_path, departments={"6150001": "護理部"}, run_date=date(2026, 7, 3))
    output, planned_outputs = _w02_output(config, date(2026, 7, 3))

    result = runner._run_w02_order_workflow(output, planned_outputs, prior_failures=[])

    assert result.ok is True
    assert config.w02_order.next_run_date == "2026/07/17"


def test_manual_w02_success_advances_from_run_date_when_forced(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    _write_r14_for_w02(tmp_path, run_date=date(2026, 7, 10), positive_order=False)
    runner = _w02_runner(
        config,
        tmp_path,
        departments={"6150001": "護理部"},
        run_date=date(2026, 7, 10),
        run_source="gui_manual",
    )
    output, planned_outputs = _w02_output(config, date(2026, 7, 10), force=True)

    result = runner._run_w02_order_workflow(output, planned_outputs, prior_failures=[])

    assert result.ok is True
    assert config.w02_order.next_run_date == "2026/07/24"


def test_w02_run_completes_when_all_order_items_are_skipped_and_reported(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    _enable_only_w02(config)
    _write_r14_for_w02(tmp_path, run_date=date(2026, 7, 3), positive_order=True)
    FakeGmailSender.sent = []
    runner = _w02_runner(config, tmp_path, departments={}, run_date=date(2026, 7, 3))

    summary = runner.run()

    assert summary.ok is True
    assert summary.completed == 1
    assert summary.failures == ()
    assert config.w02_order.next_run_date == "2026/07/17"
    assert len(FakeGmailSender.sent) == 1
    state_payload = _load_run_state_payload(config, date(2026, 7, 3))
    assert state_payload["status"] == "success"
    assert next(iter(state_payload["outputs"].values()))["status"] == "completed"


def test_w02_run_fails_when_normal_forms_exist_but_pos_submission_is_disabled(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    _enable_only_w02(config)
    _write_r14_for_w02(
        tmp_path,
        run_date=date(2026, 7, 3),
        positive_order=True,
        extra_positive_order=True,
    )
    FakeGmailSender.sent = []
    runner = _w02_runner(
        config,
        tmp_path,
        departments={"6150002": "護理部"},
        known_item_codes={"6150001", "6150002"},
        run_date=date(2026, 7, 3),
    )

    summary = runner.run()

    assert summary.ok is False
    assert summary.completed == 0
    assert len(summary.failures) == 1
    assert summary.failures[0].error_code == "W02_POS_SUBMISSION_DISABLED"
    assert config.w02_order.next_run_date == "2026/07/03"
    assert len(FakeGmailSender.sent) == 0
    state_payload = _load_run_state_payload(config, date(2026, 7, 3))
    assert state_payload["status"] == "failed"
    assert next(iter(state_payload["outputs"].values()))["status"] == "failed"


def test_w02_pos_submission_success_advances_next_run_date(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    config.w02_order.pos_submission_enabled = True
    _write_r14_for_w02(
        tmp_path,
        run_date=date(2026, 7, 3),
        positive_order=True,
        extra_positive_order=True,
    )
    FakeGmailSender.sent = []
    FakeW02PosOrderAutomator.calls = []
    FakeW02PosOrderAutomator.result = SimpleNamespace(
        ok=True,
        message="fake W02 POS orders completed",
        actions=["fake_w02_pos_submit"],
        completed_forms=1,
        skipped_forms=0,
        error_code=None,
        diagnostic_path=None,
        skipped_issues=(),
        completed_form_counts_by_branch={"站前4樓": 2},
    )
    runner = _w02_runner(
        config,
        tmp_path,
        departments={"6150001": "護理部", "6150002": "護理部"},
        known_item_codes={"6150001", "6150002"},
        run_date=date(2026, 7, 3),
    )
    output, planned_outputs = _w02_output(config, date(2026, 7, 3))

    result = runner._run_w02_order_workflow(output, planned_outputs, prior_failures=[])

    assert result.ok is True
    assert result.error_code is None
    assert config.w02_order.next_run_date == "2026/07/17"
    assert FakeW02PosOrderAutomator.calls
    assert FakeW02PosOrderAutomator.calls[0]["form_count"] == 1
    assert "fake_w02_pos_submit" in result.actions


def test_w02_successful_runner_closes_the_internal_pos_session(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    _enable_only_w02(config)
    config.w02_order.pos_submission_enabled = True
    _write_r14_for_w02(
        tmp_path,
        run_date=date(2026, 7, 3),
        positive_order=True,
        extra_positive_order=True,
    )
    FakeW02PosOrderAutomator.calls = []
    FakeW02PosOrderAutomator.result = SimpleNamespace(
        ok=True,
        message="fake W02 POS orders completed",
        actions=["fake_w02_pos_submit"],
        completed_forms=1,
        skipped_forms=0,
        error_code=None,
        diagnostic_path=None,
        skipped_issues=(),
        completed_form_counts_by_branch={"站前4樓": 2},
    )
    pos_window = FakeReadyPosWindow()
    runner = _w02_runner(
        config,
        tmp_path,
        departments={"6150001": "護理部", "6150002": "護理部"},
        known_item_codes={"6150001", "6150002"},
        run_date=date(2026, 7, 3),
        pos_window=pos_window,
    )

    summary = runner.run()

    assert summary.ok is True
    assert pos_window.closed is True


def test_w02_diagnostic_mode_submits_three_items_for_three_branches_and_keeps_next_run_date(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    config.w02_order.pos_submission_enabled = True
    config.w02_order.diagnostic_mode = True
    r14_path = _write_r14_for_w02(
        tmp_path,
        run_date=date(2026, 7, 3),
        positive_order=True,
        extra_positive_order=True,
    )
    _append_w02_branch_orders(r14_path, "站前4樓", start_row=6, item_codes=("6150003", "6150004", "6150005"))
    _append_w02_branch_orders(r14_path, "站前11樓", start_row=4, item_codes=("6150011", "6150012", "6150013", "6150014"))
    _append_w02_branch_orders(r14_path, "忠孝7樓", start_row=4, item_codes=("6150021", "6150022", "6150023", "6150024"))
    FakeGmailSender.sent = []
    FakeW02PosOrderAutomator.calls = []
    FakeW02PosOrderAutomator.result = SimpleNamespace(
        ok=True,
        message="fake W02 POS diagnostic completed",
        actions=["fake_w02_pos_diagnostic_submit"],
        completed_forms=1,
        skipped_forms=0,
        error_code=None,
        diagnostic_path=tmp_path / "w02_diag.json",
        skipped_issues=(),
        completed_form_counts_by_branch={"站前4樓": 2},
    )
    runner = _w02_runner(
        config,
        tmp_path,
        departments={
            "6150001": "護理部",
            "6150002": "護理部",
            "6150003": "護理部",
            "6150004": "護理部",
            "6150005": "護理部",
            "6150011": "護理部",
            "6150012": "護理部",
            "6150013": "護理部",
            "6150014": "護理部",
            "6150021": "護理部",
            "6150022": "護理部",
            "6150023": "護理部",
            "6150024": "護理部",
        },
        known_item_codes={
            "6150001",
            "6150002",
            "6150003",
            "6150004",
            "6150005",
            "6150011",
            "6150012",
            "6150013",
            "6150014",
            "6150021",
            "6150022",
            "6150023",
            "6150024",
        },
        run_date=date(2026, 7, 3),
    )
    output, planned_outputs = _w02_output(config, date(2026, 7, 3))

    result = runner._run_w02_order_workflow(output, planned_outputs, prior_failures=[])

    assert result.ok is True
    assert result.error_code is None
    assert config.w02_order.next_run_date == "2026/07/03"
    assert len(FakeGmailSender.sent) == 0
    assert FakeW02PosOrderAutomator.calls[0]["form_count"] == 3
    assert FakeW02PosOrderAutomator.calls[0]["item_count"] == 9
    assert set(FakeW02PosOrderAutomator.calls[0]["item_codes"]) == {
        "6150001",
        "6150002",
        "6150003",
        "6150011",
        "6150012",
        "6150013",
        "6150021",
        "6150022",
        "6150023",
    }
    assert FakeW02PosOrderAutomator.calls[0]["kwargs"]["state_dir"] == Path(config.app.state_dir) / "20260703"
    assert "w02_diagnostic_mode:enabled" in result.actions
    assert "w02_diagnostic_item_limit_per_branch:3" in result.actions
    assert "w02_next_run_date_unchanged:2026/07/03" in result.actions
    plan_payload = json.loads(Path(result.output_path).read_text(encoding="utf-8"))
    assert plan_payload["diagnostic_mode"] is True
    assert plan_payload["order_item_count"] == 13
    assert plan_payload["diagnostic_execution_form_count"] == 3
    assert plan_payload["diagnostic_execution_item_count"] == 9


def test_w02_diagnostic_mode_fails_before_pos_when_three_by_three_sample_is_unavailable(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    config.w02_order.pos_submission_enabled = True
    config.w02_order.diagnostic_mode = True
    r14_path = _write_r14_for_w02(
        tmp_path,
        run_date=date(2026, 7, 3),
        positive_order=True,
        extra_positive_order=True,
    )
    _append_w02_branch_orders(r14_path, "站前11樓", start_row=4, item_codes=("6150011", "6150012", "6150013"))
    _append_w02_branch_orders(r14_path, "忠孝國際醫學3樓", start_row=4, item_codes=("6150031",))
    _append_w02_branch_orders(r14_path, "忠孝健康7樓", start_row=4, item_codes=("6150041", "6150042"))
    FakeGmailSender.sent = []
    FakeW02PosOrderAutomator.calls = []
    runner = _w02_runner(
        config,
        tmp_path,
        departments={
            "6150001": "護理部",
            "6150002": "護理部",
            "6150011": "護理部",
            "6150012": "護理部",
            "6150013": "護理部",
            "6150031": "護理部",
            "6150041": "護理部",
            "6150042": "護理部",
        },
        known_item_codes={
            "6150001",
            "6150002",
            "6150011",
            "6150012",
            "6150013",
            "6150031",
            "6150041",
            "6150042",
        },
        run_date=date(2026, 7, 3),
    )
    output, planned_outputs = _w02_output(config, date(2026, 7, 3))

    result = runner._run_w02_order_workflow(output, planned_outputs, prior_failures=[])

    assert result.ok is False
    assert result.error_code == "W02_DIAGNOSTIC_INSUFFICIENT_ITEMS"
    assert config.w02_order.next_run_date == "2026/07/03"
    assert FakeW02PosOrderAutomator.calls == []
    assert len(FakeGmailSender.sent) == 0
    assert "w02_diagnostic_available_branch_items:站前4樓=2|站前11樓=3|忠孝國際醫學3樓=1|忠孝健康7樓=2" in result.actions
    assert "w02_diagnostic_submitted_items:3" in result.actions
    assert "w02_next_run_date_unchanged:2026/07/03" in result.actions
    plan_payload = json.loads(Path(result.output_path).read_text(encoding="utf-8"))
    assert plan_payload["order_item_count"] == 8
    assert plan_payload["diagnostic_execution_item_count"] == 3


def test_w02_diagnostic_mode_reports_skipped_probe_item_as_failure_without_advancing(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    config.w02_order.pos_submission_enabled = True
    config.w02_order.diagnostic_mode = True
    r14_path = _write_r14_for_w02(
        tmp_path,
        run_date=date(2026, 7, 3),
        positive_order=True,
        extra_positive_order=True,
    )
    _append_w02_branch_orders(r14_path, "站前4樓", start_row=6, item_codes=("6150003", "6150004", "6150005"))
    _append_w02_branch_orders(r14_path, "站前11樓", start_row=4, item_codes=("6150011", "6150012", "6150013"))
    _append_w02_branch_orders(r14_path, "忠孝7樓", start_row=4, item_codes=("6150021", "6150022", "6150023"))
    FakeGmailSender.sent = []
    FakeW02PosOrderAutomator.calls = []
    FakeW02PosOrderAutomator.result = SimpleNamespace(
        ok=True,
        message="fake W02 POS diagnostic skipped item",
        actions=["probe:w02_failure_context:item_picker_no_selectable_result:6150001"],
        completed_forms=0,
        skipped_forms=0,
        error_code=None,
        diagnostic_path=tmp_path / "w02_diag.json",
        skipped_issues=(
            W02OrderIssue(
                branch="站前4樓",
                item_code="6150001",
                item_name="診斷模式找不到商品",
                quantity=20,
                reason="POS 建單時跳過：商品選擇視窗找不到料號。",
            ),
        ),
        completed_form_counts_by_branch={},
    )
    runner = _w02_runner(
        config,
        tmp_path,
        departments={
            "6150001": "護理部",
            "6150002": "護理部",
            "6150003": "護理部",
            "6150004": "護理部",
            "6150005": "護理部",
            "6150011": "護理部",
            "6150012": "護理部",
            "6150013": "護理部",
            "6150021": "護理部",
            "6150022": "護理部",
            "6150023": "護理部",
        },
        known_item_codes={
            "6150001",
            "6150002",
            "6150003",
            "6150004",
            "6150005",
            "6150011",
            "6150012",
            "6150013",
            "6150021",
            "6150022",
            "6150023",
        },
        run_date=date(2026, 7, 3),
    )
    output, planned_outputs = _w02_output(config, date(2026, 7, 3))

    result = runner._run_w02_order_workflow(output, planned_outputs, prior_failures=[])

    assert result.ok is False
    assert result.error_code == "W02_DIAGNOSTIC_ITEM_SKIPPED"
    assert config.w02_order.next_run_date == "2026/07/03"
    assert len(FakeGmailSender.sent) == 1
    body = str(FakeGmailSender.sent[0]["body"])
    assert "W02 下單異常品項" in body
    assert "6150001" in body
    assert "診斷模式找不到商品" in body
    assert "w02_diagnostic_item_limit_per_branch:3" in result.actions
    assert "w02_next_run_date_unchanged:2026/07/03" in result.actions
    assert any(action.startswith("probe:w02_failure_context:") for action in result.actions)


def test_w02_diagnostic_mode_sends_pos_skipped_issue_email_without_advancing(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    config.w02_order.pos_submission_enabled = True
    config.w02_order.diagnostic_mode = True
    r14_path = _write_r14_for_w02(
        tmp_path,
        run_date=date(2026, 7, 3),
        positive_order=True,
        extra_positive_order=True,
    )
    _append_w02_branch_orders(r14_path, "站前4樓", start_row=6, item_codes=("6150003", "6150004", "6150005"))
    _append_w02_branch_orders(r14_path, "站前11樓", start_row=4, item_codes=("6150011", "6150012", "6150013"))
    _append_w02_branch_orders(r14_path, "忠孝7樓", start_row=4, item_codes=("6150021", "6150022", "6150023"))
    FakeGmailSender.sent = []
    FakeW02PosOrderAutomator.calls = []
    FakeW02PosOrderAutomator.result = SimpleNamespace(
        ok=True,
        message="fake W02 POS diagnostic completed with skipped item",
        actions=["fake_w02_pos_diagnostic_submit", "w02_item_skipped:6150002:W02_POS_ITEM_NOT_FOUND"],
        completed_forms=1,
        skipped_forms=0,
        error_code=None,
        diagnostic_path=tmp_path / "w02_diag.json",
        skipped_issues=(
            W02OrderIssue(
                branch="站前4樓",
                item_code="6150002",
                item_name="POS 找不到但 R14 有下單數",
                quantity=20,
                reason="POS 建單時跳過：商品選擇視窗找不到料號。",
            ),
        ),
        completed_form_counts_by_branch={"站前4樓": 2},
    )
    runner = _w02_runner(
        config,
        tmp_path,
        departments={
            "6150001": "護理部",
            "6150002": "護理部",
            "6150003": "護理部",
            "6150004": "護理部",
            "6150005": "護理部",
            "6150011": "護理部",
            "6150012": "護理部",
            "6150013": "護理部",
            "6150021": "護理部",
            "6150022": "護理部",
            "6150023": "護理部",
        },
        known_item_codes={
            "6150001",
            "6150002",
            "6150003",
            "6150004",
            "6150005",
            "6150011",
            "6150012",
            "6150013",
            "6150021",
            "6150022",
            "6150023",
        },
        run_date=date(2026, 7, 3),
    )
    output, planned_outputs = _w02_output(config, date(2026, 7, 3))

    result = runner._run_w02_order_workflow(output, planned_outputs, prior_failures=[])

    assert result.ok is True
    assert result.error_code is None
    assert config.w02_order.next_run_date == "2026/07/03"
    assert len(FakeGmailSender.sent) == 1
    body = str(FakeGmailSender.sent[0]["body"])
    assert "W02 成功建單統計" in body
    assert "成功訂單筆數" in body
    assert "成功訂單數" not in body
    assert "<td style=\"border:1px solid #444;padding:4px 8px;text-align:left;\">2</td>" in body
    assert "W02 下單異常品項" in body
    assert "6150002" in body
    assert "POS 找不到但 R14 有下單數" in body
    assert "w02_pos_skipped_items:1" in result.actions
    assert "w02_next_run_date_unchanged:2026/07/03" in result.actions
    assert "已跳過並寄送通知" in result.message


def test_w02_diagnostic_mode_no_order_forms_does_not_advance_next_run_date(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    config.w02_order.diagnostic_mode = True
    _write_r14_for_w02(tmp_path, run_date=date(2026, 7, 3), positive_order=False)
    FakeGmailSender.sent = []
    FakeW02PosOrderAutomator.calls = []
    runner = _w02_runner(config, tmp_path, departments={"6150001": "護理部"}, run_date=date(2026, 7, 3))
    output, planned_outputs = _w02_output(config, date(2026, 7, 3))

    result = runner._run_w02_order_workflow(output, planned_outputs, prior_failures=[])

    assert result.ok is True
    assert config.w02_order.next_run_date == "2026/07/03"
    assert FakeW02PosOrderAutomator.calls == []
    assert len(FakeGmailSender.sent) == 0
    assert "w02_diagnostic_mode:enabled" in result.actions
    assert "w02_next_run_date_unchanged:2026/07/03" in result.actions


def test_w02_diagnostic_mode_no_order_forms_with_issues_sends_email_without_advancing(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    config.w02_order.diagnostic_mode = True
    _write_r14_for_w02(tmp_path, run_date=date(2026, 7, 3), positive_order=True)
    FakeGmailSender.sent = []
    FakeW02PosOrderAutomator.calls = []
    runner = _w02_runner(config, tmp_path, departments={}, run_date=date(2026, 7, 3))
    output, planned_outputs = _w02_output(config, date(2026, 7, 3))

    result = runner._run_w02_order_workflow(output, planned_outputs, prior_failures=[])

    assert result.ok is True
    assert config.w02_order.next_run_date == "2026/07/03"
    assert FakeW02PosOrderAutomator.calls == []
    assert len(FakeGmailSender.sent) == 1
    body = str(FakeGmailSender.sent[0]["body"])
    assert "W02 下單異常品項" in body
    assert "6150001" in body
    assert "B 欄找不到" in body
    assert "w02_diagnostic_mode:enabled" in result.actions
    assert "w02_skipped_items:1" in result.actions
    assert "w02_next_run_date_unchanged:2026/07/03" in result.actions


def test_w02_pos_submission_failure_keeps_next_run_date(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    config.w02_order.pos_submission_enabled = True
    _write_r14_for_w02(
        tmp_path,
        run_date=date(2026, 7, 3),
        positive_order=True,
        extra_positive_order=True,
    )
    FakeGmailSender.sent = []
    FakeW02PosOrderAutomator.calls = []
    FakeW02PosOrderAutomator.result = SimpleNamespace(
        ok=False,
        message="fake W02 POS order failed",
        actions=["fake_w02_pos_failed"],
        completed_forms=0,
        skipped_forms=0,
        error_code="W02_POS_ORDER_FAILED",
        diagnostic_path=tmp_path / "w02_diag.json",
        skipped_issues=(
            W02OrderIssue(
                branch="站前4樓",
                item_code="6150002",
                item_name="已跳過但流程後續失敗",
                quantity=20,
                reason="POS 建單時跳過：商品選擇視窗找不到料號。",
            ),
        ),
        completed_form_counts_by_branch={},
    )
    runner = _w02_runner(
        config,
        tmp_path,
        departments={"6150001": "護理部", "6150002": "護理部"},
        known_item_codes={"6150001", "6150002"},
        run_date=date(2026, 7, 3),
    )
    output, planned_outputs = _w02_output(config, date(2026, 7, 3))

    result = runner._run_w02_order_workflow(output, planned_outputs, prior_failures=[])

    assert result.ok is False
    assert result.error_code == "W02_POS_ORDER_FAILED"
    assert config.w02_order.next_run_date == "2026/07/03"
    assert FakeW02PosOrderAutomator.calls
    assert len(FakeGmailSender.sent) == 0
    assert "fake_w02_pos_failed" in result.actions


def test_w02_pos_submission_sends_combined_issues_after_pos_flow(tmp_path: Path) -> None:
    config = _w02_config(tmp_path, next_run_date="2026/07/03")
    config.w02_order.pos_submission_enabled = True
    _write_r14_for_w02(
        tmp_path,
        run_date=date(2026, 7, 3),
        positive_order=True,
        extra_positive_order=True,
    )
    FakeGmailSender.sent = []
    FakeW02PosOrderAutomator.calls = []
    FakeW02PosOrderAutomator.result = SimpleNamespace(
        ok=True,
        message="fake W02 POS orders completed",
        actions=["fake_w02_pos_submit", "w02_item_skipped:6150002:W02_POS_ITEM_NOT_FOUND"],
        completed_forms=1,
        skipped_forms=0,
        error_code=None,
        diagnostic_path=None,
        skipped_issues=(
            W02OrderIssue(
                branch="站前4樓",
                item_code="6150002",
                item_name="POS找不到商品",
                quantity=20,
                reason="POS 建單時跳過：商品選擇視窗找不到料號。",
            ),
        ),
        completed_form_counts_by_branch={"站前4樓": 2},
    )
    runner = _w02_runner(
        config,
        tmp_path,
        departments={"6150002": "護理部"},
        known_item_codes={"6150001", "6150002"},
        run_date=date(2026, 7, 3),
    )
    output, planned_outputs = _w02_output(config, date(2026, 7, 3))

    result = runner._run_w02_order_workflow(output, planned_outputs, prior_failures=[])

    assert result.ok is True
    assert config.w02_order.next_run_date == "2026/07/17"
    assert len(FakeGmailSender.sent) == 1
    body = str(FakeGmailSender.sent[0]["body"])
    assert "W02 成功建單統計" in body
    assert "成功訂單筆數" in body
    assert "成功訂單數" not in body
    assert "<td style=\"border:1px solid #444;padding:4px 8px;text-align:left;\">2</td>" in body
    assert "站前4樓" in body
    assert "6150001" in body
    assert "6150002" in body


def _w02_config(tmp_path: Path, *, next_run_date: str):  # type: ignore[no-untyped-def]
    config = load_project_config(ROOT / "config_templates" / "app.template.yaml")
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.state_dir = str(tmp_path / "state")
    config.w02_order.next_run_date = next_run_date
    config.email.notify_on_failure = False
    return config


def _enable_only_w02(config) -> None:  # type: ignore[no-untyped-def]
    for report in config.reports:
        report.enabled = report.id == "W02"


def _load_run_state_payload(config, run_date: date) -> dict[str, object]:  # type: ignore[no-untyped-def]
    path = Path(config.app.state_dir) / run_date.strftime("%Y%m%d") / "run_state_latest.json"
    return json.loads(path.read_text(encoding="utf-8"))


def _w02_runner(
    config,
    tmp_path: Path,
    *,
    departments: dict[str, str],
    known_item_codes: set[str] | None = None,
    run_date: date,
    run_source: str = "windows_task_scheduler",
    pos_window: FakeReadyPosWindow | None = None,
):  # type: ignore[no-untyped-def]
    connected_window = pos_window or FakeReadyPosWindow()
    return AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        r14_inventory_client_factory=lambda _config: FakeW02DepartmentClient(departments, known_item_codes),
        gmail_sender_factory=FakeGmailSender,
        connect_pos_window_func=lambda **_kwargs: connected_window,
        w02_pos_order_automator_factory=FakeW02PosOrderAutomator,
        run_date=run_date,
        run_source=run_source,
    )


def _w02_output(config, run_date: date, *, force: bool = False):  # type: ignore[no-untyped-def]
    plan = build_dry_run_plan(
        config,
        today=run_date,
        force_weekly_report_ids={"W02"} if force else None,
    )
    return next(output for output in plan.outputs if output.task_id == "W02"), plan.outputs


def _write_r14_for_w02(
    tmp_path: Path,
    *,
    run_date: date,
    positive_order: bool,
    extra_positive_order: bool = False,
) -> Path:
    folder = tmp_path / "downloads" / "R14" / run_date.strftime("%Y%m%d")
    folder.mkdir(parents=True, exist_ok=True)
    report_date = run_date - timedelta(days=1)
    path = folder / f"診所stock status - 2026 demand planning-{report_date.strftime('%m%d')}.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["F2"] = datetime(report_date.year, report_date.month, report_date.day)

    branch = workbook.create_sheet("站前4樓")
    branch.cell(2, 7).value = run_date.strftime("%Y/%m")
    branch.cell(3, 2).value = "凱惠料號"
    branch.cell(3, 3).value = "品名"
    branch.cell(3, 4).value = "盒入數"
    branch.cell(3, 6).value = "站前4樓庫存"
    branch.cell(3, 7).value = "Actual"
    branch.cell(2, 8).value = "下單數"
    branch.cell(2, 9).value = "2026/06"
    branch.cell(3, 9).value = "Actual"
    branch.cell(4, 2).value = "6150001"
    branch.cell(4, 3).value = "測試商品"
    branch.cell(4, 4).value = 20
    branch.cell(4, 6).value = 0 if positive_order else 999
    branch.cell(4, 7).value = 28
    branch.cell(4, 8).value = 20 if positive_order else 0
    branch.cell(4, 9).value = 28
    if extra_positive_order:
        branch.cell(5, 2).value = "6150002"
        branch.cell(5, 3).value = "正常下單商品"
        branch.cell(5, 4).value = 10
        branch.cell(5, 6).value = 0
        branch.cell(5, 7).value = 28
        branch.cell(5, 8).value = 20
        branch.cell(5, 9).value = 28
    for branch_name in (
        "站前11樓",
        "忠孝國際醫學3樓",
        "忠孝7樓",
        "忠孝健康7樓",
        "忠孝預防醫學3樓",
    ):
        empty_branch = workbook.create_sheet(branch_name)
        empty_branch.cell(2, 7).value = report_date.strftime("%Y/%m")
        empty_branch.cell(3, 2).value = "凱惠料號"
        empty_branch.cell(3, 3).value = "品名"
        empty_branch.cell(3, 4).value = "盒入數"
        empty_branch.cell(3, 6).value = f"{branch_name}庫存"
        empty_branch.cell(3, 7).value = "Actual"
        empty_branch.cell(2, 8).value = "下單數"
        empty_branch.cell(2, 9).value = "2026/06"
        empty_branch.cell(3, 9).value = "Actual"
    workbook.create_sheet("領用表")
    workbook.save(path)
    return path


def _append_w02_branch_orders(
    path: Path,
    branch_name: str,
    *,
    start_row: int,
    item_codes: tuple[str, ...],
) -> None:
    workbook = load_workbook(path)
    if branch_name in workbook.sheetnames:
        branch = workbook[branch_name]
    else:
        branch = workbook.create_sheet(branch_name)
        branch.cell(2, 7).value = "2026/07"
        branch.cell(3, 2).value = "凱惠料號"
        branch.cell(3, 3).value = "品名"
        branch.cell(3, 4).value = "盒入數"
        branch.cell(3, 6).value = f"{branch_name}庫存"
        branch.cell(3, 7).value = "Actual"
        branch.cell(2, 8).value = "下單數"
    for offset, item_code in enumerate(item_codes):
        row = start_row + offset
        branch.cell(row, 2).value = item_code
        branch.cell(row, 3).value = f"{branch_name}診斷商品{offset + 1}"
        branch.cell(row, 4).value = 1
        branch.cell(row, 6).value = 0
        branch.cell(row, 7).value = 1
        branch.cell(row, 8).value = 1
    workbook.save(path)
