from contextlib import contextmanager
from datetime import UTC, date, datetime, timedelta
import json
import os
from pathlib import Path
import stat
import sys
from types import SimpleNamespace
import zipfile

from openpyxl import Workbook, load_workbook
import pytest

from pos_report_bot.app.automation_runner import (
    AutomationRunner,
    AutomationRunSummary,
    ReportRunFailure,
    forced_weekly_report_ids_for_run_source,
)
from pos_report_bot.config.loader import load_project_config
from pos_report_bot.drive.uploader import (
    DriveTemplateDownloadResult,
    DriveUploadResult,
    MockDriveUploader,
)
from pos_report_bot.google.oauth import (
    GOOGLE_DRIVE_PROFILE,
    GOOGLE_DRIVE_SCOPES,
    GOOGLE_GMAIL_PROFILE,
    GOOGLE_GMAIL_SCOPES,
    GOOGLE_SHEETS_PROFILE,
    GOOGLE_SHEETS_SCOPES,
)
from pos_report_bot.google.sheets import R14InventorySheetError
from pos_report_bot.pos.report_automation import AUTOMATION_LOGIC_FINGERPRINT, ReportAutomationError
from pos_report_bot.pos.save_as_handler import MockSaveAsHandler
from pos_report_bot.pos.ui_probe import UiProbeError
from pos_report_bot.reports.models import PlannedOutput
from pos_report_bot.reports.planner import build_dry_run_plan
from pos_report_bot.reports.r14_transformer import R14TransformError
from pos_report_bot.storage.r13_no_data import write_r13_no_data_marker
from pos_report_bot.storage.run_state import RunStateStore
import pos_report_bot.app.automation_runner as automation_runner_module
from tests.unit.test_report_automation import FakePosControl, FakeRectPosControl


ROOT = Path(__file__).resolve().parents[2]


class FakeClosablePosWindow(FakePosControl):
    def __init__(self) -> None:
        super().__init__("SPA-POS", children=[FakePosControl("統計報表", "MenuItem")])
        self.closed = False

    def close(self) -> None:
        self.closed = True


class FakeProcessBoundPosWindow(FakeClosablePosWindow):
    def __init__(self, process_id: int) -> None:
        super().__init__()
        self._process_id = process_id

    def process_id(self) -> int:
        return self._process_id


def _ready_pos_window(*roots: str) -> FakePosControl:
    menu_roots = roots or ("統計報表",)
    return FakePosControl(
        "SPA-POS",
        "Window",
        children=[FakePosControl(root, "MenuItem") for root in menu_roots],
    )


def _win32_menu_shell_without_root_menus() -> FakePosControl:
    window = FakePosControl(
        "SPA-POS Ver.1.5.18.77",
        "Window",
        children=[
            FakePosControl("menuStrip1", "Window"),
            FakePosControl("登入檢查完成!\r\r請從上方選單選取您要執行的功能.", "Static"),
        ],
    )
    window._pos_report_bot_backend = "win32"
    return window


def _pos_startup_ini_window() -> tuple[FakePosControl, FakePosControl, FakePosControl]:
    combo = FakePosControl("", "ComboBox", automation_id="M_INI")
    confirm = FakePosControl("確　定", "Button", automation_id="B_INIOK")
    window = FakePosControl(
        "SPA-POS",
        "Window",
        children=[
            FakePosControl(".", "Button", automation_id="B_CD", enabled=False),
            FakePosControl("關閉", "Button", automation_id="button1"),
            confirm,
            combo,
        ],
    )
    window._pos_report_bot_backend = "win32"
    return window, combo, confirm


class FakeKeyboardSelectableStartupIniCombo(FakePosControl):
    def __init__(self) -> None:
        super().__init__("", "ComboBox", automation_id="M_INI")
        self.items = [
            r"c:\tkhspa\tkhspa -測試.ini",
            r"c:\tkhspa\tkhspa-正式.ini",
        ]
        self.focused = False
        self.dropdown_opened = False

    def select(self, _value: str) -> None:
        return

    def ItemTexts_(self) -> list[str]:
        return self.items

    def click_input(self) -> None:
        self.focused = True

    def type_keys(self, value: str, with_spaces: bool = False) -> None:
        if value == "{F4}":
            self.dropdown_opened = True
            return
        if not self.dropdown_opened:
            return
        if "{HOME}" not in value:
            return
        index = value.count("{DOWN}")
        if 0 <= index < len(self.items):
            self.selected_value = self.items[index]


class FakeUnverifiableStartupIniCombo(FakePosControl):
    def __init__(self) -> None:
        super().__init__("", "ComboBox", automation_id="M_INI")

    def select(self, _value: str) -> None:
        return


class FakeTextsOnlyStartupIniCombo(FakeUnverifiableStartupIniCombo):
    def texts(self) -> list[str]:
        return [
            r"c:\tkhspa\tkhspa -測試.ini",
            r"c:\tkhspa\tkhspa-正式.ini",
        ]


class FakeAlreadySelectedStartupIniCombo(FakePosControl):
    def __init__(self, selected_value: str) -> None:
        super().__init__("", "ComboBox", automation_id="M_INI")
        self.selected_value = selected_value

    def select(self, _value: str) -> None:
        raise RuntimeError("select must not be called for an already selected profile")


def _pos_startup_ini_window_with_combo(combo: FakePosControl) -> tuple[FakePosControl, FakePosControl]:
    confirm = FakePosControl("確　定", "Button", automation_id="B_INIOK")
    window = FakePosControl("SPA-POS", "Window", children=[confirm, combo])
    window._pos_report_bot_backend = "win32"
    return window, confirm


def _load_runner_config(tmp_path: Path):  # type: ignore[no-untyped-def]
    config = load_project_config(ROOT / "config_templates" / "app.template.yaml")
    config.app.state_dir = str(tmp_path / "state")
    # Tests must opt in to external notification explicitly.  Production
    # defaults remain enabled in config_templates/app.template.yaml.
    config.email.enabled = False
    return config


def _disable_uploads(config) -> None:  # type: ignore[no-untyped-def]
    for report in config.reports:
        report.upload_enabled = False


def _set_runtime_dirs(config, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.app.state_dir = str(tmp_path / "state")


class FakeR14TemplateDownloader:
    def __init__(
        self,
        source_path: Path | None,
        *,
        error_code: str = "R14_TEMPLATE_DRIVE_FILE_MISSING",
    ) -> None:
        self.source_path = source_path
        self.error_code = error_code
        self.calls: list[dict[str, object]] = []

    def download_latest_template(  # type: ignore[no-untyped-def]
        self,
        folder_id_or_url: str,
        *,
        filename_glob: str,
        destination_dir: Path,
    ) -> DriveTemplateDownloadResult:
        self.calls.append(
            {
                "folder_id_or_url": folder_id_or_url,
                "filename_glob": filename_glob,
                "destination_dir": destination_dir,
            }
        )
        if self.source_path is None:
            return DriveTemplateDownloadResult(
                success=False,
                folder_id="folder123",
                error_code=self.error_code,
                message="R14 雲端模板資料夾沒有符合檔名規則的 .xlsx 檔案。",
            )
        destination_dir.mkdir(parents=True, exist_ok=True)
        destination = destination_dir / self.source_path.name
        destination.write_bytes(self.source_path.read_bytes())
        return DriveTemplateDownloadResult(
            success=True,
            drive_file_id="drive-template-1",
            folder_id="folder123",
            source_name=self.source_path.name,
            modified_time="2026-09-03T02:27:44Z",
            local_path=destination,
            size=destination.stat().st_size,
            sha256="fake-sha256",
            message="Mock R14 template download completed.",
        )


def test_forced_weekly_report_ids_are_manual_only() -> None:
    assert forced_weekly_report_ids_for_run_source("gui_manual") == {"W01", "W02"}
    assert forced_weekly_report_ids_for_run_source("manual_cli") == set()
    assert forced_weekly_report_ids_for_run_source("manual_single_task") == {"W01", "W02"}
    assert forced_weekly_report_ids_for_run_source("windows_task_scheduler") == set()
    assert forced_weekly_report_ids_for_run_source("dry_run") == set()


def test_automation_runner_refuses_to_touch_pos_when_another_process_holds_run_lock(
    monkeypatch,
    tmp_path: Path,
) -> None:  # type: ignore[no-untyped-def]
    config = _load_runner_config(tmp_path)
    for report in config.reports:
        report.enabled = report.id == "R01"
    connect_calls = 0

    def fail_connect(**_kwargs):  # type: ignore[no-untyped-def]
        nonlocal connect_calls
        connect_calls += 1
        raise AssertionError("runner must not touch POS while another process owns the lock")

    @contextmanager
    def blocked_run_lock(*_args, **_kwargs):  # type: ignore[no-untyped-def]
        raise automation_runner_module.AutomationRunAlreadyActiveError(
            owner={"pid": 1234, "run_source": "historical_backfill"}
        )
        yield

    monkeypatch.setattr(automation_runner_module, "automation_run_lock", blocked_run_lock)
    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=fail_connect,
        run_source="windows_task_scheduler",
        run_date=date(2026, 8, 11),
    )

    summary = runner.run()

    assert summary.ok is False
    assert summary.error_code == "AUTOMATION_ALREADY_RUNNING"
    assert summary.completed == 0
    assert summary.total == 0
    assert "historical_backfill" in (summary.details or "")
    assert connect_calls == 0
    assert summary.evidence_bundle_path is not None
    assert Path(summary.evidence_bundle_path).exists()


def test_gui_manual_force_is_limited_to_selected_weekly_tasks(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)

    r13_runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        run_source="gui_manual",
        selected_task_ids={"R13"},
    )
    w02_runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        run_source="gui_manual",
        selected_task_ids={"W02"},
    )

    assert r13_runner._forced_weekly_report_ids() == set()
    assert w02_runner._forced_weekly_report_ids() == {"W02"}


@pytest.mark.parametrize("run_source", ["windows_task_scheduler", "manual_cli"])
def test_automation_runner_batch_does_not_start_w02_before_next_run_date(
    tmp_path: Path,
    run_source: str,
) -> None:
    config = _load_runner_config(tmp_path)
    config.w02_order.next_run_date = "2026/07/17"
    for report in config.reports:
        report.enabled = report.id == "W02"
    connect_calls = 0

    def fail_connect(**_kwargs):  # type: ignore[no-untyped-def]
        nonlocal connect_calls
        connect_calls += 1
        raise AssertionError("W02 must not connect to POS before next_run_date")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=fail_connect,
        run_date=date(2026, 7, 15),
        run_source=run_source,
    )

    summary = runner.run()

    assert summary.ok is True
    assert summary.total == 1
    assert summary.skipped == 1
    assert connect_calls == 0


def test_automation_runner_batch_rejects_invalid_w02_next_run_date(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.w02_order.next_run_date = "not-a-date"
    for report in config.reports:
        report.enabled = report.id == "W02"

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        run_date=date(2026, 7, 15),
        run_source="manual_cli",
    )

    summary = runner.run()

    assert summary.ok is False
    assert summary.error_code == "W02_NEXT_RUN_DATE_INVALID"


def test_automation_runner_local_transform_log_records_run_source(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    for report in config.reports:
        report.enabled = report.id == "W02"
    plan = build_dry_run_plan(
        config,
        today=date(2026, 7, 15),
        force_weekly_report_ids={"W02"},
    )
    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        run_date=date(2026, 7, 15),
        run_source="gui_manual",
    )
    log_path = tmp_path / "automation_local_transform.jsonl"

    runner._write_local_transform_log_event(
        log_path,
        "local_transform_start",
        output=plan.outputs[0],
        planned_task_ids=["W02"],
    )

    record = json.loads(log_path.read_text(encoding="utf-8").splitlines()[0])
    assert record["run_source"] == "gui_manual"


def test_automation_runner_selects_pos_startup_ini_before_login(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    window, combo, confirm = _pos_startup_ini_window()

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
    )

    handled = runner._prepare_pos_window_for_login(config, window)

    assert handled is window
    assert combo.selected_value == r"c:\tkhspa\tkhspa -測試.ini"
    assert confirm.clicked is True


def test_existing_pos_session_with_different_attested_ini_is_blocked(monkeypatch, tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    _disable_uploads(config)
    config.email.enabled = False
    config.login.required = False
    config.pos.startup_ini_selection_enabled = True
    config.pos.startup_ini_profile = r"c:\tkhspa\tkhspa-正式.ini"
    for report in config.reports:
        report.enabled = report.id == "R01"
    marker_path = Path(config.app.state_dir) / "pos_session_profile.json"
    marker_path.parent.mkdir(parents=True, exist_ok=True)
    marker_path.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "profile": r"c:\tkhspa\tkhspa -測試.ini",
                "process_id": 4242,
                "process_created_at": "2026-09-02T08:00:00.000000+00:00",
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    calls: list[str] = []
    monkeypatch.setattr(
        automation_runner_module,
        "_safe_process_creation_time",
        lambda _process_id: "2026-09-02T08:00:00.000000+00:00",
    )

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: FakeProcessBoundPosWindow(4242),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run()

    assert summary.ok is False
    assert calls == []
    assert summary.failures[0].error_code == "POS_ACTIVE_SESSION_PROFILE_MISMATCH"
    assert "POS_ACTIVE_SESSION_PROFILE_MISMATCH" in summary.failures[0].message


def test_existing_process_bound_pos_without_ini_attestation_is_blocked(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    config.pos.startup_ini_selection_enabled = True
    config.pos.startup_ini_profile = r"c:\tkhspa\tkhspa-正式.ini"
    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
    )

    with pytest.raises(RuntimeError, match="POS_ACTIVE_SESSION_PROFILE_UNVERIFIED"):
        runner._validate_or_record_pos_session_profile(config, FakeProcessBoundPosWindow(4242))


def test_pos_session_ini_selection_is_recorded_with_process_identity(monkeypatch, tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    config.pos.startup_ini_selection_enabled = True
    config.pos.startup_ini_profile = r"c:\tkhspa\tkhspa-正式.ini"
    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        run_source="gui_manual",
    )
    monkeypatch.setattr(
        automation_runner_module,
        "_safe_process_creation_time",
        lambda _process_id: "2026-09-02T08:00:00.000000+00:00",
    )
    runner._pos_startup_ini_profile_verified_this_run = config.pos.startup_ini_profile
    runner._pos_startup_ini_process_id_verified_this_run = 4242
    runner._pos_startup_ini_process_created_at_verified_this_run = "2026-09-02T08:00:00.000000+00:00"

    runner._validate_or_record_pos_session_profile(config, FakeProcessBoundPosWindow(4242))

    marker_path = Path(config.app.state_dir) / "pos_session_profile.json"
    payload = json.loads(marker_path.read_text(encoding="utf-8"))
    assert payload["profile"] == r"c:\tkhspa\tkhspa-正式.ini"
    assert payload["process_id"] == 4242
    assert payload["process_created_at"] == "2026-09-02T08:00:00.000000+00:00"
    assert payload["run_source"] == "gui_manual"


def test_pos_session_ini_selection_cannot_be_rebound_to_a_different_process(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    config.pos.startup_ini_selection_enabled = True
    config.pos.startup_ini_profile = r"c:\tkhspa\tkhspa-正式.ini"
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")
    runner._pos_startup_ini_profile_verified_this_run = config.pos.startup_ini_profile
    runner._pos_startup_ini_process_id_verified_this_run = 1111

    with pytest.raises(RuntimeError, match="POS_ACTIVE_SESSION_PROFILE_UNVERIFIED"):
        runner._validate_or_record_pos_session_profile(config, FakeProcessBoundPosWindow(2222))

    assert not (Path(config.app.state_dir) / "pos_session_profile.json").exists()


def test_pos_session_ini_selection_accepts_fresh_runner_launched_process_handoff(
    monkeypatch,
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    config.pos.startup_ini_selection_enabled = True
    config.pos.startup_ini_profile = r"c:\tkhspa\tkhspa-正式.ini"
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")
    runner._pos_launch_evidence = {
        "launch_requested": True,
        "requested_at": "2026-09-02T08:00:00.000000+00:00",
    }
    runner._pos_startup_ini_profile_verified_this_run = config.pos.startup_ini_profile
    runner._pos_startup_ini_process_id_verified_this_run = 1111
    runner._pos_startup_ini_process_created_at_verified_this_run = "2026-09-02T08:00:01.000000+00:00"
    creation_times = {
        1111: "2026-09-02T08:00:01.000000+00:00",
        2222: "2026-09-02T08:00:02.000000+00:00",
        3333: "2026-09-02T08:00:03.000000+00:00",
    }
    monkeypatch.setattr(
        automation_runner_module,
        "_safe_process_creation_time",
        lambda process_id: creation_times.get(process_id),
    )

    runner._validate_or_record_pos_session_profile(config, FakeProcessBoundPosWindow(2222))

    marker_path = Path(config.app.state_dir) / "pos_session_profile.json"
    payload = json.loads(marker_path.read_text(encoding="utf-8"))
    assert payload["profile"] == config.pos.startup_ini_profile
    assert payload["process_id"] == 2222
    assert payload["process_created_at"] == "2026-09-02T08:00:02.000000+00:00"
    assert runner._active_pos_session_profile_evidence == "selected_this_run_launched_process_handoff"
    assert runner._pos_launch_evidence["profile_handoff_bound_process_id"] == 2222

    with pytest.raises(RuntimeError, match="POS_ACTIVE_SESSION_PROFILE_UNVERIFIED"):
        runner._validate_or_record_pos_session_profile(config, FakeProcessBoundPosWindow(3333))


def test_pos_session_ini_selection_accepts_same_fresh_process_when_selection_creation_time_was_unreadable(
    monkeypatch,
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    config.pos.startup_ini_selection_enabled = True
    config.pos.startup_ini_profile = r"c:\tkhspa\tkhspa-正式.ini"
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")
    runner._pos_launch_evidence = {
        "launch_requested": True,
        "requested_at": "2026-09-02T13:43:41.519161+00:00",
    }
    runner._pos_startup_ini_profile_verified_this_run = config.pos.startup_ini_profile
    runner._pos_startup_ini_process_id_verified_this_run = 8496
    runner._pos_startup_ini_process_created_at_verified_this_run = None
    creation_time_reads: list[int | None] = []

    def read_creation_time_once(process_id: int | None) -> str | None:
        creation_time_reads.append(process_id)
        if process_id == 8496 and len(creation_time_reads) == 1:
            return "2026-09-02T13:43:42.000000+00:00"
        return None

    monkeypatch.setattr(
        automation_runner_module,
        "_safe_process_creation_time",
        read_creation_time_once,
    )

    runner._validate_or_record_pos_session_profile(config, FakeProcessBoundPosWindow(8496))

    marker_path = Path(config.app.state_dir) / "pos_session_profile.json"
    payload = json.loads(marker_path.read_text(encoding="utf-8"))
    assert payload["process_id"] == 8496
    assert payload["process_created_at"] == "2026-09-02T13:43:42.000000+00:00"
    assert payload["selection_process_created_at"] is None
    assert runner._active_pos_session_profile_evidence == "selected_this_run_fresh_process_rebound"
    assert creation_time_reads == [8496]


def test_pos_session_ini_selection_rejects_same_preexisting_process_when_selection_creation_time_was_unreadable(
    monkeypatch,
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    config.pos.startup_ini_selection_enabled = True
    config.pos.startup_ini_profile = r"c:\tkhspa\tkhspa-正式.ini"
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")
    runner._pos_launch_evidence = {
        "launch_requested": True,
        "requested_at": "2026-09-02T13:43:41.519161+00:00",
    }
    runner._pos_startup_ini_profile_verified_this_run = config.pos.startup_ini_profile
    runner._pos_startup_ini_process_id_verified_this_run = 8496
    runner._pos_startup_ini_process_created_at_verified_this_run = None
    monkeypatch.setattr(
        automation_runner_module,
        "_safe_process_creation_time",
        lambda process_id: (
            "2026-09-02T13:40:00.000000+00:00" if process_id == 8496 else None
        ),
    )

    with pytest.raises(RuntimeError, match="POS_ACTIVE_SESSION_PROFILE_UNVERIFIED"):
        runner._validate_or_record_pos_session_profile(config, FakeProcessBoundPosWindow(8496))

    assert not (Path(config.app.state_dir) / "pos_session_profile.json").exists()


def test_pos_session_ini_selection_uses_run_scoped_same_pid_proof_when_creation_time_is_unavailable(
    monkeypatch,
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    config.pos.startup_ini_selection_enabled = True
    config.pos.startup_ini_profile = r"c:\tkhspa\tkhspa-正式.ini"
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")
    runner._pos_launch_evidence = {
        "launch_requested": True,
        "requested_at": datetime.now(tz=UTC).isoformat(),
    }
    runner._pos_startup_ini_profile_verified_this_run = config.pos.startup_ini_profile
    runner._pos_startup_ini_process_id_verified_this_run = 4432
    runner._pos_startup_ini_process_created_at_verified_this_run = None
    monkeypatch.setattr(
        automation_runner_module,
        "_safe_process_creation_time",
        lambda _process_id: None,
    )

    runner._validate_or_record_pos_session_profile(config, FakeProcessBoundPosWindow(4432))
    runner._validate_or_record_pos_session_profile(config, FakeProcessBoundPosWindow(4432))

    assert runner._active_pos_session_profile == config.pos.startup_ini_profile
    assert runner._active_pos_session_profile_evidence == "selected_this_run_same_pid_run_scoped"
    assert runner._pos_launch_evidence["profile_run_scoped_bound_process_id"] == 4432
    assert not (Path(config.app.state_dir) / "pos_session_profile.json").exists()


def test_pos_session_ini_selection_rejects_stale_run_scoped_launch_proof(
    monkeypatch,
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    config.pos.startup_ini_selection_enabled = True
    config.pos.startup_ini_profile = r"c:\tkhspa\tkhspa-正式.ini"
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")
    runner._pos_launch_evidence = {
        "launch_requested": True,
        "requested_at": (datetime.now(tz=UTC) - timedelta(minutes=11)).isoformat(),
    }
    runner._pos_startup_ini_profile_verified_this_run = config.pos.startup_ini_profile
    runner._pos_startup_ini_process_id_verified_this_run = 4432
    runner._pos_startup_ini_process_created_at_verified_this_run = None
    monkeypatch.setattr(
        automation_runner_module,
        "_safe_process_creation_time",
        lambda _process_id: None,
    )

    with pytest.raises(RuntimeError, match="POS_ACTIVE_SESSION_PROFILE_UNVERIFIED"):
        runner._validate_or_record_pos_session_profile(config, FakeProcessBoundPosWindow(4432))

    assert not (Path(config.app.state_dir) / "pos_session_profile.json").exists()


def test_pos_exit_confirmation_desktop_fallback_never_uses_global_uia(
    monkeypatch,
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")
    backends: list[str] = []
    dialog = SimpleNamespace(window_text=lambda: "結束程式確認")

    class FakeDesktop:
        def __init__(self, *, backend: str) -> None:
            backends.append(backend)
            if backend == "uia":
                raise AssertionError("global Desktop UIA enumeration must not run")

        def windows(self, *, title: str, visible_only: bool):
            assert title == "結束程式確認"
            assert visible_only is True
            return [dialog]

    window = SimpleNamespace(children=lambda: [], descendants=lambda: [])
    monkeypatch.setattr(automation_runner_module.sys, "platform", "win32")
    monkeypatch.setitem(sys.modules, "pywinauto", SimpleNamespace(Desktop=FakeDesktop))

    dialogs = runner._find_pos_exit_confirmation_dialogs(window)

    assert dialogs == [dialog]
    assert backends == ["win32"]


def test_pos_session_ini_selection_rejects_preexisting_process_as_launch_handoff(
    monkeypatch,
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    config.pos.startup_ini_selection_enabled = True
    config.pos.startup_ini_profile = r"c:\tkhspa\tkhspa-正式.ini"
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")
    runner._pos_launch_evidence = {
        "launch_requested": True,
        "requested_at": "2026-09-02T08:00:00.000000+00:00",
    }
    runner._pos_startup_ini_profile_verified_this_run = config.pos.startup_ini_profile
    runner._pos_startup_ini_process_id_verified_this_run = 1111
    runner._pos_startup_ini_process_created_at_verified_this_run = "2026-09-02T08:00:01.000000+00:00"
    creation_times = {
        1111: "2026-09-02T08:00:01.000000+00:00",
        2222: "2026-09-02T07:59:00.000000+00:00",
    }
    monkeypatch.setattr(
        automation_runner_module,
        "_safe_process_creation_time",
        lambda process_id: creation_times.get(process_id),
    )

    with pytest.raises(RuntimeError, match="POS_ACTIVE_SESSION_PROFILE_UNVERIFIED"):
        runner._validate_or_record_pos_session_profile(config, FakeProcessBoundPosWindow(2222))

    assert not (Path(config.app.state_dir) / "pos_session_profile.json").exists()


def test_pos_session_marker_rejects_reused_pid_with_different_creation_time(monkeypatch, tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    config.pos.startup_ini_selection_enabled = True
    config.pos.startup_ini_profile = r"c:\tkhspa\tkhspa-正式.ini"
    marker_path = Path(config.app.state_dir) / "pos_session_profile.json"
    marker_path.parent.mkdir(parents=True, exist_ok=True)
    marker_path.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "profile": config.pos.startup_ini_profile,
                "process_id": 4242,
                "process_created_at": "old-process-instance",
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(
        automation_runner_module,
        "_safe_process_creation_time",
        lambda _process_id: "new-process-instance",
        raising=False,
    )
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")

    with pytest.raises(RuntimeError, match="POS_ACTIVE_SESSION_PROFILE_UNVERIFIED"):
        runner._validate_or_record_pos_session_profile(config, FakeProcessBoundPosWindow(4242))


def test_pos_session_marker_rejects_process_identity_without_creation_time(monkeypatch, tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    config.pos.startup_ini_selection_enabled = True
    config.pos.startup_ini_profile = r"c:\tkhspa\tkhspa-正式.ini"
    marker_path = Path(config.app.state_dir) / "pos_session_profile.json"
    marker_path.parent.mkdir(parents=True, exist_ok=True)
    marker_path.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "profile": config.pos.startup_ini_profile,
                "process_id": 4242,
                "process_created_at": None,
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(automation_runner_module, "_safe_process_creation_time", lambda _process_id: None)
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")

    with pytest.raises(RuntimeError, match="POS_ACTIVE_SESSION_PROFILE_UNVERIFIED"):
        runner._validate_or_record_pos_session_profile(config, FakeProcessBoundPosWindow(4242))


def test_automation_runner_uses_keyboard_fallback_when_ini_select_is_noop(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    combo = FakeKeyboardSelectableStartupIniCombo()
    window, confirm = _pos_startup_ini_window_with_combo(combo)

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
    )

    handled = runner._prepare_pos_window_for_login(config, window)

    assert handled is window
    assert combo.focused is True
    assert combo.dropdown_opened is True
    assert combo.selected_value == r"c:\tkhspa\tkhspa -測試.ini"
    assert confirm.clicked is True


def test_automation_runner_uses_keyboard_fallback_for_production_ini(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    # 2.1.2 persisted this obsolete profile name.  The 2026-08-06 POS
    # diagnostic exposes only tkhspa-正式.ini and lists it second.
    config.pos.startup_ini_profile = r"c:\tkhspa\tkhspa-正式區.ini"
    combo = FakeKeyboardSelectableStartupIniCombo()
    window, confirm = _pos_startup_ini_window_with_combo(combo)

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
    )

    handled = runner._prepare_pos_window_for_login(config, window)

    assert handled is window
    assert combo.dropdown_opened is True
    assert combo.selected_value == r"c:\tkhspa\tkhspa-正式.ini"
    assert confirm.clicked is True


def test_automation_runner_keeps_already_selected_pos_startup_ini(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    target_profile = r"c:\tkhspa\tkhspa-正式.ini"
    config.pos.startup_ini_profile = target_profile
    combo = FakeAlreadySelectedStartupIniCombo(target_profile)
    window, confirm = _pos_startup_ini_window_with_combo(combo)

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
    )

    handled = runner._prepare_pos_window_for_login(config, window)

    assert handled is window
    assert combo.selected_value == target_profile
    assert confirm.clicked is True


def test_wait_for_reconnected_pos_window_handles_startup_ini_and_writes_log(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    chooseini_window, combo, confirm = _pos_startup_ini_window()
    ready_window = _ready_pos_window()
    windows = [chooseini_window, ready_window]

    def connect_pos_window(**_kwargs):  # type: ignore[no-untyped-def]
        return windows.pop(0)

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
    )

    window = runner._wait_for_reconnected_pos_window(config)
    log_path = runner.runtime_paths.logs_dir / f"automation_pos_startup_ini_{runner.run_date:%Y%m%d}.jsonl"
    log_text = log_path.read_text(encoding="utf-8")
    records = [json.loads(line) for line in log_text.splitlines()]
    selection_record = next(record for record in records if record["event"] == "selection_verified")

    assert window is ready_window
    assert combo.selected_value == r"c:\tkhspa\tkhspa -測試.ini"
    assert confirm.clicked is True
    assert "dialog_detected" in log_text
    assert "selection_verified" in log_text
    assert "confirm_clicked" in log_text
    assert "selection_process_id" in selection_record
    assert "selection_process_created_at" in selection_record
    assert selection_record["launch_requested"] is False
    assert "launch_requested_at" in selection_record


def test_automation_runner_does_not_confirm_unverified_pos_startup_ini_selection(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    combo = FakeUnverifiableStartupIniCombo()
    window, confirm = _pos_startup_ini_window_with_combo(combo)

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
    )

    with pytest.raises(RuntimeError, match="POS_STARTUP_INI_SELECTION_FAILED"):
        runner._prepare_pos_window_for_login(config, window)
    log_path = runner.runtime_paths.logs_dir / f"automation_pos_startup_ini_{runner.run_date:%Y%m%d}.jsonl"
    log_text = log_path.read_text(encoding="utf-8")
    probe_paths = list(runner.runtime_paths.logs_dir.glob("ui_probe_pos_startup_ini_selection_not_verified_*.json"))

    assert confirm.clicked is False
    assert "selection_not_verified" in log_text
    assert probe_paths


def test_automation_runner_does_not_treat_ini_texts_list_as_selected_value(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    combo = FakeTextsOnlyStartupIniCombo()
    window, confirm = _pos_startup_ini_window_with_combo(combo)

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
    )

    with pytest.raises(RuntimeError, match="POS_STARTUP_INI_SELECTION_FAILED"):
        runner._prepare_pos_window_for_login(config, window)

    assert confirm.clicked is False


def test_automation_runner_does_not_click_generic_confirm_without_ini_dialog(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    confirm = FakePosControl("確　定", "Button")
    window = FakePosControl("SPA-POS", "Window", children=[confirm])

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
    )

    handled = runner._prepare_pos_window_for_login(config, window)

    assert handled is window
    assert confirm.clicked is False


def _write_r14_snapshot(
    path: Path,
    *,
    report_date: date,
    item_actuals: dict[str, float],
    include_n006: bool = True,
) -> Path:
    branches = [
        "站前4樓",
        "站前11樓",
        "忠孝國際醫學3樓",
        "忠孝7樓",
        "忠孝健康7樓",
    ]
    if include_n006:
        branches.append("忠孝預防醫學3樓")
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["F2"] = report_date
    month_label = report_date.strftime("%Y/%m")
    for branch in branches:
        sheet = workbook.create_sheet(branch)
        sheet["B2"] = "凱惠料號"
        sheet["C2"] = "品名"
        sheet["H2"] = month_label
        sheet["H3"] = "Actual"
        sheet["B4"] = "ITEM001"
        sheet["C4"] = "高波動品項"
        sheet["H4"] = item_actuals.get(branch, 0)
        sheet["B5"] = "ITEM002"
        sheet["C5"] = "穩定品項"
        sheet["H5"] = 50
    workbook.create_sheet("領用表")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def test_automation_runner_r04_uses_pos_host_dated_downloads_dir(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R04"
    captured: dict[str, object] = {}

    class FakeAutomator:
        def __init__(self, *_args, output_dir: Path, **_kwargs):  # type: ignore[no-untyped-def]
            captured["output_dir"] = output_dir

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            captured["task_id"] = output.task_id
            captured["start_date"] = output.start_date
            captured["end_date"] = output.end_date
            captured["report_menu_text"] = report.report_menu_text
            captured["branch_mode"] = report.branch_mode
            output_path = Path(captured["output_dir"]) / output.output_filename
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_bytes(b"r04-xls")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        run_date=date(2026, 7, 9),
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run()

    expected_dir = tmp_path / "downloads" / "20260709"
    expected_path = expected_dir / "預約資料統計報表-20260709-20260808.xls"
    assert summary.ok is True
    assert captured == {
        "output_dir": expected_dir,
        "task_id": "R04",
        "start_date": "2026/07/09",
        "end_date": "2026/08/08",
        "report_menu_text": "預約紀錄查詢統計表",
        "branch_mode": "multi_select",
    }
    assert expected_path.exists()


def test_automation_runner_r04_uploads_from_pos_host_download_path(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    config.google_drive.upload_enabled = True
    for report in config.reports:
        report.enabled = report.id == "R04"
        report.upload_enabled = report.id == "R04"
    upload_calls: list[tuple[Path, str, str]] = []

    class FakeAutomator:
        def __init__(self, *_args, output_dir: Path, **_kwargs):  # type: ignore[no-untyped-def]
            self.output_dir = output_dir

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            output_path = self.output_dir / output.output_filename
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_bytes(b"r04-xls")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
            )

    class FakeUploader:
        def upload(self, file_path: Path, folder_id: str, name: str) -> DriveUploadResult:
            upload_calls.append((file_path, folder_id, name))
            return DriveUploadResult(
                success=True,
                drive_file_id="drive-r04",
                folder_id=folder_id,
                uploaded_name=name,
                size=file_path.stat().st_size,
                mime_type="application/vnd.ms-excel",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        run_date=date(2026, 7, 9),
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        drive_uploader_factory=lambda _config: FakeUploader(),
    )

    summary = runner.run()

    expected_path = tmp_path / "downloads" / "20260709" / "預約資料統計報表-20260709-20260808.xls"
    assert summary.ok is True
    assert upload_calls == [
        (
            expected_path,
            "1B3_KGkQ3nMNA0MMJxWVqLi1EvzKKZhGV",
            "預約資料統計報表-20260709-20260808.xls",
        )
    ]
    state_text = (tmp_path / "state" / "20260709" / "run_state_latest.json").read_text(encoding="utf-8")
    assert "uploaded" in state_text
    assert "drive-r04" in state_text


def test_automation_runner_continues_after_failed_report(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id in {"R01", "R02", "R03"}
    calls: list[str] = []
    progress_messages: list[str] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            if output.task_id == "R02":
                raise ReportAutomationError("CONTROL_NOT_CLICKABLE", "控制項無法點擊：檢視報表")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        run_source="windows_task_scheduler",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run(on_progress=lambda event: progress_messages.append(event.message))

    assert calls == ["R01", "R02", "R03"]
    assert summary.ok is False
    assert summary.completed == 2
    assert summary.total == 3
    assert summary.error_code == "PARTIAL_REPORT_RUN_FAILED"
    assert summary.details is not None
    assert "任務：R02" in summary.details
    assert any("R02 失敗" in message for message in progress_messages)
    assert progress_messages[-1].startswith("已完成 2 個報表任務")


def test_automation_runner_reconnects_after_failed_report_before_next_pos_task(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id in {"R01", "R02", "R03"}
    calls: list[str] = []
    connected_roots: list[tuple[str, ...]] = []
    automator_windows: list[FakePosControl] = []
    progress_messages: list[str] = []

    def connect_pos_window(**_kwargs):  # type: ignore[no-untyped-def]
        roots = ("統計報表",)
        connected_roots.append(roots)
        return _ready_pos_window(*roots)

    class FakeAutomator:
        def __init__(self, window, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            automator_windows.append(window)

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            if output.task_id == "R02":
                raise ReportAutomationError("BRANCH_CONTROL_NOT_FOUND", "找不到分店下拉選項：所有分店")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run(on_progress=lambda event: progress_messages.append(event.message))

    assert calls == ["R01", "R02", "R03"]
    assert len(connected_roots) >= 1
    assert len(automator_windows) >= 2
    assert summary.ok is False
    assert summary.completed == 2
    assert any(
        "重新連接並確認 SPA-POS 主選單" in message or "現有 SPA-POS 主畫面仍可讀取" in message
        for message in progress_messages
    )


def test_automation_runner_blocks_duplicate_drive_folder_and_filename_targets(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.google_drive.upload_enabled = True
    for report in config.reports:
        report.enabled = report.id in {"R02", "R03"}
    config.drive_targets.targets["R03"].folder_id_or_url = config.drive_targets.targets["R02"].folder_id_or_url
    connect_calls = 0

    def connect_pos_window(**_kwargs):  # type: ignore[no-untyped-def]
        nonlocal connect_calls
        connect_calls += 1
        return _ready_pos_window()

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
    )

    summary = runner.run()

    assert connect_calls == 0
    assert summary.ok is False
    assert summary.completed == 0
    assert {failure.task_id for failure in summary.failures} == {"R02", "R03"}
    assert {failure.error_code for failure in summary.failures} == {"DUPLICATE_DRIVE_UPLOAD_TARGET"}
    assert all("同一個 Google Drive folder 且檔名相同" in failure.message for failure in summary.failures)


def test_automation_runner_stops_pos_tasks_when_reconnect_after_failure_fails(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.app.logs_dir = str(tmp_path / "logs")
    config.pos.startup_wait_seconds = 1
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id in {"R01", "R02", "R03"}
    calls: list[str] = []
    connect_calls = 0
    progress_messages: list[str] = []
    current_window = _ready_pos_window("統計報表")

    def connect_pos_window(**_kwargs):  # type: ignore[no-untyped-def]
        nonlocal connect_calls
        connect_calls += 1
        if connect_calls > 1:
            raise UiProbeError("Cannot connect to SPA-POS")
        return current_window

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            if output.task_id == "R02":
                current_window.children_controls = []
                raise ReportAutomationError("REPORT_ROOT_MENU_NOT_FOUND", "找不到控制項：統計報表")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run(on_progress=lambda event: progress_messages.append(event.message))

    assert calls == ["R01", "R02"]
    assert summary.ok is False
    assert summary.completed == 1
    assert summary.total == 3
    assert summary.failures[0].task_id == "R02"
    assert "已停止後續 POS 任務" in summary.failures[0].message
    assert summary.failures[0].diagnostic_path is not None
    diagnostic_path = Path(summary.failures[0].diagnostic_path)
    payload = json.loads(diagnostic_path.read_text(encoding="utf-8"))
    assert payload["error"]["code"] == "POS_CONNECTION_FAILED"
    assert "重新連接階段" in payload["note"]
    assert [output["task_id"] for output in payload["outputs"]] == ["R03"]
    assert [failure.task_id for failure in summary.failures] == ["R02", "R03"]
    assert summary.failures[1].error_code == "POS_CONNECTION_FAILED"
    assert any("已停止後續 POS 任務" in message for message in progress_messages)
    state = RunStateStore.default_for_config(config, run_date=runner.run_date).load()
    assert state is not None
    r03_state = next(output for output in state.outputs.values() if output.task_id == "R03")
    assert r03_state.status == "failed"
    assert r03_state.error_code == "POS_CONNECTION_FAILED"


def test_automation_runner_reuses_current_ready_pos_window_when_fresh_reconnect_fails(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id in {"R01", "R02", "R03"}
    calls: list[str] = []
    connect_calls = 0
    progress_messages: list[str] = []
    current_window = _ready_pos_window("統計報表")

    def connect_pos_window(**_kwargs):  # type: ignore[no-untyped-def]
        nonlocal connect_calls
        connect_calls += 1
        if connect_calls > 1:
            raise UiProbeError("Cannot connect to SPA-POS")
        return current_window

    class FakeAutomator:
        def __init__(self, window, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            assert window is current_window

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            if output.task_id == "R02":
                raise ReportAutomationError("REPORT_ROOT_MENU_NOT_FOUND", "找不到控制項：統計報表")
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run(on_progress=lambda event: progress_messages.append(event.message))

    assert calls == ["R01", "R02", "R03"]
    assert connect_calls == 1
    assert summary.ok is False
    assert summary.completed == 2
    assert [failure.task_id for failure in summary.failures] == ["R02"]
    assert any("現有 SPA-POS 主畫面仍可讀取" in message for message in progress_messages)


def test_automation_runner_stops_pos_tasks_but_keeps_r14_blocked_by_r13_semantics(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.pos.startup_wait_seconds = 1
    _disable_uploads(config)
    config.r14_email.enabled = False
    for report in config.reports:
        report.enabled = report.id in {"R02", "R13", "R14"}
    calls: list[str] = []
    connect_calls = 0
    current_window = _ready_pos_window("統計報表", "庫存管理")

    def connect_pos_window(**_kwargs):  # type: ignore[no-untyped-def]
        nonlocal connect_calls
        connect_calls += 1
        if connect_calls > 1:
            raise UiProbeError("Cannot connect to SPA-POS")
        return current_window

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            if output.task_id == "R02":
                current_window.children_controls = []
                raise ReportAutomationError("REPORT_ROOT_MENU_NOT_FOUND", "找不到控制項：統計報表")
            raise AssertionError(f"{output.task_id} POS task should have been stopped")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        run_date=date(2026, 6, 17),
    )

    summary = runner.run()

    assert calls == ["R02"]
    assert summary.ok is False
    assert [failure.task_id for failure in summary.failures] == ["R02", "R13", "R14"]
    failures_by_task = {failure.task_id: failure for failure in summary.failures}
    assert failures_by_task["R13"].error_code == "POS_CONNECTION_FAILED"
    assert failures_by_task["R14"].error_code == "R14_BLOCKED_BY_R13_FAILED"
    assert "R13 raw data 未產生" in failures_by_task["R14"].message
    state = RunStateStore.default_for_config(config, run_date=runner.run_date).load()
    assert state is not None
    r14_state = next(output for output in state.outputs.values() if output.task_id == "R14")
    assert r14_state.status == "failed"
    assert r14_state.error_code == "R14_BLOCKED_BY_R13_FAILED"


def test_automation_runner_runs_r14_without_connecting_pos(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.r14_email.enabled = False
    config.r14_transform.template_path = str(ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx")
    config.r14_transform.raw_search_dir = str(ROOT / "tests" / "R14_TEST")
    for report in config.reports:
        report.enabled = report.id == "R14"
        if report.id == "R14":
            report.upload_enabled = False

    def fail_connect(**_kwargs):  # type: ignore[no-untyped-def]
        raise AssertionError("R14 offline transform must not connect to POS")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=fail_connect,
        run_date=date(2026, 6, 9),
    )

    summary = runner.run()

    assert summary.ok is True
    assert summary.completed == 1
    assert summary.total == 1
    output_path = tmp_path / "downloads" / "R14" / "20260609" / "診所stock status - 2026 demand planning-0608.xlsx"
    assert output_path.exists()


def test_automation_runner_r14_missing_month_end_snapshot_still_outputs_with_blank_planning_warning(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.email.enabled = True
    config.r14_email.enabled = True
    template_path = tmp_path / "診所stock status - 2026 demand planning-template.xlsx"
    template_path.write_bytes(
        (ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx").read_bytes()
    )
    workbook = load_workbook(template_path)
    may_group = _summary_actual_group_columns(workbook["Summary"], "2026/05", allow_legacy=True)
    workbook["Summary"].delete_cols(may_group[0], len(may_group))
    workbook.save(template_path)
    config.r14_transform.template_path = str(template_path)
    config.r14_transform.raw_search_dir = str(ROOT / "tests" / "R14_TEST")
    partial_history = (
        tmp_path
        / "downloads"
        / "R14"
        / "20260523"
        / "診所stock status - 2026 demand planning-0523.xlsx"
    )
    partial_history.parent.mkdir(parents=True)
    partial_history.write_bytes(
        (ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0608.xlsx").read_bytes()
    )
    for report in config.reports:
        report.enabled = report.id == "R14"
        report.upload_enabled = False
    sent_messages: list[tuple[str, str, list[Path]]] = []

    class FakeGmailSender:
        def send(self, _settings, *, subject, body, attachments=None):  # type: ignore[no-untyped-def]
            sent_messages.append((subject, body, list(attachments or [])))
            return SimpleNamespace(ok=True, message="sent", gmail_message_id="gmail-r14-warning")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("R14 offline transform must not connect to POS")
        ),
        gmail_sender_factory=lambda _config: FakeGmailSender(),
        run_date=date(2026, 6, 9),
    )

    summary = runner.run()

    output_path = tmp_path / "downloads" / "R14" / "20260609" / "診所stock status - 2026 demand planning-0608.xlsx"
    assert summary.ok is True
    assert output_path.exists()
    output_workbook = load_workbook(output_path, data_only=False)
    branch_sheet = output_workbook["站前4樓"]
    item_row = _find_workbook_item_row(branch_sheet, "6050010")
    planning_columns: dict[str, int] = {}
    for col in range(1, branch_sheet.max_column + 1):
        for row in (2, 3):
            label = branch_sheet.cell(row, col).value
            if label in {"Forecast", "下單數", "安庫"}:
                planning_columns[str(label)] = col
    assert branch_sheet.cell(item_row, planning_columns["Forecast"]).value is None
    assert branch_sheet.cell(item_row, planning_columns["下單數"]).value is None
    assert branch_sheet.cell(item_row, planning_columns["安庫"]).value is None
    assert len(sent_messages) == 1
    assert "缺少前月月底歷史資料" in sent_messages[0][1]
    assert "最新找到的前月歷史檔名日期為 2026/05/23" in sent_messages[0][1]
    assert "Forecast、安庫與下單數已留白" in sent_messages[0][1]
    state_template = tmp_path / "state" / "r14_templates" / template_path.name
    state_workbook = load_workbook(state_template, data_only=False)
    with pytest.raises(AssertionError, match="2026/05"):
        _summary_actual_group_columns(state_workbook["Summary"], "2026/05")


def test_automation_runner_reads_r14_inventory_before_configured_weekday_transform(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.r14_email.enabled = False
    template_path = tmp_path / "r14-template.xlsx"
    template_path.write_bytes((ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx").read_bytes())
    original_template_bytes = template_path.read_bytes()
    config.r14_transform.template_path = str(template_path)
    config.r14_transform.raw_search_dir = str(ROOT / "tests" / "R14_TEST")
    config.r14_inventory_source.enabled = True
    config.r14_inventory_source.apply_weekday = "Tuesday"
    for report in config.reports:
        report.enabled = report.id == "R14"
        if report.id == "R14":
            report.upload_enabled = False
    calls = 0

    class FakeInventoryClient:
        def read_r14_inventory(self, settings):  # type: ignore[no-untyped-def]
            nonlocal calls
            calls += 1
            assert settings.sheet_name == "Summary"
            return SimpleNamespace(inventories={"站前4樓": {"6050010": 321}})

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("R14 offline transform must not connect to POS")
        ),
        r14_inventory_client_factory=lambda _config: FakeInventoryClient(),
        run_date=date(2026, 6, 9),
    )

    summary = runner.run()

    output_path = tmp_path / "downloads" / "R14" / "20260609" / "診所stock status - 2026 demand planning-0608.xlsx"
    workbook = load_workbook(output_path, data_only=False)
    row = 4
    while str(workbook["站前4樓"].cell(row, 2).value).strip() != "6050010":
        row += 1
    assert summary.ok is True
    assert calls == 1
    assert workbook["站前4樓"].cell(row, 6).value == 321
    assert template_path.read_bytes() == original_template_bytes
    state_template = tmp_path / "state" / "r14_templates" / template_path.name
    assert state_template.exists()
    state_workbook = load_workbook(state_template, data_only=False)
    may_group = _summary_actual_group_columns(state_workbook["Summary"], "2026/05")
    june_group = _summary_actual_group_columns(state_workbook["Summary"], "2026/06")
    assert june_group == list(range(may_group[-1] + 1, may_group[-1] + 8))


def test_automation_runner_r14_uses_unique_output_path_when_same_day_file_exists(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.r14_email.enabled = False
    config.r14_transform.template_path = str(ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx")
    config.r14_transform.raw_search_dir = str(ROOT / "tests" / "R14_TEST")
    for report in config.reports:
        report.enabled = report.id == "R14"
        report.upload_enabled = False

    archive_dir = tmp_path / "downloads" / "R14" / "20260609"
    archive_dir.mkdir(parents=True)
    existing_path = archive_dir / "診所stock status - 2026 demand planning-0608.xlsx"
    existing_path.write_bytes(b"already-open-or-existing")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("R14 offline transform must not connect to POS")
        ),
        run_date=date(2026, 6, 9),
    )

    summary = runner.run()

    assert summary.ok is True
    assert existing_path.read_bytes() == b"already-open-or-existing"
    assert (archive_dir / "診所stock status - 2026 demand planning-0608_001.xlsx").exists()


def test_automation_runner_w01_updates_r14_template_without_output_upload_or_pos(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = True
    config.r14_email.enabled = False
    template_path = tmp_path / "r14-template.xlsx"
    template_path.write_bytes((ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx").read_bytes())
    config.r14_transform.template_path = str(template_path)
    config.r14_inventory_source.enabled = True
    config.r14_inventory_source.apply_weekday = "Tuesday"
    for report in config.reports:
        report.enabled = report.id == "W01"
        if report.id == "W01":
            report.upload_enabled = False
            report.output_filename = ""
    calls = 0

    class FakeInventoryClient:
        def read_r14_inventory(self, settings):  # type: ignore[no-untyped-def]
            nonlocal calls
            calls += 1
            assert settings.sheet_name == "Summary"
            return SimpleNamespace(rows_read=1, inventories={"站前4樓": {"6050010": 321}})

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("W01 must not connect to POS")
        ),
        drive_uploader_factory=lambda _config: (_ for _ in ()).throw(
            AssertionError("W01 must not upload to Google Drive")
        ),
        r14_inventory_client_factory=lambda _config: FakeInventoryClient(),
        run_date=date(2026, 6, 9),
    )

    summary = runner.run()

    workbook = load_workbook(template_path, data_only=False)
    row = 4
    while str(workbook["站前4樓"].cell(row, 2).value).strip() != "6050010":
        row += 1
    assert summary.ok is True
    assert calls == 1
    assert summary.completed == 1
    assert workbook["站前4樓"].cell(row, 6).value == 321


def test_automation_runner_manual_run_forces_w01_on_non_configured_weekday(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.r14_email.enabled = False
    template_path = tmp_path / "r14-template.xlsx"
    template_path.write_bytes((ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx").read_bytes())
    config.r14_transform.template_path = str(template_path)
    config.r14_inventory_source.enabled = True
    config.r14_inventory_source.apply_weekday = "Friday"
    for report in config.reports:
        report.enabled = report.id == "W01"
        report.upload_enabled = False
    calls = 0

    class FakeInventoryClient:
        def read_r14_inventory(self, _settings):  # type: ignore[no-untyped-def]
            nonlocal calls
            calls += 1
            return SimpleNamespace(
                rows_read=1,
                inventories={"站前4樓": {"6050010": 321, "NEWGOOGLE003": 8}},
                inventory_date=date(2026, 6, 25),
                item_names={"NEWGOOGLE003": "W01 新增品項"},
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("W01 must not connect to POS")
        ),
        r14_inventory_client_factory=lambda _config: FakeInventoryClient(),
        run_source="gui_manual",
        run_date=date(2026, 6, 9),
    )

    summary = runner.run()

    workbook = load_workbook(template_path, data_only=False)
    row = 4
    while str(workbook["站前4樓"].cell(row, 2).value).strip() != "6050010":
        row += 1
    assert summary.ok is True
    assert summary.completed == 1
    assert calls == 1
    assert workbook["站前4樓"].cell(row, 6).value == 321


def test_automation_runner_scheduler_run_does_not_force_w01_on_non_configured_weekday(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.r14_inventory_source.apply_weekday = "Friday"
    for report in config.reports:
        report.enabled = report.id == "W01"

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        r14_inventory_client_factory=lambda _config: (_ for _ in ()).throw(
            AssertionError("scheduled non-Friday run must not read W01 inventory")
        ),
        run_source="windows_task_scheduler",
        run_date=date(2026, 6, 9),
    )

    summary = runner.run()

    assert summary.ok is False
    assert summary.error_code == "NO_ENABLED_REPORTS"


def test_automation_runner_r14_uses_w01_updated_template_without_second_sheet_read(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.r14_email.enabled = False
    template_path = tmp_path / "r14-template.xlsx"
    template_path.write_bytes((ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx").read_bytes())
    config.r14_transform.template_path = str(template_path)
    config.r14_transform.raw_search_dir = str(ROOT / "tests" / "R14_TEST")
    config.r14_inventory_source.enabled = True
    config.r14_inventory_source.apply_weekday = "Tuesday"
    for report in config.reports:
        report.enabled = report.id in {"W01", "R14"}
        report.upload_enabled = False
    calls = 0

    class FakeInventoryClient:
        def read_r14_inventory(self, _settings):  # type: ignore[no-untyped-def]
            nonlocal calls
            calls += 1
            return SimpleNamespace(
                rows_read=1,
                inventories={"站前4樓": {"6050010": 321, "NEWGOOGLE003": 8}},
                inventory_date=date(2026, 6, 25),
                item_names={"NEWGOOGLE003": "W01 新增品項"},
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("W01/R14 local transforms must not connect to POS")
        ),
        r14_inventory_client_factory=lambda _config: FakeInventoryClient(),
        run_date=date(2026, 6, 9),
    )

    summary = runner.run()

    output_path = tmp_path / "downloads" / "R14" / "20260609" / "診所stock status - 2026 demand planning-0608.xlsx"
    workbook = load_workbook(output_path, data_only=False)
    row = 4
    while str(workbook["站前4樓"].cell(row, 2).value).strip() != "6050010":
        row += 1
    new_item_row = 4
    while str(workbook["站前4樓"].cell(new_item_row, 2).value).strip() != "NEWGOOGLE003":
        new_item_row += 1
    assert summary.ok is True
    assert calls == 1
    assert summary.completed == 2
    assert workbook["Summary"]["F2"].value.date() == date(2026, 6, 25)
    assert workbook["站前4樓"].cell(row, 6).value == 321
    assert workbook["站前4樓"].cell(new_item_row, 3).value == "W01 新增品項"
    assert workbook["站前4樓"].cell(new_item_row, 6).value == 8


def test_automation_runner_w01_copies_packaged_template_to_writable_dir_before_sync(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.r14_email.enabled = False
    packaged_dir = tmp_path / "_internal" / "config_templates" / "templates"
    packaged_dir.mkdir(parents=True)
    packaged_template = packaged_dir / "診所stock status - 2026 demand planning-template.xlsx"
    packaged_template.write_bytes(
        (ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx").read_bytes()
    )
    original_packaged_bytes = packaged_template.read_bytes()
    writable_template_dir = tmp_path / "programdata" / "templates"
    config.r14_transform.template_path = ""
    config.r14_transform.template_search_dir = str(writable_template_dir)
    config.r14_transform.raw_search_dir = str(ROOT / "tests" / "R14_TEST")
    config.r14_inventory_source.enabled = True
    config.r14_inventory_source.apply_weekday = "Tuesday"
    for report in config.reports:
        report.enabled = report.id in {"W01", "R14"}
        report.upload_enabled = False

    monkeypatch.setattr(automation_runner_module, "_packaged_r14_template_dirs", lambda: [packaged_dir])
    calls = 0

    class FakeInventoryClient:
        def read_r14_inventory(self, _settings):  # type: ignore[no-untyped-def]
            nonlocal calls
            calls += 1
            return SimpleNamespace(rows_read=1, inventories={"站前4樓": {"6050010": 321}})

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("W01/R14 local transforms must not connect to POS")
        ),
        r14_inventory_client_factory=lambda _config: FakeInventoryClient(),
        run_date=date(2026, 6, 9),
        run_source="gui_manual",
    )

    summary = runner.run()

    runtime_template = writable_template_dir / packaged_template.name
    output_path = tmp_path / "downloads" / "R14" / "20260609" / "診所stock status - 2026 demand planning-0608.xlsx"
    workbook = load_workbook(output_path, data_only=False)
    row = 4
    while str(workbook["站前4樓"].cell(row, 2).value).strip() != "6050010":
        row += 1
    assert summary.ok is True
    assert calls == 1
    assert runtime_template.exists()
    assert packaged_template.read_bytes() == original_packaged_bytes
    assert workbook["站前4樓"].cell(row, 6).value == 321


def test_automation_runner_w01_makes_existing_runtime_template_writable_before_sync(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.r14_email.enabled = False
    writable_template_dir = tmp_path / "programdata" / "templates"
    writable_template_dir.mkdir(parents=True)
    runtime_template = writable_template_dir / "診所stock status - 2026 demand planning-template.xlsx"
    runtime_template.write_bytes(
        (ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx").read_bytes()
    )
    runtime_template.chmod(stat.S_IREAD)
    config.r14_transform.template_path = ""
    config.r14_transform.template_search_dir = str(writable_template_dir)
    config.r14_inventory_source.enabled = True
    config.r14_inventory_source.apply_weekday = "Tuesday"
    for report in config.reports:
        report.enabled = report.id == "W01"
        report.upload_enabled = False

    class FakeInventoryClient:
        def read_r14_inventory(self, _settings):  # type: ignore[no-untyped-def]
            return SimpleNamespace(rows_read=1, inventories={"站前4樓": {"6050010": 321}})

    def assert_writable_template(template_path, branch_inventory):  # type: ignore[no-untyped-def]
        assert template_path == runtime_template
        assert template_path.stat().st_mode & stat.S_IWUSR
        return SimpleNamespace(template_path=template_path, inventory_updated=1, inventory_unmatched=0)

    monkeypatch.setattr(automation_runner_module, "sync_r14_template_inventory", assert_writable_template)

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("W01 local transform must not connect to POS")
        ),
        r14_inventory_client_factory=lambda _config: FakeInventoryClient(),
        run_date=date(2026, 6, 9),
        run_source="gui_manual",
    )

    summary = runner.run()

    assert summary.ok is True
    assert summary.completed == 1


def test_automation_runner_w01_falls_back_to_state_template_when_runtime_template_is_not_writable(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.r14_email.enabled = False
    locked_template_dir = tmp_path / "programdata" / "templates"
    locked_template_dir.mkdir(parents=True)
    locked_template = locked_template_dir / "診所stock status - 2026 demand planning-template.xlsx"
    locked_template.write_bytes(
        (ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx").read_bytes()
    )
    original_locked_bytes = locked_template.read_bytes()
    config.r14_transform.template_path = ""
    config.r14_transform.template_search_dir = str(locked_template_dir)
    config.r14_transform.raw_search_dir = str(ROOT / "tests" / "R14_TEST")
    config.r14_inventory_source.enabled = True
    config.r14_inventory_source.apply_weekday = "Tuesday"
    for report in config.reports:
        report.enabled = report.id in {"W01", "R14"}
        report.upload_enabled = False

    original_ensure_writable = automation_runner_module._ensure_writable_file_for_update

    def fail_only_for_locked_template(path: Path) -> None:
        if path == locked_template:
            raise R14TransformError(
                "W01_TEMPLATE_NOT_WRITABLE",
                f"R14 模板檔無法改成可寫入狀態，請確認檔案權限或唯讀屬性：{path}",
            )
        original_ensure_writable(path)

    monkeypatch.setattr(
        automation_runner_module,
        "_ensure_writable_file_for_update",
        fail_only_for_locked_template,
    )
    calls = 0

    class FakeInventoryClient:
        def read_r14_inventory(self, _settings):  # type: ignore[no-untyped-def]
            nonlocal calls
            calls += 1
            return SimpleNamespace(rows_read=1, inventories={"站前4樓": {"6050010": 321}})

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("W01/R14 local transforms must not connect to POS")
        ),
        r14_inventory_client_factory=lambda _config: FakeInventoryClient(),
        run_date=date(2026, 6, 9),
        run_source="gui_manual",
    )

    summary = runner.run()

    state_template = tmp_path / "state" / "r14_templates" / locked_template.name
    output_path = tmp_path / "downloads" / "R14" / "20260609" / "診所stock status - 2026 demand planning-0608.xlsx"
    workbook = load_workbook(output_path, data_only=False)
    row = 4
    while str(workbook["站前4樓"].cell(row, 2).value).strip() != "6050010":
        row += 1
    assert summary.ok is True
    assert calls == 1
    assert state_template.exists()
    assert locked_template.read_bytes() == original_locked_bytes
    assert workbook["站前4樓"].cell(row, 6).value == 321


def test_automation_runner_w01_falls_back_to_state_template_when_explicit_template_is_not_writable(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.r14_email.enabled = False
    locked_template = tmp_path / "programdata" / "templates" / "診所stock status - 2026 demand planning-template.xlsx"
    locked_template.parent.mkdir(parents=True)
    locked_template.write_bytes(
        (ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx").read_bytes()
    )
    original_locked_bytes = locked_template.read_bytes()
    config.r14_transform.template_path = str(locked_template)
    config.r14_transform.raw_search_dir = str(ROOT / "tests" / "R14_TEST")
    config.r14_inventory_source.enabled = True
    config.r14_inventory_source.apply_weekday = "Tuesday"
    for report in config.reports:
        report.enabled = report.id in {"W01", "R14"}
        report.upload_enabled = False

    original_ensure_writable = automation_runner_module._ensure_writable_file_for_update

    def fail_only_for_locked_template(path: Path) -> None:
        if path == locked_template:
            raise R14TransformError("W01_TEMPLATE_NOT_WRITABLE", f"template denied: {path}")
        original_ensure_writable(path)

    monkeypatch.setattr(
        automation_runner_module,
        "_ensure_writable_file_for_update",
        fail_only_for_locked_template,
    )

    class FakeInventoryClient:
        def read_r14_inventory(self, _settings):  # type: ignore[no-untyped-def]
            return SimpleNamespace(rows_read=1, inventories={"站前4樓": {"6050010": 321}})

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("W01/R14 local transforms must not connect to POS")
        ),
        r14_inventory_client_factory=lambda _config: FakeInventoryClient(),
        run_date=date(2026, 6, 9),
        run_source="gui_manual",
    )

    summary = runner.run()

    state_template = tmp_path / "state" / "r14_templates" / locked_template.name
    output_path = tmp_path / "downloads" / "R14" / "20260609" / "診所stock status - 2026 demand planning-0608.xlsx"
    workbook = load_workbook(output_path, data_only=False)
    row = 4
    while str(workbook["站前4樓"].cell(row, 2).value).strip() != "6050010":
        row += 1
    assert summary.ok is True
    assert state_template.exists()
    assert locked_template.read_bytes() == original_locked_bytes
    assert workbook["站前4樓"].cell(row, 6).value == 321


def test_automation_runner_r14_template_path_prefers_existing_state_template_over_explicit_path(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    explicit_template = tmp_path / "programdata" / "templates" / "診所stock status - 2026 demand planning-template.xlsx"
    explicit_template.parent.mkdir(parents=True)
    explicit_template.write_bytes(
        (ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx").read_bytes()
    )
    state_template = tmp_path / "state" / "r14_templates" / explicit_template.name
    state_template.parent.mkdir(parents=True)
    state_template.write_bytes(explicit_template.read_bytes())
    config.r14_transform.template_path = str(explicit_template)

    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")

    assert runner._resolve_r14_template_path() == state_template


def test_automation_runner_r14_cloud_template_overrides_existing_state_template(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    config.r14_transform.template_drive_folder_id_or_url = "https://drive.google.com/drive/u/3/folders/folder123"
    config.r14_transform.template_search_dir = str(tmp_path / "templates")
    cloud_template = tmp_path / "cloud" / "診所stock status - 2026 demand planning-0901.xlsx"
    cloud_template.parent.mkdir(parents=True)
    cloud_template.write_bytes(
        (ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx").read_bytes()
    )
    state_template = tmp_path / "state" / "r14_templates" / "診所stock status - 2026 demand planning-template.xlsx"
    state_template.parent.mkdir(parents=True)
    state_template.write_bytes(cloud_template.read_bytes())
    downloader = FakeR14TemplateDownloader(cloud_template)

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        r14_template_downloader_factory=lambda _config: downloader,
    )

    resolved = runner._resolve_r14_transform_template_path()

    assert resolved == tmp_path / "templates" / cloud_template.name
    assert downloader.calls == [
        {
            "folder_id_or_url": "https://drive.google.com/drive/u/3/folders/folder123",
            "filename_glob": "診所stock status - * demand planning-*.xlsx",
            "destination_dir": tmp_path / "templates",
        }
    ]


def test_automation_runner_r14_cloud_template_failure_does_not_fall_back_to_state_by_default(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    config.r14_transform.template_drive_folder_id_or_url = "folder123"
    state_template = tmp_path / "state" / "r14_templates" / "診所stock status - 2026 demand planning-template.xlsx"
    state_template.parent.mkdir(parents=True)
    state_template.write_bytes(
        (ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx").read_bytes()
    )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        r14_template_downloader_factory=lambda _config: FakeR14TemplateDownloader(None),
    )

    with pytest.raises(R14TransformError) as error:
        runner._resolve_r14_transform_template_path()

    assert error.value.error_code == "R14_TEMPLATE_DRIVE_FILE_MISSING"
    assert "沒有符合檔名規則" in error.value.message


def test_automation_runner_r14_cloud_template_rejects_invalid_workbook(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    config.r14_transform.template_drive_folder_id_or_url = "folder123"
    invalid_template = tmp_path / "cloud" / "診所stock status - 2026 demand planning-0901.xlsx"
    invalid_template.parent.mkdir(parents=True)
    invalid_template.write_bytes(b"not an excel workbook")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        r14_template_downloader_factory=lambda _config: FakeR14TemplateDownloader(invalid_template),
    )

    with pytest.raises(R14TransformError) as error:
        runner._resolve_r14_transform_template_path()

    assert error.value.error_code == "R14_TEMPLATE_DRIVE_WORKBOOK_INVALID"
    assert "雲端模板檔案格式不符合需求" in error.value.message


def test_automation_runner_repairs_r14_previous_month_state_from_archive(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    state_template = tmp_path / "state" / "r14_templates" / "診所stock status - 2026 demand planning-template.xlsx"
    state_template.parent.mkdir(parents=True)
    state_template.write_bytes(
        (ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx").read_bytes()
    )
    state_workbook = load_workbook(state_template)
    state_workbook["Summary"]["F2"] = date(2026, 6, 25)
    state_workbook.save(state_template)
    archive_report = tmp_path / "downloads" / "R14" / "20260630" / "診所stock status - 2026 demand planning-0630.xlsx"
    archive_report.parent.mkdir(parents=True)
    archive_report.write_bytes(
        (ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0608.xlsx").read_bytes()
    )
    newer_wrong_report = tmp_path / "downloads" / "R14" / "20260628" / "診所stock status - 2026 demand planning-0628.xlsx"
    newer_wrong_report.parent.mkdir(parents=True)
    newer_wrong_report.write_bytes(archive_report.read_bytes())
    newer_wrong_report.touch()
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")

    action = runner._sync_r14_previous_month_state_if_needed(state_template, "2026/07")

    repaired_workbook = load_workbook(state_template, data_only=False)
    summary = repaired_workbook["Summary"]
    group = _summary_actual_group_columns(summary, "2026/06")
    item_row = _find_workbook_item_row(summary, "6050010")
    assert action is not None
    assert "r14_template_previous_month_state_synced:2026/06" in action
    assert str(archive_report) in action
    assert str(newer_wrong_report) not in action
    assert len(group) == 7
    assert group == list(range(group[0], group[0] + 7))
    assert summary["F2"].value.date() == date(2026, 6, 25)
    assert summary.cell(item_row, group[0]).value == 1


def test_automation_runner_finds_r14_previous_month_archive_root(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    state_template = tmp_path / "state" / "r14_templates" / "診所stock status - 2026 demand planning-template.xlsx"
    state_template.parent.mkdir(parents=True)
    state_template.write_bytes(
        (ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx").read_bytes()
    )
    archive_report = tmp_path / "downloads" / "R14" / "診所stock status - 2026 demand planning-0630.xlsx"
    archive_report.parent.mkdir(parents=True)
    archive_report.write_bytes(
        (ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0608.xlsx").read_bytes()
    )
    monkeypatch.setattr(automation_runner_module, "r14_template_has_actual_month_state", lambda *_args, **_kwargs: True)
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")

    source = runner._find_r14_archive_for_report_date(date(2026, 6, 30), exclude_path=state_template)

    assert source == archive_report


def test_automation_runner_finds_r14_archive_root_when_downloads_config_is_dated(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    config = _load_runner_config(tmp_path)
    dated_downloads = tmp_path / "downloads" / "20260724"
    config.app.downloads_dir = str(dated_downloads)
    archive_report = tmp_path / "downloads" / "R14" / "診所stock status - 2026 demand planning-0630.xlsx"
    archive_report.parent.mkdir(parents=True)
    archive_report.write_bytes(
        (ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0608.xlsx").read_bytes()
    )
    monkeypatch.setattr(automation_runner_module, "r14_template_has_actual_month_state", lambda *_args, **_kwargs: True)
    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        run_date=date(2026, 7, 24),
    )

    source = runner._find_r14_archive_for_report_date(
        date(2026, 6, 30),
        exclude_path=tmp_path / "state" / "template.xlsx",
    )

    assert source == archive_report


def test_automation_runner_w01_does_not_overwrite_existing_state_template_with_seed(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.r14_email.enabled = False
    seed_template = tmp_path / "programdata" / "templates" / "診所stock status - 2026 demand planning-template.xlsx"
    seed_template.parent.mkdir(parents=True)
    seed_template.write_bytes(
        (ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx").read_bytes()
    )
    state_template = tmp_path / "state" / "r14_templates" / seed_template.name
    state_template.parent.mkdir(parents=True)
    state_template.write_bytes(seed_template.read_bytes())
    state_workbook = load_workbook(state_template)
    state_workbook["站前4樓"].cell(4, 6).value = 777
    state_workbook.save(state_template)
    config.r14_transform.template_path = str(seed_template)
    config.r14_inventory_source.enabled = True
    config.r14_inventory_source.apply_weekday = "Tuesday"
    for report in config.reports:
        report.enabled = report.id == "W01"
        report.upload_enabled = False

    class FakeInventoryClient:
        def read_r14_inventory(self, _settings):  # type: ignore[no-untyped-def]
            return SimpleNamespace(rows_read=1, inventories={"站前4樓": {"6050010": 321}})

    def assert_state_template_was_not_reseeded(template_path, branch_inventory):  # type: ignore[no-untyped-def]
        workbook = load_workbook(template_path, data_only=False)
        assert template_path == state_template
        assert workbook["站前4樓"].cell(4, 6).value == 777
        return SimpleNamespace(template_path=template_path, inventory_updated=1, inventory_unmatched=0)

    monkeypatch.setattr(
        automation_runner_module,
        "sync_r14_template_inventory",
        assert_state_template_was_not_reseeded,
    )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("W01 local transform must not connect to POS")
        ),
        r14_inventory_client_factory=lambda _config: FakeInventoryClient(),
        run_date=date(2026, 6, 9),
        run_source="gui_manual",
    )

    summary = runner.run()

    assert summary.ok is True


def test_automation_runner_blocks_r14_when_w01_cannot_copy_packaged_template(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.r14_email.enabled = False
    packaged_dir = tmp_path / "_internal" / "config_templates" / "templates"
    packaged_dir.mkdir(parents=True)
    packaged_template = packaged_dir / "診所stock status - 2026 demand planning-template.xlsx"
    packaged_template.write_bytes(
        (ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx").read_bytes()
    )
    writable_template_dir = tmp_path / "programdata" / "templates"
    config.r14_transform.template_path = ""
    config.r14_transform.template_search_dir = str(writable_template_dir)
    config.r14_transform.raw_search_dir = str(ROOT / "tests" / "R14_TEST")
    config.r14_inventory_source.enabled = True
    config.r14_inventory_source.apply_weekday = "Tuesday"
    for report in config.reports:
        report.enabled = report.id in {"W01", "R14"}
        report.upload_enabled = False

    monkeypatch.setattr(automation_runner_module, "_packaged_r14_template_dirs", lambda: [packaged_dir])

    def fail_copy(_src, _dst):  # type: ignore[no-untyped-def]
        raise PermissionError("denied")

    monkeypatch.setattr(automation_runner_module.shutil, "copyfile", fail_copy)

    class FakeInventoryClient:
        def read_r14_inventory(self, _settings):  # type: ignore[no-untyped-def]
            return SimpleNamespace(rows_read=1, inventories={"站前4樓": {"6050010": 321}})

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("W01/R14 local transforms must not connect to POS")
        ),
        r14_inventory_client_factory=lambda _config: FakeInventoryClient(),
        run_date=date(2026, 6, 9),
        run_source="gui_manual",
    )

    summary = runner.run()

    assert summary.ok is False
    assert [failure.task_id for failure in summary.failures] == ["W01", "R14"]
    assert summary.failures[0].error_code == "W01_TEMPLATE_COPY_FAILED"
    assert "可寫入資料夾" in summary.failures[0].message
    assert summary.failures[1].error_code == "R14_BLOCKED_BY_W01_FAILED"


def test_automation_runner_blocks_r14_when_w01_failed_on_configured_weekday(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.r14_email.enabled = False
    config.r14_transform.template_path = str(ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx")
    config.r14_transform.raw_search_dir = str(ROOT / "tests" / "R14_TEST")
    config.r14_inventory_source.enabled = True
    config.r14_inventory_source.apply_weekday = "Tuesday"
    for report in config.reports:
        report.enabled = report.id in {"W01", "R14"}
        report.upload_enabled = False

    class FailingInventoryClient:
        def read_r14_inventory(self, _settings):  # type: ignore[no-untyped-def]
            raise R14InventorySheetError("R14_INVENTORY_SHEET_READ_FAILED", "讀取 R14 Google Sheet 庫存失敗：denied")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("W01/R14 local transforms must not connect to POS")
        ),
        r14_inventory_client_factory=lambda _config: FailingInventoryClient(),
        run_date=date(2026, 6, 9),
    )

    summary = runner.run()

    assert summary.ok is False
    assert [failure.task_id for failure in summary.failures] == ["W01", "R14"]
    assert summary.failures[0].error_code == "R14_INVENTORY_SHEET_READ_FAILED"
    assert summary.failures[1].error_code == "R14_BLOCKED_BY_W01_FAILED"
    assert "避免用舊庫存模板" in summary.failures[1].message


def test_automation_runner_manual_forced_w01_failure_blocks_r14_on_non_configured_weekday(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.r14_email.enabled = False
    config.r14_transform.template_path = str(ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx")
    config.r14_transform.raw_search_dir = str(ROOT / "tests" / "R14_TEST")
    config.r14_inventory_source.enabled = True
    config.r14_inventory_source.apply_weekday = "Friday"
    for report in config.reports:
        report.enabled = report.id in {"W01", "R14"}
        report.upload_enabled = False

    class FailingInventoryClient:
        def read_r14_inventory(self, _settings):  # type: ignore[no-untyped-def]
            raise R14InventorySheetError("R14_INVENTORY_SHEET_READ_FAILED", "讀取 R14 Google Sheet 庫存失敗：denied")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("W01/R14 local transforms must not connect to POS")
        ),
        r14_inventory_client_factory=lambda _config: FailingInventoryClient(),
        run_date=date(2026, 6, 9),
        run_source="gui_manual",
    )

    summary = runner.run()

    assert summary.ok is False
    assert [failure.task_id for failure in summary.failures] == ["W01", "R14"]
    assert summary.failures[0].error_code == "R14_INVENTORY_SHEET_READ_FAILED"
    assert summary.failures[1].error_code == "R14_BLOCKED_BY_W01_FAILED"


def test_automation_runner_blocks_r14_when_w01_template_is_locked(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.r14_email.enabled = False
    config.r14_transform.template_path = str(ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx")
    config.r14_transform.raw_search_dir = str(ROOT / "tests" / "R14_TEST")
    config.r14_inventory_source.enabled = True
    config.r14_inventory_source.apply_weekday = "Tuesday"
    for report in config.reports:
        report.enabled = report.id in {"W01", "R14"}
        report.upload_enabled = False

    class FakeInventoryClient:
        def read_r14_inventory(self, _settings):  # type: ignore[no-untyped-def]
            return SimpleNamespace(rows_read=1, inventories={"站前4樓": {"6050010": 321}})

    def locked_template(_template_path, _inventories):  # type: ignore[no-untyped-def]
        raise R14TransformError("W01_TEMPLATE_LOCKED", "R14 模板檔目前無法寫入：locked")

    monkeypatch.setattr(automation_runner_module, "sync_r14_template_inventory", locked_template)

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("W01/R14 local transforms must not connect to POS")
        ),
        r14_inventory_client_factory=lambda _config: FakeInventoryClient(),
        run_date=date(2026, 6, 9),
    )

    summary = runner.run()

    assert summary.ok is False
    assert [failure.task_id for failure in summary.failures] == ["W01", "R14"]
    assert summary.failures[0].error_code == "W01_TEMPLATE_LOCKED"
    assert summary.failures[0].diagnostic_path is not None
    assert Path(summary.failures[0].diagnostic_path).name.startswith("automation_local_transform_")
    assert summary.failures[1].error_code == "R14_BLOCKED_BY_W01_FAILED"


def test_automation_runner_reads_r14_inventory_on_default_friday_run_date(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    output = next(item for item in build_dry_run_plan(config, today=date(2026, 6, 12)).outputs if item.task_id == "R14")
    calls = 0

    class FakeInventoryClient:
        def read_r14_inventory(self, settings):  # type: ignore[no-untyped-def]
            nonlocal calls
            calls += 1
            assert settings.apply_weekday == "Friday"
            assert settings.sheet_name == "Summary"
            return SimpleNamespace(inventories={})

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        r14_inventory_client_factory=lambda _config: FakeInventoryClient(),
        run_date=date(2026, 6, 12),
    )

    result = runner._read_r14_inventory_for_output(output)  # type: ignore[attr-defined]

    assert calls == 1
    assert result is not None


def test_automation_runner_skips_r14_inventory_on_non_friday_run_date(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    output = next(item for item in build_dry_run_plan(config, today=date(2026, 6, 9)).outputs if item.task_id == "R14")

    def fail_inventory_client(_config):  # type: ignore[no-untyped-def]
        raise AssertionError("R14 inventory client must not be built when run_date is not Friday")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        r14_inventory_client_factory=fail_inventory_client,
        run_date=date(2026, 6, 9),
    )

    result = runner._read_r14_inventory_for_output(output)  # type: ignore[attr-defined]

    assert result is None


def test_automation_runner_fails_r14_when_inventory_sheet_read_fails(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.r14_email.enabled = False
    config.r14_transform.template_path = str(ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx")
    config.r14_transform.raw_search_dir = str(ROOT / "tests" / "R14_TEST")
    config.r14_inventory_source.enabled = True
    config.r14_inventory_source.apply_weekday = "Tuesday"
    for report in config.reports:
        report.enabled = report.id == "R14"
        if report.id == "R14":
            report.upload_enabled = False

    class FailingInventoryClient:
        def read_r14_inventory(self, _settings):  # type: ignore[no-untyped-def]
            raise R14InventorySheetError("R14_INVENTORY_SHEET_READ_FAILED", "讀取 R14 Google Sheet 庫存失敗：denied")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("R14 offline transform must not connect to POS")
        ),
        r14_inventory_client_factory=lambda _config: FailingInventoryClient(),
        run_date=date(2026, 6, 9),
    )

    summary = runner.run()

    assert summary.ok is False
    assert summary.failures[0].task_id == "R14"
    assert summary.failures[0].error_code == "R14_INVENTORY_SHEET_READ_FAILED"
    assert summary.failures[0].diagnostic_path is not None
    diagnostic_path = Path(summary.failures[0].diagnostic_path)
    assert diagnostic_path.name.startswith("automation_local_transform_")
    diagnostic_text = diagnostic_path.read_text(encoding="utf-8")
    assert "local_transform_result" in diagnostic_text
    assert AUTOMATION_LOGIC_FINGERPRINT in diagnostic_text
    assert "denied" in summary.failures[0].message


def test_automation_runner_sends_r14_completion_email_with_r14_settings(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.email.enabled = True
    config.email.notify_on_failure = True
    config.email.recipients = ["ops-failure@example.com"]
    config.r14_email.enabled = True
    config.r14_email.subject_template = "R14 {date} custom"
    config.r14_email.body = "custom body for {filename}"
    config.r14_transform.template_path = str(
        ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx"
    )
    config.r14_transform.raw_search_dir = str(ROOT / "tests" / "R14_TEST")
    for report in config.reports:
        report.enabled = report.id == "R14"
        if report.id == "R14":
            report.upload_enabled = False
    sent_messages: list[tuple[str, str, list[Path], list[str]]] = []

    class FakeGmailSender:
        def send(self, settings, *, subject, body, attachments=None):  # type: ignore[no-untyped-def]
            sent_messages.append((subject, body, list(attachments or []), list(settings.recipients)))
            return SimpleNamespace(ok=True, message="Gmail API message sent.", gmail_message_id="gmail-r14")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("R14 offline transform must not connect to POS")
        ),
        gmail_sender_factory=lambda _config: FakeGmailSender(),
        run_date=date(2026, 6, 9),
    )

    summary = runner.run()

    assert summary.ok is True
    assert len(sent_messages) == 1
    subject, body, attachments, recipients = sent_messages[0]
    assert subject == "R14 20260608 custom"
    assert body == "custom body for 診所stock status - 2026 demand planning-0608.xlsx"
    assert recipients == [
        "joe.little7208@gmail.com",
        "mickey.chen@mikobeaute.com",
        "rae.hsu@mikobeaute.com",
        "miko_03@mikobeaute.com",
        "bbone_pu@bebetterone.com",
    ]
    assert attachments == [tmp_path / "downloads" / "R14" / "20260609" / "診所stock status - 2026 demand planning-0608.xlsx"]
    assert attachments[0].exists()


def test_automation_runner_r14_friday_email_reports_insufficient_history(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    current = _write_r14_snapshot(
        tmp_path / "downloads" / "R14" / "20260619" / "current.xlsx",
        report_date=date(2026, 6, 18),
        item_actuals={"站前4樓": 130},
    )
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test", run_date=date(2026, 6, 19))

    section = runner._build_r14_friday_analysis_section(current)  # type: ignore[attr-defined]

    assert "日平均量累積成長超過10%:" not in section
    assert "週耗用量暴漲/暴跌超過30%:" in section
    assert 'style="border-collapse:collapse;border:1px solid #444;"' in section
    assert section.count("數據量累積不足，暫無法提供") == 1


def test_automation_runner_r14_friday_email_adds_growth_tables_from_archived_r14_files(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    _write_r14_snapshot(
        tmp_path / "downloads" / "R14" / "20260605" / "two_weeks_ago.xlsx",
        report_date=date(2026, 6, 4),
        item_actuals={"站前4樓": 90},
    )
    _write_r14_snapshot(
        tmp_path / "downloads" / "R14" / "20260612" / "previous.xlsx",
        report_date=date(2026, 6, 11),
        item_actuals={"站前4樓": 100},
    )
    current = _write_r14_snapshot(
        tmp_path / "downloads" / "R14" / "20260619" / "current.xlsx",
        report_date=date(2026, 6, 18),
        item_actuals={"站前4樓": 130},
    )
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test", run_date=date(2026, 6, 19))

    section = runner._build_r14_friday_analysis_section(current)  # type: ignore[attr-defined]

    assert "日平均量累積成長超過10%:" not in section
    assert "<table" in section
    assert "border:1px solid #444" in section
    assert "<th" in section
    assert "凱惠料號" in section
    assert "品名" in section
    assert "站前4樓" in section
    assert "站前11樓" in section
    assert "ITEM001" in section
    assert "高波動品項" in section
    assert "30.0%" not in section
    assert "200.0%" in section
    assert "ITEM002" not in section


def test_automation_runner_r14_friday_email_reads_named_snapshots_from_daily_r14_folders(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    _write_r14_snapshot(
        tmp_path / "downloads" / "R14" / "20260605" / "診所stock status - 2026 demand planning-0604.xlsx",
        report_date=date(2026, 6, 4),
        item_actuals={"站前4樓": 90},
        include_n006=False,
    )
    _write_r14_snapshot(
        tmp_path / "downloads" / "R14" / "20260612" / "診所stock status - 2026 demand planning-0611.xlsx",
        report_date=date(2026, 6, 11),
        item_actuals={"站前4樓": 100},
        include_n006=False,
    )
    current = _write_r14_snapshot(
        tmp_path / "downloads" / "R14" / "20260619" / "診所stock status - 2026 demand planning-0618.xlsx",
        report_date=date(2026, 6, 18),
        item_actuals={"站前4樓": 130},
        include_n006=False,
    )
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test", run_date=date(2026, 6, 19))

    section = runner._build_r14_friday_analysis_section(current)  # type: ignore[attr-defined]

    assert "數據量累積不足，暫無法提供" not in section
    assert "200.0%" in section


def test_automation_runner_r14_friday_email_pivots_items_by_branch_and_colors_rates() -> None:
    def snapshot(*items: SimpleNamespace) -> SimpleNamespace:
        return SimpleNamespace(items=list(items))

    def item(branch: str, item_code: str, item_name: str, actual: float) -> SimpleNamespace:
        return SimpleNamespace(branch=branch, item_code=item_code, item_name=item_name, actual=actual)

    two_weeks_ago = snapshot(
        item("站前4樓", "ITEM_A", "共同品項", 100),
        item("站前11樓", "ITEM_A", "共同品項", 100),
        item("忠孝7樓", "ITEM_B", "忠孝負成長", 100),
    )
    previous = snapshot(
        item("站前4樓", "ITEM_A", "共同品項", 110),
        item("站前11樓", "ITEM_A", "共同品項", 110),
        item("忠孝7樓", "ITEM_B", "忠孝負成長", 110),
    )
    current = snapshot(
        item("站前4樓", "ITEM_A", "共同品項", 130),
        item("站前11樓", "ITEM_A", "共同品項", 150),
        item("忠孝7樓", "ITEM_B", "忠孝負成長", 105),
    )

    rows = AutomationRunner._r14_weekly_usage_growth_rows(current, previous, two_weeks_ago)  # type: ignore[attr-defined]
    section = AutomationRunner._format_r14_anomaly_table("週耗用量暴漲/暴跌超過30%:", rows)  # type: ignore[attr-defined]

    assert [(row.branch, row.item_code) for row in rows] == [
        ("站前4樓", "ITEM_A"),
        ("站前11樓", "ITEM_A"),
        ("忠孝7樓", "ITEM_B"),
    ]
    assert rows[0].growth_rate == pytest.approx(1.0)
    assert rows[1].growth_rate == pytest.approx(3.0)
    assert rows[2].growth_rate == pytest.approx(-1.5)
    assert section.count("ITEM_A") == 1
    assert section.index("ITEM_A") < section.index("ITEM_B")
    assert section.index("站前4樓") < section.index("站前11樓") < section.index("忠孝7樓")
    assert 'color:#c00000;font-weight:600;">100.0%</td>' in section
    assert 'color:#c00000;font-weight:600;">300.0%</td>' in section
    assert 'color:#008000;font-weight:600;">-150.0%</td>' in section


def test_automation_runner_r14_weekly_analysis_treats_missing_legacy_branch_as_zero() -> None:
    def snapshot(*items: SimpleNamespace, missing_branches: tuple[str, ...] = ()) -> SimpleNamespace:
        return SimpleNamespace(items=list(items), missing_branches=missing_branches)

    def item(actual: float) -> SimpleNamespace:
        return SimpleNamespace(branch="忠孝預防醫學3樓", item_code="ITEM_N006", item_name="新增分館品項", actual=actual)

    rows = AutomationRunner._r14_weekly_usage_growth_rows(
        snapshot(item(40)),
        snapshot(item(10)),
        snapshot(missing_branches=("忠孝預防醫學3樓",)),
    )

    assert len(rows) == 1
    assert rows[0].growth_rate == pytest.approx(2.0)


def test_automation_runner_skips_r14_completion_email_when_r14_email_disabled(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    config.app.logs_dir = str(tmp_path / "logs")
    config.app.screenshots_dir = str(tmp_path / "screenshots")
    config.google_drive.upload_enabled = False
    config.email.enabled = True
    config.email.notify_on_failure = True
    config.r14_email.enabled = False
    config.r14_transform.template_path = str(
        ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx"
    )
    config.r14_transform.raw_search_dir = str(ROOT / "tests" / "R14_TEST")
    for report in config.reports:
        report.enabled = report.id == "R14"
        if report.id == "R14":
            report.upload_enabled = False
    send_calls = 0

    class FakeGmailSender:
        def send(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            nonlocal send_calls
            send_calls += 1
            return SimpleNamespace(ok=True, message="Gmail API message sent.", gmail_message_id="gmail-r14")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("R14 offline transform must not connect to POS")
        ),
        gmail_sender_factory=lambda _config: FakeGmailSender(),
        run_date=date(2026, 6, 9),
    )

    summary = runner.run()

    assert summary.ok is True
    assert send_calls == 0


def test_automation_runner_r14_raw_selection_prefers_planned_r13_filename(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    raw_dir = tmp_path / "raw"
    raw_dir.mkdir()
    exact_raw = raw_dir / "診所stock status - 2026 demand planning-0608-rawdata.xls"
    newer_glob_raw = raw_dir / "診所stock status - 2026 demand planning-9999-rawdata.xls"
    fixture_raw = ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0609-rawdata.xls"
    exact_raw.write_bytes(fixture_raw.read_bytes())
    newer_glob_raw.write_bytes(fixture_raw.read_bytes())
    newer_glob_raw.touch()
    config.r14_transform.raw_search_dir = str(raw_dir)
    for report in config.reports:
        report.enabled = report.id in {"R13", "R14"}
    plan = build_dry_run_plan(config, today=date(2026, 6, 9))

    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test", run_date=date(2026, 6, 9))

    assert runner._resolve_r14_raw_path(plan.outputs, expected_end_date=date(2026, 6, 8)) == exact_raw


def test_automation_runner_r14_current_raw_mismatch_does_not_fallback_to_old_raw(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    config = _load_runner_config(tmp_path)
    raw_dir = tmp_path / "raw"
    raw_dir.mkdir()
    current_raw = raw_dir / "current-r13.xls"
    fallback_raw = raw_dir / "診所stock status - 2026 demand planning-0608-rawdata.xls"
    current_raw.write_bytes(b"current-r13")
    fallback_raw.write_bytes(b"old-r13")
    config.r14_transform.raw_search_dir = str(raw_dir)
    for report in config.reports:
        report.enabled = report.id in {"R13", "R14"}
    plan = build_dry_run_plan(config, today=date(2026, 6, 9))

    def fake_parse(path: Path):  # type: ignore[no-untyped-def]
        return SimpleNamespace(end_date=date(2026, 6, 7) if path == current_raw else date(2026, 6, 8))

    monkeypatch.setattr(automation_runner_module, "parse_r13_usage_summary", fake_parse)
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test", run_date=date(2026, 6, 9))

    with pytest.raises(R14TransformError) as exc_info:
        runner._resolve_r14_raw_path(
            plan.outputs,
            expected_end_date=date(2026, 6, 8),
            current_r13_output_path=current_raw,
        )

    assert exc_info.value.error_code == "R14_SOURCE_FILE_MISSING"
    assert "未回退到其他檔案" in exc_info.value.message


def test_automation_runner_w02_r14_resolver_requires_valid_expected_date_and_current_provenance(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    for report in config.reports:
        report.enabled = report.id in {"R14", "W02"}
    run_date = date(2026, 6, 9)
    plan = build_dry_run_plan(config, today=run_date)
    archive_dir = tmp_path / "downloads" / "R14" / "20260609"
    archive_dir.mkdir(parents=True)
    fixture = ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0608.xlsx"
    valid_archive = archive_dir / "診所stock status - 2026 demand planning-0608_001.xlsx"
    wrong_date_archive = archive_dir / "診所stock status - 2026 demand planning-0607.xlsx"
    valid_archive.write_bytes(fixture.read_bytes())
    wrong_date_archive.write_bytes(fixture.read_bytes())
    wrong_date_archive.touch()

    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test", run_date=run_date)

    assert runner._resolve_w02_r14_output_path(plan.outputs) == valid_archive

    invalid_current_output = tmp_path / "current-r14.xlsx"
    invalid_current_output.write_bytes(b"stale-or-invalid")
    runner._latest_r14_output_path = invalid_current_output
    with pytest.raises(R14TransformError) as exc_info:
        runner._resolve_w02_r14_output_path(plan.outputs)

    assert exc_info.value.error_code == "W02_R14_OUTPUT_INVALID"
    assert "未回退到封存資料夾中的其他檔案" in exc_info.value.message


def test_automation_runner_w02_r14_resolver_finds_same_day_archive_without_planned_r14(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    run_date = date(2026, 6, 9)
    for report in config.reports:
        report.enabled = report.id == "W02"

    plan = build_dry_run_plan(
        config,
        today=run_date,
        force_weekly_report_ids={"W02"},
    )
    assert [output.task_id for output in plan.outputs] == ["W02"]

    archive_dir = tmp_path / "downloads" / "R14" / "20260609"
    archive_dir.mkdir(parents=True)
    valid_archive = archive_dir / "診所stock status - 2026 demand planning-0608.xlsx"
    fixture = ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0608.xlsx"
    valid_archive.write_bytes(fixture.read_bytes())
    wrong_date_archive = archive_dir / "診所stock status - 2026 demand planning-0607.xlsx"
    wrong_date_archive.write_bytes(fixture.read_bytes())
    wrong_date_archive.touch()

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        run_date=run_date,
        run_source="gui_manual",
    )

    assert runner._resolve_w02_r14_output_path(plan.outputs) == valid_archive


def test_automation_runner_w02_r14_resolver_does_not_use_old_date_archive_without_planned_r14(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path / "downloads")
    run_date = date(2026, 6, 9)
    for report in config.reports:
        report.enabled = report.id == "W02"
    plan = build_dry_run_plan(config, today=run_date, force_weekly_report_ids={"W02"})

    old_archive = tmp_path / "downloads" / "R14" / "20260608"
    old_archive.mkdir(parents=True)
    (old_archive / "診所stock status - 2026 demand planning-0608.xlsx").write_bytes(
        (ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0608.xlsx").read_bytes()
    )
    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        run_date=run_date,
        run_source="gui_manual",
    )

    with pytest.raises(R14TransformError) as exc_info:
        runner._resolve_w02_r14_output_path(plan.outputs)

    assert exc_info.value.error_code == "W02_R14_OUTPUT_MISSING"


def test_automation_runner_marks_r14_blocked_when_r13_failed_in_same_run(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.app.logs_dir = str(tmp_path / "logs")
    config.r14_email.enabled = False
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id in {"R13", "R14"}
    calls: list[str] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            return SimpleNamespace(
                ok=False,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code="EXPORT_PROGRESS_TIMEOUT",
                message="POS 正在匯出超過 300 秒，尚未出現另存新檔視窗。",
                actions=[],
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window("庫存管理"),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        run_date=date(2026, 6, 16),
    )

    summary = runner.run()

    assert calls == ["R13"]
    assert summary.ok is False
    assert summary.completed == 0
    assert summary.details is not None
    assert "任務：R13" in summary.details
    assert "任務：R14" in summary.details
    assert "R14_BLOCKED_BY_R13_FAILED" in summary.details
    assert "R13 raw data 未產生" in summary.details


def test_automation_runner_runs_r14_when_r13_local_file_exists_but_drive_upload_fails(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    downloads_dir = tmp_path / "downloads"
    config.app.downloads_dir = str(downloads_dir)
    config.app.logs_dir = str(tmp_path / "logs")
    config.r14_email.enabled = False
    config.google_drive.upload_enabled = True
    config.r14_transform.template_path = str(ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx")
    config.r14_transform.raw_search_dir = str(downloads_dir)
    for report in config.reports:
        report.enabled = report.id in {"R13", "R14"}
        report.upload_enabled = report.id in {"R13", "R14"}
    calls: list[str] = []
    uploads: list[str] = []
    upload_paths: list[Path] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            output_path = downloads_dir / output.output_filename
            output_path.parent.mkdir(parents=True, exist_ok=True)
            fixture = ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0609-rawdata.xls"
            output_path.write_bytes(fixture.read_bytes())
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
                actions=[],
            )

    class FakeUploader:
        def upload(self, file_path: Path, folder_id: str, name: str) -> DriveUploadResult:
            uploads.append(name)
            upload_paths.append(file_path)
            if name.endswith("-rawdata.xls"):
                return DriveUploadResult(
                    success=False,
                    folder_id=folder_id,
                    uploaded_name=name,
                    size=file_path.stat().st_size,
                    error_code="GOOGLE_OAUTH_REAUTH_REQUIRED",
                    message="Google OAuth 授權範圍已失效，請重新連接 Google Drive。",
                )
            return DriveUploadResult(
                success=True,
                drive_file_id="drive-r14",
                folder_id=folder_id,
                uploaded_name=name,
                size=file_path.stat().st_size,
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window("庫存管理"),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        drive_uploader_factory=lambda _config: FakeUploader(),
        run_date=date(2026, 6, 9),
    )

    summary = runner.run()

    assert calls == ["R13"]
    assert summary.ok is False
    assert summary.completed == 1
    assert [failure.task_id for failure in summary.failures] == ["R13"]
    assert summary.failures[0].error_code == "GOOGLE_OAUTH_REAUTH_REQUIRED"
    assert any(name.endswith("-rawdata.xls") for name in uploads)
    assert any(name.endswith(".xlsx") for name in uploads)
    assert any(path.suffix == ".xlsx" and path.exists() for path in upload_paths)


def test_automation_runner_google_clients_use_service_specific_oauth_scopes(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
    )

    drive_uploader = runner._build_google_drive_uploader(config)  # type: ignore[attr-defined]
    gmail_sender = runner._build_gmail_sender(config)  # type: ignore[attr-defined]
    sheets_client = runner._build_r14_inventory_client(config)  # type: ignore[attr-defined]

    assert tuple(drive_uploader.oauth.scopes) == GOOGLE_DRIVE_SCOPES
    assert tuple(gmail_sender.oauth.scopes) == GOOGLE_GMAIL_SCOPES
    assert tuple(sheets_client.oauth.scopes) == GOOGLE_SHEETS_SCOPES
    assert drive_uploader.oauth.profile == GOOGLE_DRIVE_PROFILE
    assert gmail_sender.oauth.profile == GOOGLE_GMAIL_PROFILE
    assert sheets_client.oauth.profile == GOOGLE_SHEETS_PROFILE


def test_automation_runner_blocks_r14_after_r13_download_failure_even_if_stale_raw_exists(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    downloads_dir = tmp_path / "downloads"
    config.app.downloads_dir = str(downloads_dir)
    config.app.logs_dir = str(tmp_path / "logs")
    config.r14_email.enabled = False
    config.r14_transform.template_path = str(ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx")
    config.r14_transform.raw_search_dir = str(downloads_dir)
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id in {"R13", "R14"}
    stale_raw = downloads_dir / "診所stock status - 2026 demand planning-0609-rawdata.xls"
    stale_raw.parent.mkdir(parents=True, exist_ok=True)
    stale_raw.write_bytes((ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0609-rawdata.xls").read_bytes())
    calls: list[str] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            return SimpleNamespace(
                ok=False,
                task_id=output.task_id,
                output_path=downloads_dir / output.output_filename,
                error_code="EXPORT_PROGRESS_TIMEOUT",
                message="POS 正在匯出超過 300 秒，尚未出現另存新檔視窗。",
                actions=[],
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window("庫存管理"),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        run_date=date(2026, 6, 9),
    )

    summary = runner.run()

    assert calls == ["R13"]
    assert summary.ok is False
    assert [failure.task_id for failure in summary.failures] == ["R13", "R14"]
    assert summary.failures[1].error_code == "R14_BLOCKED_BY_R13_FAILED"
    r14_dir = downloads_dir / "R14"
    assert not (r14_dir.exists() and any(r14_dir.rglob("*.xlsx")))


def test_automation_runner_outputs_zero_usage_r14_after_r13_no_data_without_using_stale_raw(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    downloads_dir = tmp_path / "downloads"
    config.app.downloads_dir = str(downloads_dir)
    config.app.logs_dir = str(tmp_path / "logs")
    config.r14_email.enabled = False
    config.r14_transform.template_path = str(ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx")
    config.r14_transform.raw_search_dir = str(downloads_dir)
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id in {"R13", "R14"}
    stale_raw = downloads_dir / "診所stock status - 2026 demand planning-0609-rawdata.xls"
    stale_raw.parent.mkdir(parents=True, exist_ok=True)
    stale_raw.write_bytes((ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0609-rawdata.xls").read_bytes())
    calls: list[str] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            return SimpleNamespace(
                ok=False,
                task_id=output.task_id,
                output_path=downloads_dir / output.output_filename,
                error_code="NO_REPORT_DATA",
                message="POS 顯示目前並無符合的資料；已按下確定並停止本輪 R13 依賴。",
                actions=["dismiss_warning:目前並無符合的資料"],
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window("庫存管理"),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        run_date=date(2026, 6, 9),
    )

    summary = runner.run()

    assert calls == ["R13"]
    assert summary.ok is False
    assert summary.completed == 1
    assert [failure.task_id for failure in summary.failures] == ["R13"]
    assert summary.failures[0].error_code == "NO_REPORT_DATA"
    assert "已記錄可驗證的零領用狀態" in summary.failures[0].message
    r14_dir = downloads_dir / "R14"
    assert any(r14_dir.rglob("診所stock status - 2026 demand planning-0608.xlsx"))
    marker_path = tmp_path / "state" / "20260609" / "r13_no_report_data.json"
    assert marker_path.exists()
    r14_logs = list((tmp_path / "logs" / "20260609").glob("automation_local_transform_*_R14.jsonl"))
    assert len(r14_logs) == 1
    log_text = r14_logs[0].read_text(encoding="utf-8")
    assert "r14_no_data_marker" in log_text
    assert "r14_imported_rows:0" in log_text


def test_automation_runner_r14_only_child_consumes_exact_prior_r13_no_data_marker(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    downloads_dir = tmp_path / "downloads"
    config.app.downloads_dir = str(downloads_dir)
    config.app.logs_dir = str(tmp_path / "logs")
    config.r14_email.enabled = False
    config.r14_transform.template_path = str(
        ROOT / "tests" / "R14_TEST" / "診所stock status - 2026 demand planning-0531.xlsx"
    )
    config.r14_transform.raw_search_dir = str(downloads_dir)
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R14"
    marker_path = tmp_path / "state" / "20260609" / "r13_no_report_data.json"
    write_r13_no_data_marker(
        marker_path,
        app_version="test",
        run_date=date(2026, 6, 9),
        start_date="2026/06/01",
        end_date="2026/06/08",
        output_filename="診所stock status - 2026 demand planning-0608-rawdata.xls",
        message="POS 明確回覆無領用資料。",
    )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("R14 marker transform must not connect to POS")
        ),
        run_date=date(2026, 6, 9),
    )

    summary = runner.run()

    assert summary.ok is True
    assert summary.completed == 1
    assert any((downloads_dir / "R14").rglob("診所stock status - 2026 demand planning-0608.xlsx"))


def test_automation_runner_closes_pos_after_all_outputs_when_enabled(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.pos.close_after_run = True
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id in {"R01", "R02", "R03"}
    pos_window = FakeClosablePosWindow()
    progress_messages: list[str] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
                actions=[],
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: pos_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run(on_progress=lambda event: progress_messages.append(event.message))

    assert summary.ok is True
    assert pos_window.closed is True
    assert any("正在關閉 SPA-POS" in message for message in progress_messages)


def test_automation_runner_reconnects_after_success_with_stale_post_save_actions(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id in {"R01", "R02"}
    connect_calls: list[str] = []
    task_calls: list[str] = []
    runtime_sources: list[str] = []

    def connect_pos(**_kwargs):  # type: ignore[no-untyped-def]
        connect_calls.append("connect")
        return _ready_pos_window()

    class FakeAutomator:
        def __init__(self, *_args, **kwargs):  # type: ignore[no-untyped-def]
            runtime_sources.append(kwargs["runtime_metadata"]["run_source"])

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            task_calls.append(output.task_id)
            actions = []
            if output.task_id == "R01":
                actions = [
                    "skip_close_report_viewer:post_save_success:"
                    "(-2147220991, '事件無法啟動任何訂閱者', (None, None, None, 0, None))"
                ]
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
                actions=actions,
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        run_source="windows_task_scheduler",
    )

    summary = runner.run()

    assert summary.ok is True
    assert task_calls == ["R01", "R02"]
    assert len(connect_calls) >= 2
    assert runtime_sources == ["windows_task_scheduler", "windows_task_scheduler"]


def test_automation_runner_reconnects_through_login_after_stale_post_save_actions(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id in {"R01", "R02"}
    ready_window = _ready_pos_window()
    login_window = FakePosControl("帳號登入", "Window")
    account_edit = FakePosControl("帳號", "Edit")
    secret_edit = FakePosControl("密碼", "Edit")
    current_window = ready_window
    task_calls: list[str] = []
    automator_windows: list[object] = []

    def complete_login() -> None:
        nonlocal current_window
        current_window = ready_window

    login_window.children_controls = [
        secret_edit,
        account_edit,
        FakePosControl("登入", "Button", on_click=complete_login),
    ]

    def connect_pos(**_kwargs):  # type: ignore[no-untyped-def]
        return current_window

    class FakeAutomator:
        def __init__(self, window, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            automator_windows.append(window)

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            nonlocal current_window
            task_calls.append(output.task_id)
            actions = []
            if output.task_id == "R01":
                actions = [
                    "skip_close_report_viewer:post_save_success:"
                    "(-2147220991, '事件無法啟動任何訂閱者', (None, None, None, 0, None))"
                ]
                current_window = login_window
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
                actions=actions,
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_login_secret_provider=lambda: "1234",
    )

    summary = runner.run()

    assert summary.ok is True
    assert task_calls == ["R01", "R02"]
    assert automator_windows == [ready_window, ready_window]
    assert account_edit.text_value == "A0042"
    assert secret_edit.text_value == "1234"


def test_automation_runner_closes_each_r06_branch_window(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R06"
    close_flags: list[tuple[str, bool]] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            close_flags.append((output.branch_code, close_after_success))
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
                actions=[],
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run()

    assert summary.ok is True
    assert [branch_code for branch_code, _close in close_flags] == ["N001", "N002", "N003", "N004", "N005", "N006"]
    assert all(close_after_success is True for _branch_code, close_after_success in close_flags)


def test_automation_runner_treats_any_no_data_as_failure_and_sends_notification(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.email.enabled = True
    config.email.notify_on_failure = True
    config.email.recipients = ["ops@example.com"]
    for report in config.reports:
        report.enabled = report.id == "R06"
    progress_messages: list[str] = []
    sent_messages: list[tuple[str, str, list[str]]] = []
    run_state_store = RunStateStore(tmp_path / "state" / "run_state_latest.json")

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            if output.branch_code == "N006":
                return SimpleNamespace(
                    ok=False,
                    task_id=output.task_id,
                    output_path=tmp_path / output.output_filename,
                    error_code="NO_REPORT_DATA",
                    message="POS 顯示目前並無符合的療程殘值資料；已按下確定並跳過此輸出。",
                    actions=["dismiss_warning:目前並無符合的療程殘值資料"],
                )
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
                actions=[],
            )

    class FakeGmailSender:
        def send(self, settings, *, subject, body):  # type: ignore[no-untyped-def]
            sent_messages.append((subject, body, list(settings.recipients)))
            return SimpleNamespace(ok=True, message="Gmail API message sent.", gmail_message_id="gmail-1")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        run_state_store=run_state_store,
        gmail_sender_factory=lambda _config: FakeGmailSender(),
    )

    summary = runner.run(on_progress=lambda event: progress_messages.append(event.message))
    snapshot = RunStateStore(tmp_path / "state" / "run_state_latest.json").load()

    assert summary.ok is False
    assert summary.completed == 5
    assert summary.skipped == 0
    assert len(summary.failures) == 1
    assert summary.failures[0].task_id == "R06"
    assert summary.failures[0].error_code == "NO_REPORT_DATA"
    assert "R06 失敗：NO_REPORT_DATA" in progress_messages
    assert len(sent_messages) == 1
    subject, body, recipients = sent_messages[0]
    assert subject == "POSReportBot 報表自動化失敗通知"
    assert "任務：R06" in body
    assert "錯誤代碼：NO_REPORT_DATA" in body
    assert recipients == ["ops@example.com"]
    assert snapshot is not None
    assert snapshot.status == "partial_failed"
    n006_state = next(output for output in snapshot.outputs.values() if output.branch_code == "N006")
    assert n006_state.status == "failed"
    assert n006_state.error_code == "NO_REPORT_DATA"
    assert n006_state.local_file_path is None


def test_automation_runner_uses_gmail_api_for_failure_notification(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.email.enabled = True
    config.email.notify_on_failure = True
    config.email.recipients = ["ops@example.com"]
    config.r14_email.recipients = ["r14-report@example.com"]
    for report in config.reports:
        report.enabled = report.id == "R01"
    sent_messages: list[tuple[str, str, list[str]]] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            raise ReportAutomationError("CONTROL_NOT_CLICKABLE", "控制項無法點擊：檢視報表")

    class FakeGmailSender:
        def send(self, settings, *, subject, body):  # type: ignore[no-untyped-def]
            sent_messages.append((subject, body, list(settings.recipients)))
            return SimpleNamespace(ok=True, message="Gmail API message sent.", gmail_message_id="gmail-1")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        gmail_sender_factory=lambda _config: FakeGmailSender(),
    )

    summary = runner.run()

    assert summary.ok is False
    assert sent_messages
    subject, body, recipients = sent_messages[0]
    assert subject == "POSReportBot 報表自動化失敗通知"
    assert "任務：R01" in body
    assert recipients == ["ops@example.com"]
    assert "Gmail API message sent." in summary.message


def test_automation_runner_sends_configured_success_summary_with_output_evidence(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.email.enabled = True
    config.email.notify_on_success_summary = True
    config.email.recipients = ["ops@example.com"]
    for report in config.reports:
        report.enabled = report.id == "R01"
    sent_messages: list[tuple[str, str, list[str]]] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
                actions=[],
            )

    class FakeGmailSender:
        def send(self, settings, *, subject, body):  # type: ignore[no-untyped-def]
            sent_messages.append((subject, body, list(settings.recipients)))
            return SimpleNamespace(ok=True, message="Gmail API message sent.", gmail_message_id="gmail-1")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        run_date=date(2026, 9, 3),
        run_source="windows_task_scheduler",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        gmail_sender_factory=lambda _config: FakeGmailSender(),
    )

    summary = runner.run()

    assert summary.ok is True
    assert len(sent_messages) == 1
    subject, body, recipients = sent_messages[0]
    assert subject == "POSReportBot 每日自動化完成摘要"
    assert "執行日期：2026-09-03" in body
    assert "R01" in body
    assert "課程服務明細表" in body
    assert recipients == ["ops@example.com"]


def test_automation_runner_does_not_hide_success_summary_delivery_failure(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.email.enabled = True
    config.email.notify_on_failure = False
    config.email.notify_on_success_summary = True
    config.email.recipients = ["ops@example.com"]
    for report in config.reports:
        report.enabled = report.id == "R01"
    run_state_store = RunStateStore(tmp_path / "state" / "run_state_latest.json")

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
                actions=[],
            )

    class FailedGmailSender:
        @staticmethod
        def send(settings, *, subject, body):  # type: ignore[no-untyped-def]
            del settings, subject, body
            return SimpleNamespace(ok=False, message="Gmail API rejected the message")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        run_source="windows_task_scheduler",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        run_state_store=run_state_store,
        gmail_sender_factory=lambda _config: FailedGmailSender(),
    )

    summary = runner.run()
    snapshot = run_state_store.load()

    assert summary.ok is False
    assert summary.error_code == "SUCCESS_SUMMARY_NOTIFICATION_FAILED"
    assert "Gmail API rejected the message" in summary.message
    assert snapshot is not None
    assert snapshot.status == "partial_failed"


def test_automation_runner_finishes_run_state_on_unexpected_task_error(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R01"

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            raise RuntimeError("unexpected stall")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run()
    state = RunStateStore.default_for_config(config).load()

    assert summary.ok is False
    assert summary.error_code == "REPORT_RUN_FAILED"
    assert state is not None
    assert state.status == "failed"
    r01_state = next(iter(state.outputs.values()))
    assert r01_state.status == "failed"
    assert r01_state.error_code == "UNEXPECTED_RUNNER_ERROR"


def test_automation_runner_automatically_bundles_failure_evidence(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R01"

    unrelated_path = Path(config.app.logs_dir) / "unrelated_other_execution.json"
    sensitive_path = Path(config.app.logs_dir) / "oauth_token.json"
    production_token_path = Path(config.app.state_dir) / "google_drive_user_token.bin"

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            diagnostic_path = Path(config.app.logs_dir) / "automatic_failure_source.json"
            diagnostic_path.parent.mkdir(parents=True, exist_ok=True)
            unrelated_path.write_text('{"execution_id": "other"}', encoding="utf-8")
            sensitive_path.write_text('{"access_token": "must-not-escape"}', encoding="utf-8")
            production_token_path.parent.mkdir(parents=True, exist_ok=True)
            production_token_path.write_bytes(b"dpapi-token-must-not-escape")
            diagnostic_path.write_text(
                json.dumps(
                    {
                        "failure": True,
                        "sensitive_reference": str(sensitive_path),
                        "production_token_reference": str(production_token_path),
                        "password": "alpha beta gamma",
                    }
                ),
                encoding="utf-8",
            )
            return SimpleNamespace(
                ok=False,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code="TEST_REPORT_FAILED",
                message="simulated report failure",
                actions=["simulated_failure"],
                diagnostic_path=diagnostic_path,
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run()

    assert summary.ok is False
    assert summary.evidence_bundle_path is not None
    bundle_path = Path(summary.evidence_bundle_path)
    assert bundle_path.exists()
    assert bundle_path.suffix == ".zip"
    with zipfile.ZipFile(bundle_path) as archive:
        names = archive.namelist()
        assert "manifest.json" in names
        assert any(name.endswith("automatic_failure_source.json") for name in names)
        assert any("automation_runtime_" in name for name in names)
        assert any("run_state_" in name for name in names)
        assert not any("unrelated_other_execution" in name for name in names)
        assert not any("oauth_token" in name for name in names)
        assert not any("google_drive_user_token" in name for name in names)
        diagnostic_member = next(name for name in names if name.endswith("automatic_failure_source.json"))
        diagnostic_text = archive.read(diagnostic_member).decode("utf-8")
        assert "alpha beta gamma" not in diagnostic_text
        assert '"password": "<redacted>"' in diagnostic_text
        manifest = json.loads(archive.read("manifest.json"))
        denied_items = [
            item
            for item in manifest["omitted_files"]
            if item["reason"] == "sensitive_filename_denied"
        ]
        assert len(denied_items) >= 2
        assert all(item["path"] == "<redacted-sensitive-path>" for item in denied_items)
        assert all("path_sha256" not in item for item in denied_items)
        manifest_text = json.dumps(manifest, ensure_ascii=False)
        assert "oauth_token.json" not in manifest_text
        assert "google_drive_user_token.bin" not in manifest_text


def test_failure_evidence_total_limit_uses_redacted_archive_size(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R01"
    max_total_size = 8 * 1024
    monkeypatch.setattr(
        automation_runner_module,
        "_FAILURE_EVIDENCE_MAX_TOTAL_SIZE",
        max_total_size,
    )
    diagnostic_path = Path(config.app.logs_dir) / "expanding_failure.json"
    raw_diagnostic = json.dumps(
        {"values": [0] * 1500},
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode("utf-8")
    assert len(raw_diagnostic) < max_total_size
    assert (
        len(automation_runner_module._redacted_evidence_bytes(diagnostic_path, raw_diagnostic))
        > max_total_size
    )

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            diagnostic_path.parent.mkdir(parents=True, exist_ok=True)
            diagnostic_path.write_bytes(raw_diagnostic)
            return SimpleNamespace(
                ok=False,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code="TEST_REPORT_FAILED",
                message="simulated report failure",
                actions=["simulated_failure"],
                diagnostic_path=diagnostic_path,
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run()

    assert summary.evidence_bundle_path is not None
    with zipfile.ZipFile(summary.evidence_bundle_path) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        assert sum(item["size"] for item in manifest["included_files"]) <= max_total_size
        assert not any(name.endswith("expanding_failure.json") for name in archive.namelist())
        assert any(
            item["reason"] == "size_limit_after_redaction"
            for item in manifest["omitted_files"]
        )


def test_failure_manifest_redacts_sensitive_summary_path_fields() -> None:
    manifest = {
        "summary": {
            "failures": [
                {
                    "diagnostic_path": r"C:\ProgramData\POSReportBot\logs\client_secret.txt",
                    "evidence_paths": [
                        r"C:\ProgramData\POSReportBot\state\oauth_token.json",
                        r"C:\ProgramData\POSReportBot\logs\safe_diagnostic.json",
                    ],
                }
            ]
        }
    }

    redacted = automation_runner_module._redact_evidence_object(manifest)
    failure = redacted["summary"]["failures"][0]

    assert failure["diagnostic_path"] == "<redacted-sensitive-path>"
    assert failure["evidence_paths"] == [
        "<redacted-sensitive-path>",
        r"C:\ProgramData\POSReportBot\logs\safe_diagnostic.json",
    ]
    text = json.dumps(redacted, ensure_ascii=False)
    assert "client_secret.txt" not in text
    assert "oauth_token.json" not in text


def test_automation_runner_bundles_planner_initialization_failure(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    monkeypatch.setattr(
        automation_runner_module,
        "build_dry_run_plan",
        lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("planner exploded")),
    )
    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
    )

    summary = runner.run()

    assert summary.ok is False
    assert summary.error_code == "RUNNER_INITIALIZATION_FAILED"
    assert summary.evidence_bundle_path is not None
    with zipfile.ZipFile(summary.evidence_bundle_path) as archive:
        assert any("automation_runner_outer_failure_" in name for name in archive.namelist())


def test_failure_bundle_includes_explicit_w02_plan_but_excludes_stale_primary_state(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        run_date=date(2026, 9, 1),
    )
    plan_path = Path(config.app.downloads_dir) / "W02" / "20260901" / "w02_order_plan_20260901.json"
    plan_path.parent.mkdir(parents=True, exist_ok=True)
    plan_path.write_text('{"forms": [{"branch": "站前4樓"}]}', encoding="utf-8")
    diagnostic_path = Path(config.app.logs_dir) / "automation_w02_pos_order_test.json"
    diagnostic_path.parent.mkdir(parents=True, exist_ok=True)
    diagnostic_path.write_text('{"error_code": "W02_TEST"}', encoding="utf-8")
    primary_path = Path(config.app.state_dir) / "20260901" / "run_state_latest.json"
    recovery_path = primary_path.with_name("run_state_recovery_current.json")
    primary_path.parent.mkdir(parents=True, exist_ok=True)
    primary_path.write_text('{"execution_id": "stale"}', encoding="utf-8")
    recovery_path.write_text('{"execution_id": "current"}', encoding="utf-8")
    journal_path = Path(config.app.logs_dir) / "automation_runtime_current.jsonl"
    journal_path.write_text(
        json.dumps(
            {
                "execution_id": "current",
                "run_state_path": str(recovery_path),
                "run_state_primary_path": str(primary_path),
            }
        )
        + "\n",
        encoding="utf-8",
    )
    store = RunStateStore(primary_path)
    store.path = recovery_path
    store.recovery_reason = "LATEST_STATE_REPLACE_DENIED"
    runner._active_run_state_store = store
    runner._runtime_phase_journal = SimpleNamespace(
        path=journal_path,
        execution_id="current",
    )
    summary = AutomationRunSummary(
        ok=False,
        completed=0,
        total=1,
        error_code="REPORT_RUN_FAILED",
        message="W02 failed password=manifest-secret-value",
        failures=(
            ReportRunFailure(
                task_id="W02",
                output_filename="",
                error_code="W02_TEST",
                message="failed password=manifest-secret-value",
                diagnostic_path=str(diagnostic_path),
                evidence_paths=(str(plan_path),),
            ),
        ),
    )

    bundled = runner._attach_failure_evidence(summary)

    assert bundled.evidence_bundle_path is not None
    with zipfile.ZipFile(bundled.evidence_bundle_path) as archive:
        names = archive.namelist()
        assert any(name.endswith("w02_order_plan_20260901.json") for name in names)
        assert any(name.endswith("run_state_recovery_current.json") for name in names)
        assert not any(name.endswith("run_state_latest.json") for name in names)
        manifest = json.loads(archive.read("manifest.json"))
        assert "manifest-secret-value" not in json.dumps(manifest, ensure_ascii=False)
        assert any(
            item["reason"] == "historical_locked_primary_not_current_execution"
            for item in manifest["omitted_files"]
        )


def test_automation_runner_continues_after_unexpected_task_error(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id in {"R09", "R10", "R11"}
    calls: list[str] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            if output.task_id == "R09":
                raise RuntimeError("(-2146233083, None, (None, None, None, 0, None))")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
                actions=[],
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run()
    state = RunStateStore.default_for_config(config).load()

    assert calls == ["R09", "R10", "R11"]
    assert summary.ok is False
    assert summary.completed == 2
    assert summary.error_code == "PARTIAL_REPORT_RUN_FAILED"
    assert len(summary.failures) == 1
    assert summary.failures[0].task_id == "R09"
    assert summary.failures[0].error_code == "UNEXPECTED_RUNNER_ERROR"
    assert state is not None
    r09_state = next(output for output in state.outputs.values() if output.task_id == "R09")
    r10_state = next(output for output in state.outputs.values() if output.task_id == "R10")
    r11_state = next(output for output in state.outputs.values() if output.task_id == "R11")
    assert r09_state.status == "failed"
    assert r10_state.status == "completed"
    assert r11_state.status == "completed"


def test_automation_runner_stops_remaining_pos_tasks_when_unexpected_error_reconnect_fails(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id in {"R09", "R10", "R11"}
    calls: list[str] = []
    connection_count = 0
    stale_window = False

    def connect_pos(**_kwargs):  # type: ignore[no-untyped-def]
        nonlocal connection_count
        connection_count += 1
        if connection_count == 1:
            return _ready_pos_window()
        raise UiProbeError("POS error dialog remained open")

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            nonlocal stale_window
            calls.append(output.task_id)
            stale_window = True
            raise RuntimeError("unexpected POS modal")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )
    original_prepare = runner._prepare_pos_window_for_automation

    def prepare_pos(config_value, window):  # type: ignore[no-untyped-def]
        if stale_window:
            raise RuntimeError("stale POS modal")
        return original_prepare(config_value, window)

    runner._prepare_pos_window_for_automation = prepare_pos  # type: ignore[method-assign]

    summary = runner.run()
    state = RunStateStore.default_for_config(config).load()

    assert calls == ["R09"]
    assert summary.ok is False
    assert summary.completed == 0
    assert len(summary.failures) == 3
    assert state is not None
    states = {output.task_id: output for output in state.outputs.values()}
    assert states["R09"].error_code == "UNEXPECTED_RUNNER_ERROR"
    assert states["R10"].error_code == "POS_CONNECTION_FAILED"
    assert states["R11"].error_code == "POS_CONNECTION_FAILED"


def test_automation_runner_returns_no_enabled_reports_without_pos_connection(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.pos.executable_path = ""
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = False
    connected = False

    def connect_pos(**_kwargs):  # type: ignore[no-untyped-def]
        nonlocal connected
        connected = True
        return _ready_pos_window()

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos,
    )

    summary = runner.run()

    assert summary.ok is False
    assert summary.error_code == "NO_ENABLED_REPORTS"
    assert connected is False


def test_automation_runner_marks_all_outputs_failed_when_pos_preparation_fails(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.pos.executable_path = ""
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id in {"R01", "R02", "R03"}
    connect_calls = 0

    def connect_pos(**_kwargs):  # type: ignore[no-untyped-def]
        nonlocal connect_calls
        connect_calls += 1
        raise UiProbeError("找不到 SPA-POS 視窗")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos,
    )

    summary = runner.run()
    state = RunStateStore.default_for_config(config).load()

    assert summary.ok is False
    assert summary.error_code == "POS_CONNECTION_FAILED"
    assert state is not None
    assert state.status == "failed"
    assert len(state.outputs) == 3
    assert {output.status for output in state.outputs.values()} == {"failed"}
    assert {output.error_code for output in state.outputs.values()} == {"POS_CONNECTION_FAILED"}
    assert connect_calls == 2


def test_automation_runner_recovers_pos_and_retries_current_task(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.pos_recovery.enabled = True
    config.pos_recovery.max_restarts_per_run = 1
    for report in config.reports:
        report.enabled = report.id in {"R09", "R10"}
    calls: list[str] = []
    recoveries: list[str] = []
    progress_events: list[str] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            if output.task_id == "R09" and calls.count("R09") == 1:
                raise ReportAutomationError("POS_NOT_RESPONDING", "POS 無回應")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
            )

    def recover_pos(_config, _on_progress):  # type: ignore[no-untyped-def]
        recoveries.append("recover")
        return _ready_pos_window()

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_recovery_func=recover_pos,
    )

    summary = runner.run(on_progress=lambda event: progress_events.append(event.event))

    assert calls == ["R09", "R09", "R10"]
    assert recoveries == ["recover"]
    assert "recovery" in progress_events
    assert summary.ok is True
    assert summary.completed == 2
    assert summary.failures == ()


def test_automation_runner_restarts_pos_for_export_menu_not_opened_and_retries_current_task(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.pos_recovery.enabled = True
    config.pos_recovery.max_restarts_per_run = 1
    for report in config.reports:
        report.enabled = report.id in {"R01", "R02"}
    calls: list[str] = []
    recoveries: list[str] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            if output.task_id == "R01" and len(calls) == 1:
                raise ReportAutomationError(
                    "EXPORT_MENU_NOT_OPENED",
                    "已點擊報表工具列的「匯出」，但未確認匯出格式選單。",
                )
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
            )

    def recover_pos(_config, _on_progress):  # type: ignore[no-untyped-def]
        recoveries.append("recover")
        return _ready_pos_window()

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_recovery_func=recover_pos,
    )

    summary = runner.run()

    assert calls == ["R01", "R01", "R02"]
    assert recoveries == ["recover"]
    assert summary.ok is True
    assert summary.completed == 2
    assert summary.failures == ()


@pytest.mark.parametrize(
    "error_code",
    ["EXPORT_MENU_NOT_OPENED", "VIEW_REPORT_NOT_TRIGGERED", "SAVE_AS_DIALOG_NOT_FOUND"],
)
def test_automation_runner_restarts_pos_for_recoverable_failed_result_and_retries_current_task(
    error_code: str,
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.pos_recovery.enabled = True
    config.pos_recovery.max_restarts_per_run = 1
    for report in config.reports:
        report.enabled = report.id in {"R11", "R12"}
    calls: list[str] = []
    recoveries: list[str] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            if output.task_id == "R11" and calls.count("R11") == 1:
                return SimpleNamespace(
                    ok=False,
                    task_id=output.task_id,
                    output_path=tmp_path / output.output_filename,
                    error_code=error_code,
                    message="POS 報表預覽或匯出工作流未完成",
                    actions=["strategy:匯出:R11使用有界原生popup探測避免UIA阻塞"],
                    diagnostic_path=None,
                )
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
                actions=[],
                diagnostic_path=None,
            )

    def recover_pos(_config, _on_progress):  # type: ignore[no-untyped-def]
        recoveries.append("recover")
        return _ready_pos_window()

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_recovery_func=recover_pos,
    )

    summary = runner.run()

    assert calls == ["R11", "R11", "R12"]
    assert recoveries == ["recover"]
    assert summary.ok is True
    assert summary.completed == 2
    assert summary.failures == ()


@pytest.mark.parametrize("error_code", ["EXPORT_BUTTON_NOT_READY", "VIEW_REPORT_NOT_TRIGGERED"])
def test_automation_runner_recovers_preview_failure_and_retries_current_task(
    error_code: str,
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.pos_recovery.enabled = True
    config.pos_recovery.max_restarts_per_run = 1
    for report in config.reports:
        report.enabled = report.id == "R01"
    calls: list[str] = []
    recoveries: list[str] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            if len(calls) == 1:
                raise ReportAutomationError(
                    error_code,
                    "報表已按下「檢視報表」，但工具列的「匯出」沒有啟用。",
                )
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
            )

    def recover_pos(_config, _on_progress):  # type: ignore[no-untyped-def]
        recoveries.append("recover")
        return _ready_pos_window()

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_recovery_func=recover_pos,
    )

    summary = runner.run()

    assert calls == ["R01", "R01"]
    assert recoveries == ["recover"]
    assert summary.ok is True
    assert summary.completed == 1
    assert summary.failures == ()


def test_automation_runner_recovers_export_format_not_activated_and_retries_current_task(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.pos_recovery.enabled = True
    config.pos_recovery.max_restarts_per_run = 1
    for report in config.reports:
        report.enabled = report.id == "R01"
    calls: list[str] = []
    recoveries: list[str] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            if len(calls) == 1:
                raise ReportAutomationError(
                    "EXPORT_FORMAT_NOT_ACTIVATED",
                    "已看到 Excel 匯出選項，但點擊後未出現另存新檔。",
                )
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
            )

    def recover_pos(_config, _on_progress):  # type: ignore[no-untyped-def]
        recoveries.append("recover")
        return _ready_pos_window()

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_recovery_func=recover_pos,
    )

    summary = runner.run()

    assert calls == ["R01", "R01"]
    assert recoveries == ["recover"]
    assert summary.ok is True
    assert summary.completed == 1
    assert summary.failures == ()


def test_automation_runner_recovers_export_format_not_found_and_retries_current_task(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.pos_recovery.enabled = True
    config.pos_recovery.max_restarts_per_run = 1
    for report in config.reports:
        report.enabled = report.id == "R05"
    calls: list[str] = []
    recoveries: list[str] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            if len(calls) == 1:
                raise ReportAutomationError(
                    "EXPORT_FORMAT_NOT_FOUND",
                    "已點擊報表工具列的匯出按鈕，但找不到 Excel 匯出選項。",
                )
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
            )

    def recover_pos(_config, _on_progress):  # type: ignore[no-untyped-def]
        recoveries.append("recover")
        return _ready_pos_window()

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_recovery_func=recover_pos,
    )

    summary = runner.run()

    assert calls == ["R05", "R05"]
    assert recoveries == ["recover"]
    assert summary.ok is True
    assert summary.completed == 1
    assert summary.failures == ()


def test_automation_runner_recovers_export_menu_open_failed_and_retries_current_task(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.pos_recovery.enabled = True
    config.pos_recovery.max_restarts_per_run = 1
    for report in config.reports:
        report.enabled = report.id == "R01"
    calls: list[str] = []
    recoveries: list[str] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            if len(calls) == 1:
                raise ReportAutomationError(
                    "EXPORT_MENU_OPEN_FAILED",
                    "匯出選單開啟失敗：NoPatternInterfaceError",
                )
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
            )

    def recover_pos(_config, _on_progress):  # type: ignore[no-untyped-def]
        recoveries.append("recover")
        return _ready_pos_window()

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_recovery_func=recover_pos,
    )

    summary = runner.run()

    assert calls == ["R01", "R01"]
    assert recoveries == ["recover"]
    assert summary.ok is True
    assert summary.completed == 1
    assert summary.failures == ()


def test_automation_runner_recovers_pos_session_invalid_and_retries_current_task(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.pos_recovery.enabled = True
    config.pos_recovery.max_restarts_per_run = 1
    for report in config.reports:
        report.enabled = report.id == "R05"
    calls: list[str] = []
    recoveries: list[str] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            if len(calls) == 1:
                raise ReportAutomationError(
                    "POS_SESSION_INVALID",
                    "SPA-POS 視窗連線已失效，只剩不可見的空白視窗控制項。",
                )
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
            )

    def recover_pos(_config, _on_progress):  # type: ignore[no-untyped-def]
        recoveries.append("recover")
        return _ready_pos_window()

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_recovery_func=recover_pos,
    )

    summary = runner.run()

    assert calls == ["R05", "R05"]
    assert recoveries == ["recover"]
    assert summary.ok is True
    assert summary.completed == 1
    assert summary.failures == ()


def test_automation_runner_forces_uia_reconnect_for_win32_menu_session_invalid(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.pos.backend = "auto"
    config.pos_recovery.enabled = True
    config.pos_recovery.max_restarts_per_run = 1
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R01"

    first_window = _ready_pos_window("統計報表")
    first_window._pos_report_bot_backend = "win32"
    uia_window = _ready_pos_window("統計報表")
    uia_window._pos_report_bot_backend = "uia"
    connect_backends: list[str] = []
    calls: list[str] = []
    automator_windows: list[FakePosControl] = []

    def connect_pos_window(**kwargs):  # type: ignore[no-untyped-def]
        backend = kwargs["backend"]
        connect_backends.append(backend)
        if backend == "uia":
            return uia_window
        return first_window

    class FakeAutomator:
        def __init__(self, window, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            automator_windows.append(window)

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            if len(calls) == 1:
                raise ReportAutomationError(
                    "POS_SESSION_INVALID",
                    "SPA-POS 已在主畫面，但目前 automation backend=win32 沒有列出根選單「統計報表」；需要重新連接或重啟 POS 後重試。",
                )
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run()

    assert summary.ok is True
    assert calls == ["R01", "R01"]
    assert connect_backends == ["auto", "uia"]
    assert automator_windows == [first_window, uia_window]


def test_automation_runner_uia_backend_switch_bypasses_disabled_pos_recovery(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.pos.backend = "auto"
    config.pos_recovery.enabled = False
    config.pos_recovery.max_restarts_per_run = 0
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R01"

    first_window = _ready_pos_window("統計報表")
    first_window._pos_report_bot_backend = "win32"
    uia_window = _ready_pos_window("統計報表")
    uia_window._pos_report_bot_backend = "uia"
    connect_backends: list[str] = []
    calls: list[str] = []

    def connect_pos_window(**kwargs):  # type: ignore[no-untyped-def]
        backend = kwargs["backend"]
        connect_backends.append(backend)
        return uia_window if backend == "uia" else first_window

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            if len(calls) == 1:
                raise ReportAutomationError(
                    "POS_SESSION_INVALID",
                    "SPA-POS 已在主畫面，但目前 automation backend=win32 沒有列出根選單「統計報表」；需要重新連接或重啟 POS 後重試。",
                )
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run()

    assert summary.ok is True
    assert calls == ["R01", "R01"]
    assert connect_backends == ["auto", "uia"]


def test_automation_runner_stops_recovery_after_restart_limit(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.pos_recovery.enabled = True
    config.pos_recovery.max_restarts_per_run = 1
    for report in config.reports:
        report.enabled = report.id == "R09"
    calls: list[str] = []
    recoveries: list[str] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            raise ReportAutomationError("POS_NOT_RESPONDING", "POS 無回應")

    def recover_pos(_config, _on_progress):  # type: ignore[no-untyped-def]
        recoveries.append("recover")
        return _ready_pos_window()

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_recovery_func=recover_pos,
    )

    summary = runner.run()

    assert calls == ["R09", "R09"]
    assert recoveries == ["recover"]
    assert summary.ok is False
    assert summary.completed == 0
    assert summary.error_code == "REPORT_RUN_FAILED"
    assert len(summary.failures) == 1
    assert summary.failures[0].error_code == "POS_NOT_RESPONDING"


def test_automation_runner_recovery_requires_pos_executable_path(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _disable_uploads(config)
    config.pos.executable_path = ""
    config.pos_recovery.relaunch_after_kill = True
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")

    try:
        runner._launch_pos_process(config)
    except RuntimeError as exc:
        assert "找不到 POS 啟動檔" in str(exc)
    else:
        raise AssertionError("Expected missing POS executable path to fail")


def test_automation_runner_launches_latest_clickonce_pos_when_configured_path_is_stale(
    monkeypatch,
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    stale = tmp_path / "stale" / "SPA1.exe"
    latest = tmp_path / "local" / "Apps" / "2.0" / "latest" / "SPA1.exe"
    latest.parent.mkdir(parents=True)
    latest.write_bytes(b"exe")
    config.pos.executable_path = str(stale)
    launched: list[list[str]] = []

    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "local"))
    monkeypatch.setattr(
        "pos_report_bot.app.automation_runner.subprocess.Popen",
        lambda command, cwd=None: launched.append(command),
    )

    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")
    runner._launch_pos_process(config)

    assert launched == [[str(latest)]]


def test_preparation_failure_diagnostic_records_pos_launch_lifecycle(
    monkeypatch,
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    pos_executable = tmp_path / "SPA1.exe"
    pos_executable.write_bytes(b"fake-pos")
    config.pos.executable_path = str(pos_executable)
    config.pos.startup_wait_seconds = 1
    _set_runtime_dirs(config, tmp_path)

    class ExitedProcess:
        pid = 4321

        @staticmethod
        def poll() -> int:
            return 23

    monkeypatch.setattr(
        "pos_report_bot.app.automation_runner.subprocess.Popen",
        lambda command, cwd=None: ExitedProcess(),
    )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
    )
    runner._launch_pos_process(config)

    output = SimpleNamespace(
        task_id="R01",
        output_filename="r01.xls",
        start_date="2026/07/01",
        end_date="2026/07/28",
        branch_code=None,
        branch_display_name=None,
        drive_folder_id=None,
        upload_enabled=False,
    )
    diagnostic_path = runner._write_preparation_failure_diagnostic(
        [output],
        error_code="POS_CONNECTION_FAILED",
        message="POS 視窗未出現",
    )

    assert diagnostic_path is not None
    payload = json.loads(diagnostic_path.read_text(encoding="utf-8"))
    runtime = payload["runtime"]
    assert runtime["rpa_executable_path"]
    assert runtime["pos_executable_path_configured"] == str(pos_executable)
    assert runtime["pos_executable_path_resolved"] == str(pos_executable)
    assert runtime["pos_launch"]["pid"] == 4321
    assert runtime["pos_launch"]["returncode"] == 23
    assert runtime["pos_launch"]["alive"] is False


def test_automation_runner_recovery_terminates_spa1_when_launch_path_is_appref(monkeypatch, tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.pos.executable_path = str(tmp_path / "SPA POS.appref-ms")
    calls: list[list[str]] = []

    monkeypatch.setattr("pos_report_bot.app.automation_runner.sys.platform", "win32")
    monkeypatch.setattr(
        "pos_report_bot.app.automation_runner.subprocess.run",
        lambda command, **_kwargs: calls.append(command),
    )

    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")
    runner._terminate_pos_process(config)

    assert calls == [["taskkill", "/F", "/T", "/IM", "SPA1.exe"]]


def test_automation_runner_logs_in_before_first_report_when_login_screen_is_visible(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.login.required = True
    config.login.username = "A0042"
    for report in config.reports:
        report.enabled = report.id == "R01"
    account_edit = FakePosControl("", "Edit")
    secret_edit = FakePosControl("", "Edit")
    login_window = FakePosControl("帳號登入", "Window")

    def complete_login() -> None:
        login_window.name = "SPA-POS 主畫面"
        login_window.children_controls = [FakePosControl("統計報表", "MenuItem")]

    login_button = FakePosControl("登入\r\nLogin", "Button", on_click=complete_login)
    login_window.children_controls = [
        FakePosControl("帳號", "Text"),
        account_edit,
        FakePosControl("密碼", "Text"),
        secret_edit,
        login_button,
    ]
    calls: list[str] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: login_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_login_secret_provider=lambda: "fake-login-secret",
    )

    summary = runner.run()

    assert account_edit.text_value == "A0042"
    assert secret_edit.text_value == "fake-login-secret"
    assert login_button.clicked is True
    assert calls == ["R01"]
    assert summary.ok is True


def test_automation_runner_logs_in_with_a_single_unnamed_login_button(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    account_edit = FakePosControl("", "Edit")
    secret_edit = FakePosControl("", "Edit")
    login_window = FakePosControl("帳號登入", "Window")

    def complete_login() -> None:
        login_window.name = "SPA-POS 主畫面"
        login_window.children_controls = [FakePosControl("統計報表", "MenuItem")]

    unnamed_login_button = FakePosControl("", "Button", on_click=complete_login)
    login_window.children_controls = [
        FakePosControl("帳號", "Text"),
        account_edit,
        FakePosControl("密碼", "Text"),
        secret_edit,
        unnamed_login_button,
    ]

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: login_window,
        pos_login_secret_provider=lambda: "fake-login-secret",
    )

    result = runner._login_if_required(config, login_window)

    assert result is login_window
    assert unnamed_login_button.clicked is True
    assert account_edit.text_value == "A0042"
    assert secret_edit.text_value == "fake-login-secret"


def test_automation_runner_treats_open_report_child_as_logged_in_after_reconnect(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    report_toolbar = FakePosControl("ReportToolBar", "Pane", automation_id="reportToolBar")
    report_form = FakePosControl(
        "商品銷售明細表",
        "Dialog",
        automation_id="ProdSale_Report",
        children=[report_toolbar],
    )
    report_window = FakePosControl(
        "SPA-POS Ver.1.5.19.15",
        "Dialog",
        automation_id="MainForm",
        children=[report_form],
    )

    def unexpected_password_lookup() -> str:
        raise AssertionError("a logged-in report child must not trigger login")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: report_window,
        pos_login_secret_provider=unexpected_password_lookup,
    )

    assert runner._login_if_required(config, report_window) is report_window


def test_automation_runner_ignores_hidden_stale_login_controls_on_logged_in_pos(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True

    class HiddenControl(FakePosControl):
        def is_visible(self) -> bool:
            return False

    main_window = FakePosControl(
        "SPA-POS 主畫面",
        "Window",
        children=[
            FakePosControl("統計報表", "MenuItem"),
            HiddenControl("帳號", "Text"),
            HiddenControl("密碼", "Text"),
            HiddenControl("登入", "Button"),
            HiddenControl("", "Edit"),
            HiddenControl("", "Edit"),
        ],
    )

    def unexpected_password_lookup() -> str:
        raise AssertionError("hidden stale login controls must not trigger credential lookup")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: main_window,
        pos_login_secret_provider=unexpected_password_lookup,
    )

    assert runner._login_if_required(config, main_window) is main_window


def test_automation_runner_closes_pos_even_when_legacy_flag_is_false(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.pos.close_after_run = False
    window = FakeClosablePosWindow()
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")

    runner._close_pos_after_run_if_configured(window, None)

    assert window.closed is True


def test_automation_runner_keeps_pos_open_for_verified_w02_unsaved_draft_resume(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    window = FakeClosablePosWindow()
    progress_messages: list[str] = []
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")
    runner._preserve_pos_after_run_reason = "w02_unsaved_draft_available_for_verified_resume"

    runner._close_pos_after_run_if_configured(
        window,
        lambda event: progress_messages.append(event.message),
    )

    assert window.closed is False
    assert any("已保留 SPA-POS 與未儲存草稿" in message for message in progress_messages)


def test_automation_runner_confirms_pos_exit_dialog_with_yes(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    window = FakeClosablePosWindow()
    confirmed: list[str] = []

    def confirm_exit() -> bool:
        confirmed.append("yes")
        return True

    window.confirm_exit = confirm_exit  # type: ignore[attr-defined]
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")

    runner._close_pos_after_run_if_configured(window, None)

    assert window.closed is True
    assert confirmed == ["yes"]


def test_pos_close_fallback_terminates_only_the_current_window_process(monkeypatch, tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    marker_path = Path(config.app.state_dir) / "pos_session_profile.json"
    marker_path.parent.mkdir(parents=True, exist_ok=True)
    marker_path.write_text('{"profile":"formal","process_id":4242}', encoding="utf-8")
    commands: list[list[str]] = []

    class StubbornWindow(FakeProcessBoundPosWindow):
        def close(self) -> None:
            return

        def is_visible(self) -> bool:
            return True

    monkeypatch.setattr(automation_runner_module.sys, "platform", "win32")
    monkeypatch.setattr(
        automation_runner_module.subprocess,
        "run",
        lambda command, **_kwargs: commands.append(command) or SimpleNamespace(returncode=0),
    )
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")
    monkeypatch.setattr(runner, "_confirm_pos_exit_dialog", lambda *_args, **_kwargs: False)

    runner._close_pos_after_run_if_configured(StubbornWindow(4242), None)

    assert commands == [["taskkill", "/F", "/T", "/PID", "4242"]]
    assert not marker_path.exists()


def test_pos_close_fallback_never_kills_all_by_image_when_window_pid_is_unknown(monkeypatch, tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    marker_path = Path(config.app.state_dir) / "pos_session_profile.json"
    marker_path.parent.mkdir(parents=True, exist_ok=True)
    marker_path.write_text('{"profile":"formal","process_id":4242}', encoding="utf-8")
    commands: list[list[str]] = []

    class StubbornWindow(FakeClosablePosWindow):
        def close(self) -> None:
            return

        def is_visible(self) -> bool:
            return True

    monkeypatch.setattr(automation_runner_module.sys, "platform", "win32")
    monkeypatch.setattr(
        automation_runner_module.subprocess,
        "run",
        lambda command, **_kwargs: commands.append(command) or SimpleNamespace(returncode=0),
    )
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")
    monkeypatch.setattr(runner, "_confirm_pos_exit_dialog", lambda *_args, **_kwargs: False)

    runner._close_pos_after_run_if_configured(StubbornWindow(), None)

    assert commands == []
    assert marker_path.exists()


def test_pos_recovery_with_existing_unknown_pid_window_disables_image_fallback(monkeypatch, tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    config.pos.startup_ini_selection_enabled = False
    config.login.required = False
    config.pos_recovery.kill_process_on_hang = True
    config.pos_recovery.relaunch_after_kill = False
    config.pos_recovery.restart_delay_seconds = 0
    current_window = FakeClosablePosWindow()
    ready_window = _ready_pos_window()
    calls: list[dict[str, object]] = []
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")
    runner._last_pos_window = current_window
    monkeypatch.setattr(
        runner,
        "_terminate_pos_process",
        lambda _config, **kwargs: calls.append(kwargs) or False,
    )
    monkeypatch.setattr(runner, "_wait_for_reconnected_pos_window", lambda _config: ready_window)

    assert runner._recover_pos_session(config, None) is ready_window
    assert calls == [{"window": current_window, "allow_image_fallback": False}]


def test_automation_runner_handles_pos_update_dialog_before_typing_login(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    config.pos.startup_wait_seconds = 1
    config.pos_update.restart_wait_seconds = 0
    config.pos_update.max_restart_wait_seconds = 1

    stale_account_edit = FakePosControl("", "Edit")
    stale_secret_edit = FakePosControl("", "Edit")
    account_edit = FakePosControl("", "Edit")
    secret_edit = FakePosControl("", "Edit")
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")])
    updated_login_window = FakePosControl("帳號登入", "Window")
    current_window: FakePosControl = updated_login_window

    def complete_login() -> None:
        nonlocal current_window
        current_window = main_window

    login_button = FakePosControl("登入", "Button", on_click=complete_login)
    updated_login_window.children_controls = [
        FakePosControl("帳號", "Text"),
        account_edit,
        FakePosControl("密碼", "Text"),
        secret_edit,
        login_button,
    ]

    def accept_update() -> None:
        nonlocal current_window
        current_window = updated_login_window

    yes_button = FakePosControl("是(Y)", "Button", on_click=accept_update)
    update_dialog = FakePosControl(
        "程式更新需重新啟動",
        "Window",
        children=[
            FakePosControl("新版程式已經下載安裝完成(1.5.18.77),需要重新啟動程式!", "Text"),
            yes_button,
            FakePosControl("否(N)", "Button"),
        ],
    )
    blocked_login_window = FakePosControl(
        "SPA-POS",
        "Window",
        children=[
            FakePosControl("帳號登入 (台灣凱惠SPA資訊系統 Ver.1.5.18.73)", "Text"),
            FakePosControl("帳號", "Text"),
            stale_account_edit,
            FakePosControl("密碼", "Text"),
            stale_secret_edit,
            FakePosControl("登入", "Button"),
            update_dialog,
        ],
    )
    current_window = blocked_login_window

    def connect_pos(**_kwargs):  # type: ignore[no-untyped-def]
        return current_window

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos,
        pos_login_secret_provider=lambda: "1234",
    )

    result = runner._login_if_required(config, blocked_login_window)

    assert result is main_window
    assert yes_button.clicked is True
    assert stale_account_edit.text_value == ""
    assert stale_secret_edit.text_value == ""
    assert account_edit.text_value == "A0042"
    assert secret_edit.text_value == "1234"
    assert login_button.clicked is True




def test_automation_runner_handles_top_level_pos_update_dialog_before_typing_login(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    config.pos.startup_wait_seconds = 1
    config.pos_update.restart_wait_seconds = 0
    config.pos_update.max_restart_wait_seconds = 1

    stale_account_edit = FakePosControl("", "Edit")
    stale_secret_edit = FakePosControl("", "Edit")
    blocked_login_window = FakePosControl(
        "帳號登入",
        "Window",
        children=[
            FakePosControl("帳號", "Text"),
            stale_account_edit,
            FakePosControl("密碼", "Text"),
            stale_secret_edit,
            FakePosControl("登入", "Button"),
        ],
        enabled=False,
    )
    account_edit = FakePosControl("", "Edit")
    secret_edit = FakePosControl("", "Edit")
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")])
    updated_login_window = FakePosControl("帳號登入", "Window")
    current_window: FakePosControl = blocked_login_window
    update_visible = True

    def complete_login() -> None:
        nonlocal current_window
        current_window = main_window

    login_button = FakePosControl("登入", "Button", on_click=complete_login)
    updated_login_window.children_controls = [
        FakePosControl("帳號", "Text"),
        account_edit,
        FakePosControl("密碼", "Text"),
        secret_edit,
        login_button,
    ]

    def accept_update() -> None:
        nonlocal current_window, update_visible
        update_visible = False
        current_window = updated_login_window

    yes_button = FakePosControl("是(Y)", "Button", on_click=accept_update)
    update_dialog = FakePosControl(
        "程式更新需重新啟動",
        "Window",
        children=[
            FakePosControl("新版程式已經下載安裝完成(1.5.18.89),需要重新啟動程式!", "Text"),
            yes_button,
            FakePosControl("否(N)", "Button"),
        ],
    )

    def connect_pos_window(**kwargs):  # type: ignore[no-untyped-def]
        title = str(kwargs.get("window_title_contains", ""))
        if config.pos_update.update_dialog_title_contains in title:
            if update_visible:
                return update_dialog
            raise UiProbeError("update dialog is closed")
        return current_window

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        pos_login_secret_provider=lambda: "1234",
    )

    result = runner._login_if_required(config, blocked_login_window)

    assert result is main_window
    assert yes_button.clicked is True
    assert stale_account_edit.text_value == ""
    assert stale_secret_edit.text_value == ""
    assert account_edit.text_value == "A0042"
    assert secret_edit.text_value == "1234"
    assert login_button.clicked is True


def test_automation_runner_relaunches_pos_when_update_does_not_auto_restart(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    config.pos.startup_wait_seconds = 1
    config.pos_update.restart_wait_seconds = 0
    config.pos_update.max_restart_wait_seconds = 1

    blocked_login_window = FakePosControl("帳號登入", "Window", enabled=False)
    account_edit = FakePosControl("", "Edit")
    secret_edit = FakePosControl("", "Edit")
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")])
    relaunched_login_window = FakePosControl("帳號登入", "Window")
    update_visible = True
    pos_available = False
    launched = False
    current_window: FakePosControl = blocked_login_window

    def complete_login() -> None:
        nonlocal current_window
        current_window = main_window

    relaunched_login_window.children_controls = [
        FakePosControl("帳號", "Text"),
        account_edit,
        FakePosControl("密碼", "Text"),
        secret_edit,
        FakePosControl("登入", "Button", on_click=complete_login),
    ]

    def accept_update() -> None:
        nonlocal update_visible, pos_available
        update_visible = False
        pos_available = False

    update_dialog = FakePosControl(
        "程式更新需重新啟動",
        "Window",
        children=[
            FakePosControl("新版程式已經下載安裝完成，需要重新啟動程式", "Text"),
            FakePosControl("是(Y)", "Button", on_click=accept_update),
        ],
    )

    def connect_pos_window(**kwargs):  # type: ignore[no-untyped-def]
        title = str(kwargs.get("window_title_contains", ""))
        if config.pos_update.update_dialog_title_contains in title:
            if update_visible:
                return update_dialog
            raise UiProbeError("update dialog is closed")
        if pos_available:
            return current_window
        raise UiProbeError("POS is still restarting")

    def launch_pos_process(_config) -> None:  # type: ignore[no-untyped-def]
        nonlocal launched, pos_available, current_window
        launched = True
        pos_available = True
        current_window = relaunched_login_window

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        pos_login_secret_provider=lambda: "1234",
    )
    runner._launch_pos_process = launch_pos_process  # type: ignore[method-assign]

    result = runner._login_if_required(config, blocked_login_window)

    assert result is main_window
    assert launched is True
    assert account_edit.text_value == "A0042"
    assert secret_edit.text_value == "1234"




def test_automation_runner_caps_update_restart_wait_to_five_minutes(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    config.pos.startup_wait_seconds = 1
    window = FakePosControl("帳號登入", "Window")
    launched = False
    pos_available = False
    elapsed = 0.0
    sleeps: list[float] = []

    def fake_monotonic() -> float:
        return elapsed

    def fake_sleep(seconds: float) -> None:
        nonlocal elapsed
        sleeps.append(seconds)
        elapsed += seconds

    def connect_pos_window(**_kwargs):  # type: ignore[no-untyped-def]
        if pos_available:
            return window
        raise UiProbeError("POS is still restarting")

    def launch_pos_process(_config) -> None:  # type: ignore[no-untyped-def]
        nonlocal launched, pos_available
        launched = True
        pos_available = True

    monkeypatch.setattr(automation_runner_module, "monotonic", fake_monotonic)
    monkeypatch.setattr(automation_runner_module, "sleep", fake_sleep)
    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
    )
    runner._launch_pos_process = launch_pos_process  # type: ignore[method-assign]

    result = runner._wait_for_pos_after_update_restart(
        config,
        SimpleNamespace(restart_wait_seconds=0, max_restart_wait_seconds=600),
    )

    assert result is window
    assert launched is True
    assert sum(sleeps) == 300


def test_automation_runner_dismisses_login_error_dialog_before_retyping_credentials(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    config.login.timeout_seconds = 1

    account_edit = FakePosControl("", "Edit")
    secret_edit = FakePosControl("", "Edit")
    account_edit.text_value = "1234"
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")])
    login_window = FakePosControl("帳號登入", "Window")
    current_window: FakePosControl = login_window

    def close_error() -> None:
        login_window.children_controls = [child for child in login_window.children_controls if child is not error_dialog]

    ok_button = FakePosControl("確定", "Button", on_click=close_error)
    error_dialog = FakePosControl(
        "錯誤警告",
        "Window",
        children=[
            FakePosControl("帳號輸入錯誤,查無此帳號!", "Text"),
            ok_button,
        ],
    )

    def complete_login() -> None:
        nonlocal current_window
        current_window = main_window

    login_button = FakePosControl("登入", "Button", on_click=complete_login)
    login_window.children_controls = [
        FakePosControl("帳號", "Text"),
        account_edit,
        FakePosControl("密碼", "Text"),
        secret_edit,
        login_button,
        error_dialog,
    ]

    def connect_pos(**_kwargs):  # type: ignore[no-untyped-def]
        return current_window

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos,
        pos_login_secret_provider=lambda: "1234",
    )

    result = runner._login_if_required(config, login_window)

    assert result is main_window
    assert ok_button.clicked is True
    assert account_edit.text_value == "A0042"
    assert secret_edit.text_value == "1234"
    assert login_button.clicked is True


def test_automation_runner_stops_when_login_error_dialog_cannot_be_dismissed(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    config.login.timeout_seconds = 1

    account_edit = FakePosControl("", "Edit")
    secret_edit = FakePosControl("", "Edit")
    error_dialog = FakePosControl(
        "錯誤警告",
        "Window",
        children=[FakePosControl("帳號輸入錯誤,查無此帳號!", "Text")],
    )
    login_window = FakePosControl(
        "帳號登入",
        "Window",
        children=[
            FakePosControl("帳號", "Text"),
            account_edit,
            FakePosControl("密碼", "Text"),
            secret_edit,
            FakePosControl("登入", "Button"),
            error_dialog,
        ],
    )
    sent_keys: list[str] = []

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: login_window,
        pos_login_secret_provider=lambda: "1234",
        keyboard_sender=lambda keys, **_kwargs: sent_keys.append(keys),  # type: ignore[arg-type]
    )

    with pytest.raises(RuntimeError, match="POS 顯示登入失敗警告"):
        runner._login_if_required(config, login_window)
    assert account_edit.text_value == ""
    assert secret_edit.text_value == ""
    assert sent_keys == []


def test_automation_runner_login_does_not_trigger_dynamic_login_pos_lookup(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.login.required = True
    config.login.username = "A0042"
    for report in config.reports:
        report.enabled = report.id == "R01"

    class DynamicLookupLoginWindow(FakePosControl):
        def __getattr__(self, name: str):  # type: ignore[no-untyped-def]
            if name == "login_pos":
                raise RuntimeError("Neither GUI element (wrapper) nor wrapper method 'login_pos' were found")
            raise AttributeError(name)

    account_edit = FakePosControl("", "Edit")
    secret_edit = FakePosControl("", "Edit")
    login_window = DynamicLookupLoginWindow("帳號登入", "Window")

    def complete_login() -> None:
        login_window.name = "SPA-POS 主畫面"
        login_window.children_controls = [FakePosControl("統計報表", "MenuItem")]

    login_window.children_controls = [
        FakePosControl("帳號", "Text"),
        account_edit,
        FakePosControl("密碼", "Text"),
        secret_edit,
        FakePosControl("登入", "Button", on_click=complete_login),
    ]

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: login_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_login_secret_provider=lambda: "fake-login-secret",
    )

    summary = runner.run()

    assert summary.ok is True
    assert account_edit.text_value == "A0042"
    assert secret_edit.text_value == "fake-login-secret"


def test_automation_runner_reconnects_when_login_window_handle_becomes_invalid(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.login.required = True
    config.login.username = "A0042"
    config.pos.startup_wait_seconds = 2
    for report in config.reports:
        report.enabled = report.id == "R01"
    calls: list[str] = []
    first_account_edit = FakePosControl("", "Edit")
    second_account_edit = FakePosControl("", "Edit")
    second_secret_edit = FakePosControl("", "Edit")
    second_login_window = FakePosControl("帳號登入", "Window")

    class InvalidHandleEdit(FakePosControl):
        def set_edit_text(self, value: str) -> None:
            raise RuntimeError("Handle 1377002 is not a vaild window handle")

    def complete_login() -> None:
        second_login_window.name = "SPA-POS 主畫面"
        second_login_window.children_controls = [FakePosControl("統計報表", "MenuItem")]

    first_login_window = FakePosControl(
        "帳號登入",
        "Window",
        children=[
            FakePosControl("帳號", "Text"),
            first_account_edit,
            FakePosControl("密碼", "Text"),
            InvalidHandleEdit("", "Edit"),
            FakePosControl("登入", "Button"),
        ],
    )
    second_login_window.children_controls = [
        FakePosControl("帳號", "Text"),
        second_account_edit,
        FakePosControl("密碼", "Text"),
        second_secret_edit,
        FakePosControl("登入", "Button", on_click=complete_login),
    ]
    windows = [first_login_window, second_login_window, second_login_window]

    def connect_pos(**_kwargs):  # type: ignore[no-untyped-def]
        return windows.pop(0) if windows else second_login_window

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_login_secret_provider=lambda: "fake-login-secret",
    )

    summary = runner.run()

    assert summary.ok is True
    assert second_account_edit.text_value == "A0042"
    assert second_secret_edit.text_value == "fake-login-secret"
    assert calls == ["R01"]


def test_automation_runner_control_login_waits_for_main_menu_not_unreadable_window(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    config.login.timeout_seconds = 2

    class UnreadableWindow(FakePosControl):
        def window_text(self) -> str:
            return ""

        def children(self) -> list[FakePosControl]:
            return []

        def descendants(self) -> list[FakePosControl]:
            return []

    account_edit = FakePosControl("", "Edit")
    secret_edit = FakePosControl("", "Edit")
    login_window = FakePosControl("帳號登入", "Window")
    unreadable_window = UnreadableWindow("SPA-POS", "Window")
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")])
    login_clicked = False
    post_click_connects = 0

    def complete_login() -> None:
        nonlocal login_clicked
        login_clicked = True

    def connect_pos_window(**_kwargs):  # type: ignore[no-untyped-def]
        nonlocal post_click_connects
        if login_clicked:
            post_click_connects += 1
            return unreadable_window if post_click_connects == 1 else main_window
        return login_window

    login_window.children_controls = [
        FakePosControl("帳號", "Text"),
        account_edit,
        FakePosControl("密碼", "Text"),
        secret_edit,
        FakePosControl("登入", "Button", on_click=complete_login),
    ]
    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        pos_login_secret_provider=lambda: "fake-login-secret",
        keyboard_sender=None,
    )

    result = runner._login_if_required(config, login_window)

    assert result is main_window
    assert account_edit.text_value == "A0042"
    assert secret_edit.text_value == "fake-login-secret"
    assert post_click_connects == 2


def test_automation_runner_uses_keyboard_login_when_login_handle_is_invalid(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.login.required = True
    config.login.username = "A0042"
    config.pos.startup_wait_seconds = 2
    for report in config.reports:
        report.enabled = report.id == "R01"

    class InvalidHandleEdit(FakePosControl):
        def set_edit_text(self, value: str) -> None:
            raise RuntimeError("Handle 4196182 is not a vaild window handle")

    login_window = FakePosControl(
        "帳號登入",
        "Window",
        children=[
            FakePosControl("帳號", "Text"),
            FakePosControl("", "Edit"),
            FakePosControl("密碼", "Text"),
            InvalidHandleEdit("", "Edit"),
            FakePosControl("登入", "Button"),
        ],
    )
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")])
    current_window = login_window
    sent_keys: list[str] = []
    calls: list[str] = []

    def send_keys(keys: str, **_kwargs) -> None:  # type: ignore[no-untyped-def]
        nonlocal current_window
        sent_keys.append(keys)
        if keys == "{ENTER}":
            current_window = main_window

    def connect_pos(**_kwargs):  # type: ignore[no-untyped-def]
        return current_window

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_login_secret_provider=lambda: "test-password",
        keyboard_sender=send_keys,
    )

    summary = runner.run()

    assert summary.ok is True
    assert sent_keys == ["^a{BACKSPACE}", "A0042", "{TAB}", "^a{BACKSPACE}", "test-password", "{ENTER}"]
    assert calls == ["R01"]


def test_automation_runner_falls_back_to_keyboard_when_control_login_does_not_reach_main_screen(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.login.required = True
    config.login.username = "A0042"
    for report in config.reports:
        report.enabled = report.id == "R01"
    account_edit = FakePosControl("", "Edit")
    secret_edit = FakePosControl("", "Edit")
    login_button = FakePosControl("登入", "Button")
    login_window = FakePosControl(
        "帳號登入",
        "Window",
        children=[
            FakePosControl("帳號", "Text"),
            account_edit,
            FakePosControl("密碼", "Text"),
            secret_edit,
            login_button,
        ],
    )
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")])
    current_window = login_window
    sent_keys: list[str] = []
    calls: list[str] = []

    def send_keys(keys: str, **_kwargs) -> None:  # type: ignore[no-untyped-def]
        nonlocal current_window
        sent_keys.append(keys)
        if keys == "{ENTER}":
            current_window = main_window

    def connect_pos(**_kwargs):  # type: ignore[no-untyped-def]
        return current_window

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_login_secret_provider=lambda: "test-password",
        keyboard_sender=send_keys,
    )

    summary = runner.run()

    assert summary.ok is True
    assert sent_keys == ["^a{BACKSPACE}", "A0042", "{TAB}", "^a{BACKSPACE}", "test-password", "{ENTER}"]
    assert account_edit.text_value == "A0042"
    assert secret_edit.text_value == "test-password"
    assert login_button.clicked is True
    assert calls == ["R01"]


def test_automation_runner_uses_control_login_before_keyboard_when_edits_are_available(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.login.required = True
    config.login.username = "A0042"
    for report in config.reports:
        report.enabled = report.id == "R01"
    unrelated_edit = FakePosControl("", "Edit")
    account_edit = FakePosControl("", "Edit")
    secret_edit = FakePosControl("", "Edit")
    login_window = FakePosControl("帳號登入", "Window")
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")])
    current_window = login_window
    sent_keys: list[str] = []
    calls: list[str] = []

    def complete_login() -> None:
        nonlocal current_window
        current_window = main_window

    login_button = FakePosControl("登入", "Button", on_click=complete_login)
    login_window.children_controls = [
        unrelated_edit,
        FakePosControl("帳號", "Text"),
        account_edit,
        FakePosControl("密碼", "Text"),
        secret_edit,
        login_button,
    ]

    def connect_pos(**_kwargs):  # type: ignore[no-untyped-def]
        return current_window

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_login_secret_provider=lambda: "test-password",
        keyboard_sender=lambda keys, **_kwargs: sent_keys.append(keys),  # type: ignore[arg-type]
    )

    summary = runner.run()

    assert summary.ok is True
    assert sent_keys == []
    assert unrelated_edit.text_value == ""
    assert account_edit.text_value == "A0042"
    assert secret_edit.text_value == "test-password"
    assert login_button.clicked is True
    assert calls == ["R01"]


def test_automation_runner_maps_login_edits_by_label_row_when_child_order_is_reversed(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.login.required = True
    config.login.username = "A0042"
    for report in config.reports:
        report.enabled = report.id == "R01"
    account_label = FakeRectPosControl("帳號", "Text", rect=(440, 246, 475, 272))
    secret_label = FakeRectPosControl("密碼", "Text", rect=(440, 286, 475, 312))
    account_edit = FakeRectPosControl("", "Edit", rect=(484, 246, 660, 272))
    secret_edit = FakeRectPosControl("", "Edit", rect=(484, 286, 660, 312))
    login_window = FakePosControl("帳號登入", "Window")
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")])
    current_window = login_window
    sent_keys: list[str] = []
    calls: list[str] = []

    def complete_login() -> None:
        nonlocal current_window
        current_window = main_window

    login_button = FakeRectPosControl("登入", "Button", rect=(690, 248, 784, 310), on_click=complete_login)
    login_window.children_controls = [
        secret_edit,
        account_edit,
        secret_label,
        account_label,
        login_button,
    ]

    def connect_pos(**_kwargs):  # type: ignore[no-untyped-def]
        return current_window

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_login_secret_provider=lambda: "1234",
        keyboard_sender=lambda keys, **_kwargs: sent_keys.append(keys),  # type: ignore[arg-type]
    )

    summary = runner.run()

    assert summary.ok is True
    assert sent_keys == []
    assert account_edit.text_value == "A0042"
    assert secret_edit.text_value == "1234"
    assert calls == ["R01"]


def test_automation_runner_preserves_partial_label_login_mapping_when_filling_missing_edits(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.login.required = True
    config.login.username = "A0042"
    for report in config.reports:
        report.enabled = report.id == "R01"
    account_label = FakeRectPosControl("帳號", "Text", rect=(440, 246, 475, 272))
    account_edit = FakeRectPosControl("", "Edit", rect=(484, 246, 660, 272))
    secret_edit = FakeRectPosControl("", "Edit", rect=(484, 286, 660, 312))
    login_window = FakePosControl("帳號登入", "Window")
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")])
    current_window = login_window
    sent_keys: list[str] = []
    calls: list[str] = []

    def complete_login() -> None:
        nonlocal current_window
        current_window = main_window

    login_button = FakeRectPosControl("登入", "Button", rect=(690, 248, 784, 310), on_click=complete_login)
    login_window.children_controls = [
        secret_edit,
        account_edit,
        account_label,
        login_button,
    ]

    def connect_pos(**_kwargs):  # type: ignore[no-untyped-def]
        return current_window

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_login_secret_provider=lambda: "1234",
        keyboard_sender=lambda keys, **_kwargs: sent_keys.append(keys),  # type: ignore[arg-type]
    )

    summary = runner.run()

    assert summary.ok is True
    assert sent_keys == []
    assert account_edit.text_value == "A0042"
    assert secret_edit.text_value == "1234"
    assert calls == ["R01"]


def test_automation_runner_uses_named_edit_login_fields_even_when_child_order_is_reversed(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.login.required = True
    config.login.username = "A0042"
    for report in config.reports:
        report.enabled = report.id == "R01"
    account_edit = FakePosControl("帳號", "Edit")
    secret_edit = FakePosControl("密碼", "Edit")
    login_window = FakePosControl("帳號登入", "Window")
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")])
    current_window = login_window
    sent_keys: list[str] = []
    calls: list[str] = []

    def complete_login() -> None:
        nonlocal current_window
        current_window = main_window

    login_button = FakePosControl("登入", "Button", on_click=complete_login)
    login_window.children_controls = [
        secret_edit,
        account_edit,
        login_button,
    ]

    def connect_pos(**_kwargs):  # type: ignore[no-untyped-def]
        return current_window

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_login_secret_provider=lambda: "1234",
        keyboard_sender=lambda keys, **_kwargs: sent_keys.append(keys),  # type: ignore[arg-type]
    )

    summary = runner.run()

    assert summary.ok is True
    assert sent_keys == []
    assert account_edit.text_value == "A0042"
    assert secret_edit.text_value == "1234"
    assert calls == ["R01"]


def test_automation_runner_maps_decorated_login_labels_by_normalized_text(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.login.required = True
    config.login.username = "A0042"
    for report in config.reports:
        report.enabled = report.id == "R01"
    account_label = FakeRectPosControl("帳 號(&U)", "Text", rect=(440, 246, 475, 272))
    secret_label = FakeRectPosControl("密 碼(&P)", "Text", rect=(440, 286, 475, 312))
    account_edit = FakeRectPosControl("", "Edit", rect=(484, 246, 660, 272))
    secret_edit = FakeRectPosControl("", "Edit", rect=(484, 286, 660, 312))
    login_window = FakePosControl("帳號登入", "Window")
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")])
    current_window = login_window
    calls: list[str] = []

    def complete_login() -> None:
        nonlocal current_window
        current_window = main_window

    login_button = FakeRectPosControl("登入", "Button", rect=(690, 248, 784, 310), on_click=complete_login)
    login_window.children_controls = [
        secret_edit,
        account_edit,
        secret_label,
        account_label,
        login_button,
    ]

    def connect_pos(**_kwargs):  # type: ignore[no-untyped-def]
        return current_window

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_login_secret_provider=lambda: "1234",
    )

    summary = runner.run()

    assert summary.ok is True
    assert account_edit.text_value == "A0042"
    assert secret_edit.text_value == "1234"
    assert calls == ["R01"]


def test_automation_runner_keyboard_login_failure_does_not_touch_stale_controls(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    config.login.timeout_seconds = 1

    class StaleEdit(FakePosControl):
        def set_edit_text(self, value: str) -> None:
            raise RuntimeError("stale edit wrapper should not be used")

    class StaleButton(FakePosControl):
        def click(self) -> None:
            raise RuntimeError("stale login button should not be used")

        def click_input(self) -> None:
            raise RuntimeError("stale login button should not be used")

    login_window = FakePosControl(
        "帳號登入",
        "Window",
        children=[
            FakePosControl("帳號", "Text"),
            StaleEdit("", "Edit"),
            FakePosControl("密碼", "Text"),
            StaleEdit("", "Edit"),
            StaleButton("登入", "Button"),
        ],
    )
    sent_keys: list[str] = []
    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: login_window,
        pos_login_secret_provider=lambda: "test-password",
        keyboard_sender=lambda keys, **_kwargs: sent_keys.append(keys),  # type: ignore[arg-type]
    )

    try:
        runner._login_if_required(config, login_window)
    except RuntimeError as exc:
        assert "POS_LOGIN_FAILED" in str(exc)
    else:
        raise AssertionError("keyboard login failure should stop without touching stale controls")

    assert sent_keys.count("{ENTER}") == 2


def test_automation_runner_keyboard_login_when_login_detection_wrapper_is_stale(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    config.login.timeout_seconds = 1

    class StaleLoginWindow(FakePosControl):
        def window_text(self) -> str:
            raise RuntimeError("Handle 2361392 is not a vaild window handle")

        def children(self) -> list[FakePosControl]:
            raise RuntimeError("Handle 2361392 is not a vaild window handle")

        def descendants(self) -> list[FakePosControl]:
            raise RuntimeError("Handle 2361392 is not a vaild window handle")

    login_window = StaleLoginWindow("帳號登入", "Window")
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")])
    current_window = login_window
    sent_keys: list[str] = []

    def send_keys(keys: str, **_kwargs) -> None:  # type: ignore[no-untyped-def]
        nonlocal current_window
        sent_keys.append(keys)
        if keys == "{ENTER}":
            current_window = main_window

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: current_window,
        pos_login_secret_provider=lambda: "test-password",
        keyboard_sender=send_keys,
    )

    result = runner._login_if_required(config, login_window)

    assert result is main_window
    assert sent_keys == ["^a{BACKSPACE}", "A0042", "{TAB}", "^a{BACKSPACE}", "test-password", "{ENTER}"]


def test_automation_runner_keyboard_login_waits_for_main_menu_not_unreadable_window(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    config.login.timeout_seconds = 2

    class UnreadableLoginWindow(FakePosControl):
        def window_text(self) -> str:
            return ""

        def children(self) -> list[FakePosControl]:
            return []

        def descendants(self) -> list[FakePosControl]:
            return []

    login_window = FakePosControl("帳號登入", "Window")
    unreadable_window = UnreadableLoginWindow("帳號登入", "Window")
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")])
    current_window = login_window
    post_enter_connects = 0
    sent_keys: list[str] = []

    def connect_pos_window(**_kwargs):  # type: ignore[no-untyped-def]
        nonlocal post_enter_connects
        if sent_keys and sent_keys[-1] == "{ENTER}":
            post_enter_connects += 1
            return unreadable_window if post_enter_connects == 1 else main_window
        return current_window

    def send_keys(keys: str, **_kwargs) -> None:  # type: ignore[no-untyped-def]
        sent_keys.append(keys)

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        pos_login_secret_provider=lambda: "test-password",
        keyboard_sender=send_keys,
    )

    result = runner._keyboard_login_and_wait(config, "test-password", window=login_window)

    assert result is main_window
    assert post_enter_connects == 2
    assert sent_keys == ["^a{BACKSPACE}", "A0042", "{TAB}", "^a{BACKSPACE}", "test-password", "{ENTER}"]


def test_automation_runner_keyboard_login_focuses_first_edit_before_typing(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    config.login.timeout_seconds = 1

    class FocusableEdit(FakePosControl):
        def __init__(self) -> None:
            super().__init__("", "Edit")
            self.focused = False

        def set_focus(self) -> None:
            self.focused = True

    account_edit = FocusableEdit()
    login_window = FakePosControl(
        "帳號登入",
        "Window",
        children=[
            FakePosControl("帳號", "Text"),
            account_edit,
            FakePosControl("密碼", "Text"),
            FakePosControl("", "Edit"),
            FakePosControl("登入", "Button"),
        ],
    )
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")])
    sent_keys: list[str] = []

    def connect_pos_window(**_kwargs):  # type: ignore[no-untyped-def]
        return main_window if sent_keys and sent_keys[-1] == "{ENTER}" and account_edit.focused else login_window

    def send_keys(keys: str, **_kwargs) -> None:  # type: ignore[no-untyped-def]
        sent_keys.append(keys)

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        pos_login_secret_provider=lambda: "test-password",
        keyboard_sender=send_keys,
    )

    result = runner._keyboard_login_and_wait(config, "test-password", window=login_window)

    assert result is main_window
    assert account_edit.focused is True


def test_automation_runner_keyboard_login_focuses_named_account_edit_when_child_order_is_reversed(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    config.login.timeout_seconds = 1

    class FocusableEdit(FakePosControl):
        def __init__(self, name: str) -> None:
            super().__init__(name, "Edit")
            self.focused = False

        def set_focus(self) -> None:
            self.focused = True

    secret_edit = FocusableEdit("密碼")
    account_edit = FocusableEdit("帳號")
    login_window = FakePosControl(
        "帳號登入",
        "Window",
        children=[
            secret_edit,
            account_edit,
            FakePosControl("登入", "Button"),
        ],
    )
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")])
    sent_keys: list[str] = []

    def connect_pos_window(**_kwargs):  # type: ignore[no-untyped-def]
        return main_window if sent_keys and sent_keys[-1] == "{ENTER}" and account_edit.focused else login_window

    def send_keys(keys: str, **_kwargs) -> None:  # type: ignore[no-untyped-def]
        sent_keys.append(keys)

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        pos_login_secret_provider=lambda: "test-password",
        keyboard_sender=send_keys,
    )

    result = runner._keyboard_login_and_wait(config, "test-password", window=login_window)

    assert result is main_window
    assert account_edit.focused is True
    assert secret_edit.focused is False


def test_automation_runner_keyboard_login_skips_disabled_login_edit(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    config.login.timeout_seconds = 1

    class FocusableEdit(FakePosControl):
        def __init__(self, *, enabled: bool) -> None:
            super().__init__("", "Edit", enabled=enabled)
            self.focused = False

        def set_focus(self) -> None:
            self.focused = True

    disabled_edit = FocusableEdit(enabled=False)
    enabled_edit = FocusableEdit(enabled=True)
    login_window = FakePosControl(
        "帳號登入",
        "Window",
        children=[
            FakePosControl("帳號", "Text"),
            disabled_edit,
            enabled_edit,
            FakePosControl("密碼", "Text"),
            FakePosControl("", "Edit"),
            FakePosControl("登入", "Button"),
        ],
    )
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")])
    sent_keys: list[str] = []

    def connect_pos_window(**_kwargs):  # type: ignore[no-untyped-def]
        return main_window if sent_keys and sent_keys[-1] == "{ENTER}" and enabled_edit.focused else login_window

    def send_keys(keys: str, **_kwargs) -> None:  # type: ignore[no-untyped-def]
        sent_keys.append(keys)

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        pos_login_secret_provider=lambda: "test-password",
        keyboard_sender=send_keys,
    )

    result = runner._keyboard_login_and_wait(config, "test-password", window=login_window)

    assert result is main_window
    assert disabled_edit.focused is False
    assert enabled_edit.focused is True


def test_automation_runner_keyboard_login_accepts_decorated_main_menu_text(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    config.login.timeout_seconds = 1

    login_window = FakePosControl("帳號登入", "Window")
    main_window = FakePosControl(
        "SPA-POS 主畫面",
        "Window",
        children=[FakePosControl(" 統計 報表(&R) ", "MenuItem")],
    )
    sent_keys: list[str] = []

    def connect_pos_window(**_kwargs):  # type: ignore[no-untyped-def]
        return main_window if sent_keys and sent_keys[-1] == "{ENTER}" else login_window

    def send_keys(keys: str, **_kwargs) -> None:  # type: ignore[no-untyped-def]
        sent_keys.append(keys)

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        pos_login_secret_provider=lambda: "test-password",
        keyboard_sender=send_keys,
    )

    result = runner._keyboard_login_and_wait(config, "test-password", window=login_window)

    assert result is main_window


def test_automation_runner_keyboard_login_does_not_hide_screen_check_errors(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    config.login.timeout_seconds = 1
    login_window = FakePosControl("帳號登入", "Window")
    sent_keys: list[str] = []

    class BrokenRunner(AutomationRunner):
        def _pos_main_screen_visible(self, window):  # type: ignore[no-untyped-def]
            raise AssertionError("screen check bug")

    def send_keys(keys: str, **_kwargs) -> None:  # type: ignore[no-untyped-def]
        sent_keys.append(keys)

    runner = BrokenRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: login_window,
        pos_login_secret_provider=lambda: "test-password",
        keyboard_sender=send_keys,
    )

    try:
        runner._keyboard_login_and_wait(config, "test-password", window=login_window)
    except AssertionError as exc:
        assert str(exc) == "screen check bug"
    else:
        raise AssertionError("screen check programming error was hidden")


def test_automation_runner_keyboard_login_uses_current_main_window_when_reconnect_fails(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    config.login.timeout_seconds = 1
    login_window = FakePosControl("帳號登入", "Window")
    sent_keys: list[str] = []

    def send_keys(keys: str, **_kwargs) -> None:  # type: ignore[no-untyped-def]
        sent_keys.append(keys)
        if keys == "{ENTER}":
            login_window.name = "SPA-POS 主畫面"
            login_window.children_controls = [FakePosControl(" 庫存 管理(&I) ", "MenuItem")]

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(UiProbeError("temporary reconnect failure")),
        pos_login_secret_provider=lambda: "test-password",
        keyboard_sender=send_keys,
    )

    result = runner._keyboard_login_and_wait(config, "test-password", window=login_window)

    assert result is login_window


def test_automation_runner_keyboard_login_does_not_type_when_window_already_main_screen(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.required = True
    config.login.username = "A0042"
    config.login.timeout_seconds = 1
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl(" 庫存 管理(&I) ", "MenuItem")])
    sent_keys: list[str] = []

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: main_window,
        pos_login_secret_provider=lambda: "test-password",
        keyboard_sender=lambda keys, **_kwargs: sent_keys.append(keys),  # type: ignore[arg-type]
    )

    result = runner._keyboard_login_and_wait(config, "test-password", window=main_window)

    assert result is main_window
    assert sent_keys == []


def test_automation_runner_wait_for_login_complete_uses_current_main_window_when_reconnect_fails(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.login.timeout_seconds = 1
    main_window = FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl(" 庫存 管理(&I) ", "MenuItem")])
    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(UiProbeError("temporary reconnect failure")),
    )

    result = runner._wait_for_login_complete(config, main_window)

    assert result is main_window


def test_automation_runner_does_not_require_login_when_main_window_is_already_open(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _disable_uploads(config)
    config.login.required = True
    config.login.username = "A0042"
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")

    runner._login_if_required(
        config,
        FakePosControl("SPA-POS 主畫面", "Window", children=[FakePosControl("統計報表", "MenuItem")]),
    )


def test_automation_runner_waits_for_required_menu_before_running_r13(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.pos.startup_wait_seconds = 2
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R13"

    splash_window = FakePosControl(
        "SPA-POS",
        "Window",
        children=[
            FakePosControl("稍候程式將自動關閉!", "Text"),
            FakePosControl("系統", "MenuItem"),
        ],
    )
    ready_window = FakePosControl(
        "SPA-POS",
        "Window",
        children=[
            FakePosControl("常用表單", "MenuItem"),
            FakePosControl("統計報表", "MenuItem"),
            FakePosControl("庫存管理", "MenuItem"),
        ],
    )
    connect_calls = 0
    downloads: list[str] = []

    def connect_pos_window(**_kwargs):  # type: ignore[no-untyped-def]
        nonlocal connect_calls
        connect_calls += 1
        return splash_window if connect_calls == 1 else ready_window

    class FakeAutomator:
        def __init__(self, window, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            assert window is ready_window

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            downloads.append(output.task_id)
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
                actions=[],
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run()

    assert summary.ok is True
    assert downloads == ["R13"]
    assert connect_calls >= 2


def test_automation_runner_switches_to_uia_when_auto_connect_returns_win32_menu_shell_without_roots(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.pos.backend = "auto"
    config.pos.startup_wait_seconds = 1
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R01"

    win32_shell = _win32_menu_shell_without_root_menus()
    uia_window = _ready_pos_window("統計報表")
    uia_window._pos_report_bot_backend = "uia"
    connect_backends: list[str] = []
    downloads: list[str] = []

    def connect_pos_window(**kwargs):  # type: ignore[no-untyped-def]
        backend = kwargs["backend"]
        connect_backends.append(backend)
        if backend == "uia":
            return uia_window
        return win32_shell

    class FakeAutomator:
        def __init__(self, window, *_args, **kwargs):  # type: ignore[no-untyped-def]
            assert window is uia_window
            assert kwargs["runtime_metadata"]["connected_backend"] == "uia"
            assert (
                kwargs["runtime_metadata"]["automation_logic_fingerprint"]
                == AUTOMATION_LOGIC_FINGERPRINT
            )
            assert (
                kwargs["runtime_metadata"]["export_format_probe"]
                == "desktop-menu-plus-bounded-report-scope"
            )

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            downloads.append(output.task_id)
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
                actions=[],
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run()

    assert summary.ok is True
    assert downloads == ["R01"]
    assert connect_backends[:2] == ["auto", "uia"]


def test_automation_runner_lets_automator_handle_win32_menu_shell_when_uia_unavailable(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.pos.backend = "auto"
    config.pos.startup_wait_seconds = 1
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R13"

    win32_shell = _win32_menu_shell_without_root_menus()
    connect_backends: list[str] = []
    downloads: list[str] = []

    def connect_pos_window(**kwargs):  # type: ignore[no-untyped-def]
        backend = kwargs["backend"]
        connect_backends.append(backend)
        if backend == "uia":
            raise UiProbeError("uia unavailable")
        return win32_shell

    class FakeAutomator:
        def __init__(self, window, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            assert window is win32_shell

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            downloads.append(output.task_id)
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
                actions=[],
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run()

    assert summary.ok is True
    assert downloads == ["R13"]
    assert connect_backends[:2] == ["auto", "uia"]


def test_automation_runner_waits_for_required_menu_with_normalized_ui_text(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.pos.startup_wait_seconds = 1
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R13"

    ready_window = FakePosControl(
        "SPA-POS",
        "Window",
        children=[
            FakePosControl("&統計報表", "MenuItem"),
            FakePosControl("庫存 管理(&I)", "MenuItem"),
        ],
    )
    downloads: list[str] = []

    class FakeAutomator:
        def __init__(self, window, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            assert window is ready_window

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            downloads.append(output.task_id)
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
                actions=[],
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: ready_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,
        drive_uploader_factory=lambda _config: MockDriveUploader(),
    )

    summary = runner.run()

    assert summary.ok is True
    assert downloads == ["R13"]


def test_automation_runner_detects_required_menu_nested_under_menu_strip(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.pos.startup_wait_seconds = 1
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R13"

    ready_window = FakePosControl(
        "SPA-POS Ver.1.5.18.69",
        "Window",
        children=[
            FakePosControl(
                "menuStrip1",
                "MenuBar",
                children=[
                    FakePosControl("統計報表", "MenuItem"),
                    FakePosControl("庫存管理", "MenuItem"),
                ],
            ),
            FakePosControl("登入檢查完成!請從上方選單選取您要執行的功能.", "Text"),
        ],
    )
    downloads: list[str] = []

    class FakeAutomator:
        def __init__(self, window, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            assert window is ready_window

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            downloads.append(output.task_id)
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
                actions=[],
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: ready_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,
    )

    summary = runner.run()

    assert summary.ok is True
    assert downloads == ["R13"]
    assert "庫存管理" in runner._visible_control_names(ready_window)


def test_automation_runner_control_scan_supplements_children_with_descendants(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.pos.startup_wait_seconds = 1
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R13"

    menu_strip = FakePosControl("menuStrip1", "MenuBar")
    inventory_menu = FakePosControl("庫存管理", "MenuItem")

    class DescendantOnlyMenuWindow(FakePosControl):
        def children(self) -> list[FakePosControl]:
            return [menu_strip]

        def descendants(self) -> list[FakePosControl]:
            return [menu_strip, inventory_menu]

    ready_window = DescendantOnlyMenuWindow("SPA-POS", "Window")
    downloads: list[str] = []

    class FakeAutomator:
        def __init__(self, window, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            assert window is ready_window

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            downloads.append(output.task_id)
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
                actions=[],
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: ready_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,
    )

    summary = runner.run()

    assert summary.ok is True
    assert downloads == ["R13"]
    assert runner._visible_control_names(ready_window) == ["SPA-POS", "menuStrip1", "庫存管理"]


def test_automation_runner_accepts_open_pos_with_ready_status_but_unenumerated_root_menu(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.pos.startup_wait_seconds = 1
    config.login.required = True
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R01"

    ready_window = FakePosControl(
        "SPA-POS Ver.1.5.18.77 美力時尚診所 HQ01-營運總部",
        "Window",
        children=[
            FakePosControl("menuStrip1", "MenuBar"),
            FakePosControl("登入檢查完成!請從上方選單選取您要執行的功能.", "Text"),
        ],
    )
    downloads: list[str] = []

    class FakeAutomator:
        def __init__(self, window, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            assert window is ready_window

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            downloads.append(output.task_id)
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
                actions=[],
            )

    def unexpected_password_lookup() -> str:
        raise AssertionError("already logged-in POS must not request a password")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: ready_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,
        pos_login_secret_provider=unexpected_password_lookup,
    )

    summary = runner.run()

    assert summary.ok is True
    assert downloads == ["R01"]
    assert runner._visible_control_names(ready_window) == [
        "SPA-POS Ver.1.5.18.77 美力時尚診所 HQ01-營運總部",
        "menuStrip1",
        "登入檢查完成!請從上方選單選取您要執行的功能.",
    ]


def test_automation_runner_blocks_update_dialog_during_readiness_when_login_disabled(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.pos.startup_wait_seconds = 1
    config.login.required = False
    config.pos_update.action = "detect_only"
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R01"

    update_dialog = FakePosControl(
        "程式更新需重新啟動",
        "Window",
        children=[
            FakePosControl("新版程式已經下載安裝完成(1.5.18.77),需要重新啟動程式!", "Text"),
            FakePosControl("是(Y)", "Button"),
        ],
    )
    blocked_window = FakePosControl(
        "SPA-POS",
        "Window",
        children=[
            FakePosControl("統計報表", "MenuItem"),
            update_dialog,
        ],
    )

    class UnexpectedAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            raise AssertionError("report automation must not start while the update dialog is blocking POS")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: blocked_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=UnexpectedAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run()

    assert summary.ok is False
    assert summary.error_code == "POS_CONNECTION_FAILED"
    assert "POS_UPDATE_PENDING" in summary.message


def test_automation_runner_r14_template_path_falls_back_to_packaged_template(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.r14_transform.template_path = ""
    config.r14_transform.template_search_dir = str(tmp_path / "missing_templates")
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")

    template_path = runner._resolve_r14_template_path()

    assert template_path.name == "診所stock status - 2026 demand planning-template.xlsx"
    assert template_path.exists()


def test_automation_runner_launch_wait_handles_splash_then_login_window(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.pos.startup_wait_seconds = 2
    config.login.required = True
    config.login.username = "A0042"
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R13"

    splash_window = FakePosControl(
        "SPA-POS",
        "Window",
        children=[
            FakePosControl("稍候程式將自動關閉!", "Text"),
            FakePosControl("系統", "MenuItem"),
        ],
    )
    login_window = FakePosControl("帳號登入", "Window")
    account_edit = FakePosControl("", "Edit")
    secret_edit = FakePosControl("", "Edit")

    def complete_login() -> None:
        login_window.name = "SPA-POS 主畫面"
        login_window.children_controls = [
            FakePosControl("統計報表", "MenuItem"),
            FakePosControl("庫存管理", "MenuItem"),
        ]

    login_window.children_controls = [
        FakePosControl("帳號", "Text"),
        account_edit,
        FakePosControl("密碼", "Text"),
        secret_edit,
        FakePosControl("登入", "Button", on_click=complete_login),
    ]
    connect_calls = 0
    downloads: list[str] = []
    launched: list[str] = []

    def connect_pos_window(**_kwargs):  # type: ignore[no-untyped-def]
        nonlocal connect_calls
        connect_calls += 1
        if connect_calls == 1:
            raise UiProbeError("POS not running")
        if connect_calls == 2:
            return splash_window
        return login_window

    class FakeAutomator:
        def __init__(self, window, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            assert window is login_window

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            downloads.append(output.task_id)
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
                actions=[],
            )

    class TestRunner(AutomationRunner):
        def _launch_pos_process(self, _config):  # type: ignore[no-untyped-def]
            launched.append("launch")

    runner = TestRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,
        drive_uploader_factory=lambda _config: MockDriveUploader(),
        pos_login_secret_provider=lambda: "test-password",
    )

    summary = runner.run()

    assert summary.ok is True
    assert launched == ["launch"]
    assert account_edit.text_value == "A0042"
    assert secret_edit.text_value == "test-password"
    assert downloads == ["R13"]


def test_automation_runner_writes_diagnostic_when_pos_never_reaches_required_menu(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.app.logs_dir = str(tmp_path / "logs")
    config.pos.startup_wait_seconds = 1
    config.login.required = False
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R13"

    class UnreadableWindow(FakePosControl):
        def window_text(self) -> str:
            return ""

        def children(self) -> list[FakePosControl]:
            return []

        def descendants(self) -> list[FakePosControl]:
            return []

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: UnreadableWindow("SPA-POS"),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
    )

    summary = runner.run()
    diagnostics = list((tmp_path / "logs" / date.today().strftime("%Y%m%d")).glob("automation_prepare_failure_*.json"))

    assert summary.ok is False
    assert summary.error_code == "POS_CONNECTION_FAILED"
    assert diagnostics
    payload = json.loads(diagnostics[0].read_text(encoding="utf-8"))
    assert payload["error"]["code"] == "POS_CONNECTION_FAILED"
    assert payload["required_root_menus"] == ["庫存管理"]
    assert payload["visible_control_names"] == []
    assert "desktop_windows" in payload
    evidence = payload["failure_evidence"]
    assert evidence["ui_probe_path"] is not None
    assert Path(evidence["ui_probe_path"]).exists()
    probe_payload = json.loads(Path(evidence["ui_probe_path"]).read_text(encoding="utf-8"))
    assert probe_payload["window_title"] == "SPA-POS"
    assert "診斷檔" in summary.message


def test_automation_runner_preparation_diagnostic_uses_cached_visible_controls_when_reconnect_fails(
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    config.app.logs_dir = str(tmp_path / "logs")
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R13"
    outputs = build_dry_run_plan(config).outputs
    main_window = FakePosControl(
        "SPA-POS 主畫面",
        "Window",
        children=[
            FakePosControl("常用表單", "MenuItem"),
            FakePosControl("庫存管理", "MenuItem"),
        ],
    )
    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(UiProbeError("temporary reconnect failure")),
    )
    assert runner._visible_control_names(main_window) == ["SPA-POS 主畫面", "常用表單", "庫存管理"]

    diagnostic_path = runner._write_preparation_failure_diagnostic(
        outputs,
        error_code="POS_CONNECTION_FAILED",
        message="準備 POS 失敗",
    )

    assert diagnostic_path is not None
    payload = json.loads(diagnostic_path.read_text(encoding="utf-8"))
    assert payload["visible_control_names"] == ["SPA-POS 主畫面", "常用表單", "庫存管理"]


def test_automation_runner_required_roots_include_default_and_explicit_report_menus(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    for report in config.reports:
        report.enabled = report.id in {"R01", "R13"}
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")
    outputs = [output for output in build_dry_run_plan(config).outputs if output.task_id in {"R01", "R13"}]

    roots = runner._required_report_root_menus(config, outputs)

    assert roots == ("統計報表", "庫存管理")


def test_automation_runner_required_roots_treat_w02_as_inventory_pos_workflow(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    for report in config.reports:
        report.enabled = report.id == "W02"
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")
    outputs = [
        PlannedOutput(
            task_id="W02",
            task_name="雙週五 R14 下單數 POS 分店訂貨",
            frequency="biweekly",
            handler="w02_pos_order_creation",
            report_menu_text="分店訂貨單",
            menu_path=["庫存管理", "分店訂貨單"],
            branch_mode="all",
            branch_code=None,
            branch_display_name=None,
            start_date="2026/07/04",
            end_date="2026/07/04",
            output_filename="",
            drive_folder_id=None,
            drive_target_status="disabled",
            upload_enabled=False,
            real_pos_validation_status="pending_real_pos_validation",
        )
    ]

    roots = runner._required_report_root_menus(config, outputs)

    assert roots == ("庫存管理",)


def test_automation_runner_waits_for_all_required_roots_in_mixed_batch(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.pos.startup_wait_seconds = 2
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id in {"R01", "R13"}

    stats_only_window = _ready_pos_window("統計報表")
    ready_window = _ready_pos_window("統計報表", "庫存管理")
    connect_calls = 0
    downloads: list[str] = []

    def connect_pos_window(**_kwargs):  # type: ignore[no-untyped-def]
        nonlocal connect_calls
        connect_calls += 1
        return stats_only_window if connect_calls == 1 else ready_window

    class FakeAutomator:
        def __init__(self, window, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            assert window is ready_window

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            downloads.append(output.task_id)
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
                actions=[],
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run()

    assert summary.ok is True
    assert downloads == ["R01", "R13"]
    assert connect_calls >= 2


def test_automation_runner_waits_for_r13_root_after_pos_recovery(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.pos.startup_wait_seconds = 2
    config.pos_recovery.enabled = True
    config.pos_recovery.max_restarts_per_run = 1
    _disable_uploads(config)
    for report in config.reports:
        report.enabled = report.id == "R13"

    splash_window = FakePosControl(
        "SPA-POS",
        "Window",
        children=[
            FakePosControl("稍候程式將自動關閉!", "Text"),
            FakePosControl("系統", "MenuItem"),
        ],
    )
    ready_window = _ready_pos_window("統計報表", "庫存管理")
    connect_after_recovery_calls = 0
    task_calls: list[str] = []

    class FakeAutomator:
        def __init__(self, window, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            self.window = window

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            task_calls.append(output.task_id)
            if len(task_calls) == 1:
                raise ReportAutomationError("POS_NOT_RESPONDING", "POS 無回應")
            assert self.window is ready_window
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
                actions=[],
            )

    def connect_pos_window(**_kwargs):  # type: ignore[no-untyped-def]
        nonlocal connect_after_recovery_calls
        if task_calls:
            connect_after_recovery_calls += 1
            return splash_window if connect_after_recovery_calls == 1 else ready_window
        return ready_window

    def recover_pos(_config, _on_progress):  # type: ignore[no-untyped-def]
        return splash_window

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos_window,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_recovery_func=recover_pos,
    )

    summary = runner.run()

    assert summary.ok is True
    assert task_calls == ["R13", "R13"]
    assert connect_after_recovery_calls >= 2


def test_automation_runner_uploads_downloaded_report_before_marking_success(tmp_path: Path) -> None:
    run_date = date(2026, 6, 8)
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.app.logs_dir = str(tmp_path / "logs")
    config.google_drive.upload_enabled = True
    for report in config.reports:
        report.enabled = report.id == "R01"
    config.drive_targets.targets["R01"].folder_id_or_url = "folder_r01"
    downloads: list[str] = []

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            downloads.append(output.task_id)
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        drive_uploader_factory=lambda _config: MockDriveUploader(),
        run_date=run_date,
    )

    progress_events: list[str] = []

    summary = runner.run(on_progress=lambda event: progress_events.append(event.event))

    assert downloads == ["R01"]
    assert summary.ok is True
    assert summary.completed == 1
    assert "下載與必要上傳" in summary.message
    assert "upload" in progress_events
    state_text = (tmp_path / "state" / "20260608" / "run_state_latest.json").read_text(encoding="utf-8")
    assert "uploaded" in state_text
    assert "mock-folder_r01" in state_text
    upload_logs = list((tmp_path / "logs" / "20260608").glob("automation_upload_*_R01.jsonl"))
    assert len(upload_logs) == 1
    records = [json.loads(line) for line in upload_logs[0].read_text(encoding="utf-8").splitlines()]
    assert [record["event"] for record in records] == [
        "upload_start",
        "upload_execute_start",
        "upload_result",
    ]
    assert records[0]["local_file_size"] == len(b"excel-bytes")
    assert records[1]["drive_folder_id"] == "folder_r01"
    assert records[2]["success"] is True
    assert records[2]["drive_file_id"].startswith("mock-folder_r01")
    assert records[2]["elapsed_seconds"] >= 0


def test_automation_runner_fails_downloaded_report_when_upload_target_is_missing(tmp_path: Path) -> None:
    run_date = date(2026, 6, 8)
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.app.logs_dir = str(tmp_path / "logs")
    config.google_drive.upload_enabled = True
    for report in config.reports:
        report.enabled = report.id == "R01"
        if report.id == "R01":
            report.drive_folder_id = ""
    config.drive_targets.targets["R01"].folder_id_or_url = ""

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        drive_uploader_factory=lambda _config: MockDriveUploader(),
        run_date=run_date,
    )

    summary = runner.run()

    assert summary.ok is False
    assert summary.completed == 0
    assert summary.failures[0].error_code == "DRIVE_FOLDER_ID_MISSING"
    state_text = (tmp_path / "state" / "20260608" / "run_state_latest.json").read_text(encoding="utf-8")
    assert "failed" in state_text
    assert "DRIVE_FOLDER_ID_MISSING" in state_text
    assert summary.failures[0].diagnostic_path is not None
    diagnostic_path = Path(summary.failures[0].diagnostic_path)
    assert diagnostic_path.name.startswith("automation_upload_")
    records = [json.loads(line) for line in diagnostic_path.read_text(encoding="utf-8").splitlines()]
    assert [record["event"] for record in records] == ["upload_start", "upload_result"]
    assert records[-1]["success"] is False
    assert records[-1]["error_code"] == "DRIVE_FOLDER_ID_MISSING"


def test_automation_runner_rejects_upload_success_without_drive_file_id(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.app.logs_dir = str(tmp_path / "logs")
    config.google_drive.upload_enabled = True
    for report in config.reports:
        report.enabled = report.id == "R01"
    config.drive_targets.targets["R01"].folder_id_or_url = "folder_r01"

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
            )

    class FakeUploader:
        def upload(self, file_path, folder_id, name):  # type: ignore[no-untyped-def]
            return DriveUploadResult(
                success=True,
                drive_file_id=None,
                folder_id=folder_id,
                uploaded_name=name,
                message="fake success without evidence",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        drive_uploader_factory=lambda _config: FakeUploader(),
    )

    summary = runner.run()

    assert summary.ok is False
    assert summary.completed == 0
    assert summary.failures[0].error_code == "DRIVE_FILE_ID_MISSING"
    assert summary.failures[0].diagnostic_path is not None
    diagnostic_path = Path(summary.failures[0].diagnostic_path)
    assert diagnostic_path.name.startswith("automation_upload_")
    records = [json.loads(line) for line in diagnostic_path.read_text(encoding="utf-8").splitlines()]
    assert records[-1]["event"] == "upload_result"
    assert records[-1]["success"] is True
    assert records[-1]["drive_file_id"] is None


def test_automation_runner_logs_uploader_exception_diagnostic(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.app.logs_dir = str(tmp_path / "logs")
    config.google_drive.upload_enabled = True
    for report in config.reports:
        report.enabled = report.id == "R01"
    config.drive_targets.targets["R01"].folder_id_or_url = "folder_r01"

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
            )

    class FakeUploader:
        def upload(self, file_path, folder_id, name):  # type: ignore[no-untyped-def]
            raise RuntimeError("drive timeout")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        drive_uploader_factory=lambda _config: FakeUploader(),
    )

    summary = runner.run()

    assert summary.ok is False
    assert summary.completed == 0
    assert summary.failures[0].error_code == "DRIVE_UPLOAD_FAILED"
    assert summary.failures[0].diagnostic_path is not None
    diagnostic_path = Path(summary.failures[0].diagnostic_path)
    records = [json.loads(line) for line in diagnostic_path.read_text(encoding="utf-8").splitlines()]
    assert [record["event"] for record in records] == [
        "upload_start",
        "upload_execute_start",
        "upload_result",
    ]
    assert records[-1]["success"] is False
    assert records[-1]["error_code"] == "DRIVE_UPLOAD_FAILED"
    assert "drive timeout" in records[-1]["message"]


def test_automation_runner_global_drive_upload_disabled_downloads_without_upload_requirement(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.google_drive.upload_enabled = False
    for report in config.reports:
        report.enabled = report.id == "R01"
    config.drive_targets.targets["R01"].folder_id_or_url = "folder_r01"
    download_calls = 0
    upload_calls = 0

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            nonlocal download_calls
            download_calls += 1
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
            )

    class FakeUploader:
        def upload(self, file_path, folder_id, name):  # type: ignore[no-untyped-def]
            nonlocal upload_calls
            upload_calls += 1
            return DriveUploadResult(success=True, drive_file_id="drive-file-1")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        drive_uploader_factory=lambda _config: FakeUploader(),
    )

    summary = runner.run()

    assert summary.ok is True
    assert summary.completed == 1
    assert summary.failures == ()
    assert download_calls == 1
    assert upload_calls == 0
    state_payload = json.loads(
        (tmp_path / "state" / date.today().strftime("%Y%m%d") / "run_state_latest.json").read_text(encoding="utf-8")
    )
    output_state = next(iter(state_payload["outputs"].values()))
    assert output_state["error_code"] is None
    assert output_state["status"] == "completed"
    assert output_state["drive_file_id"] is None


def test_automation_runner_global_drive_upload_disabled_allows_explicit_local_only_reports(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    config.google_drive.upload_enabled = False
    for report in config.reports:
        report.enabled = report.id == "R01"
        report.upload_enabled = False
    download_calls = 0

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            nonlocal download_calls
            download_calls += 1
            output_path = tmp_path / output.output_filename
            output_path.write_bytes(b"excel-bytes")
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=output_path,
                error_code=None,
                message="saved",
            )

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: _ready_pos_window(),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
    )

    summary = runner.run()

    assert summary.ok is True
    assert summary.completed == 1
    assert download_calls == 1


def test_automation_runner_launches_pos_when_not_running_then_logs_in_and_runs(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.app.downloads_dir = str(tmp_path)
    _disable_uploads(config)
    config.pos.executable_path = str(tmp_path / "SPA1.exe")
    config.login.required = True
    config.login.username = "A0042"
    for report in config.reports:
        report.enabled = report.id == "R01"
    account_edit = FakePosControl("", "Edit")
    secret_edit = FakePosControl("", "Edit")
    login_window = FakePosControl("帳號登入", "Window")

    def complete_login() -> None:
        login_window.name = "SPA-POS 主畫面"
        login_window.children_controls = [FakePosControl("統計報表", "MenuItem")]

    login_button = FakePosControl("登入", "Button", on_click=complete_login)
    login_window.children_controls = [
        FakePosControl("帳號", "Text"),
        account_edit,
        FakePosControl("密碼", "Text"),
        secret_edit,
        login_button,
    ]
    connect_attempts = 0
    launched: list[str] = []
    calls: list[str] = []

    def connect_pos(**_kwargs):  # type: ignore[no-untyped-def]
        nonlocal connect_attempts
        connect_attempts += 1
        if connect_attempts == 1:
            raise UiProbeError("POS not running")
        return login_window

    class FakeAutomator:
        def __init__(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            pass

        def download_report(self, output, report, *, close_after_success=True):  # type: ignore[no-untyped-def]
            calls.append(output.task_id)
            return SimpleNamespace(
                ok=True,
                task_id=output.task_id,
                output_path=tmp_path / output.output_filename,
                error_code=None,
                message="saved",
            )

    class TestRunner(AutomationRunner):
        def _launch_pos_process(self, _config):  # type: ignore[no-untyped-def]
            launched.append("launch")

    runner = TestRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=connect_pos,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        automator_factory=FakeAutomator,  # type: ignore[arg-type]
        pos_login_secret_provider=lambda: "fake-login-secret",
    )

    summary = runner.run()

    assert launched == ["launch"]
    assert connect_attempts >= 2
    assert account_edit.text_value == "A0042"
    assert secret_edit.text_value == "fake-login-secret"
    assert login_button.clicked is True
    assert calls == ["R01"]
    assert summary.ok is True


def test_automation_runner_reports_startup_failure_when_pos_not_running_and_path_missing(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    config.pos.executable_path = ""
    for report in config.reports:
        report.enabled = report.id == "R01"

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=lambda **_kwargs: (_ for _ in ()).throw(UiProbeError("POS not running")),
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
    )

    summary = runner.run()

    assert summary.ok is False
    assert summary.error_code == "POS_CONNECTION_FAILED"
    assert "找不到 POS 啟動檔" in summary.message


def test_automation_runner_records_cancellation_then_propagates_system_exit(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    _disable_uploads(config)
    config.google_drive.upload_enabled = False
    config.email.enabled = False
    for report in config.reports:
        report.enabled = report.id == "R01"
        if report.id == "R06":
            report.frequency = "weekly"
    for branch in config.branches:
        branch.enabled = False

    def exit_during_pos_connect(**_kwargs):  # type: ignore[no-untyped-def]
        raise SystemExit("simulated dependency exit")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=exit_during_pos_connect,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
        run_source="gui_manual",
        selected_task_ids={"R01", "R06"},
    )

    with pytest.raises(SystemExit, match="simulated dependency exit"):
        runner.run()

    state_payload = json.loads(
        (tmp_path / "state" / date.today().strftime("%Y%m%d") / "run_state_latest.json").read_text(
            encoding="utf-8"
        )
    )
    output_state = next(iter(state_payload["outputs"].values()))
    assert state_payload["status"] == "failed"
    assert output_state["status"] == "failed"
    assert output_state["error_code"] == "RUN_CANCELLED"
    journal_path = next((tmp_path / "logs").rglob("automation_runtime_*.jsonl"))
    phase_records = [json.loads(line) for line in journal_path.read_text(encoding="utf-8").splitlines()]
    assert phase_records[0]["selected_task_ids"] == ["R01", "R06"]
    assert phase_records[0]["planned_task_ids"] == ["R01", "R06"]
    assert phase_records[0]["disabled_but_explicitly_selected_task_ids"] == ["R06"]
    assert phase_records[0]["all_branches_disabled_but_explicitly_selected_task_ids"] == ["R06"]
    selected_diagnostics = {
        item["task_id"]: item for item in phase_records[0]["selected_task_plan_diagnostics"]
    }
    assert selected_diagnostics["R01"]["planned"] is True
    assert selected_diagnostics["R06"]["planned"] is True
    assert selected_diagnostics["R06"]["frequency"] == "weekly"
    assert selected_diagnostics["R06"]["branch_mode"] == "each_branch"
    assert selected_diagnostics["R06"]["configured_branch_count"] == 6
    assert phase_records[0]["output_count"] == 7
    assert [record["phase"] for record in phase_records] == [
        "run_started",
        "task_started",
        "pos_connect_start",
        "pos_connect_error",
        "run_cancelled",
    ]


@pytest.mark.skipif(os.name != "nt", reason="Windows file sharing semantics only")
def test_automation_runner_records_recovery_state_path_when_latest_is_locked(tmp_path: Path) -> None:
    config = _load_runner_config(tmp_path)
    _set_runtime_dirs(config, tmp_path)
    _disable_uploads(config)
    config.google_drive.upload_enabled = False
    config.email.enabled = False
    for report in config.reports:
        report.enabled = report.id == "R01"

    state_dir = tmp_path / "state" / date.today().strftime("%Y%m%d")
    state_dir.mkdir(parents=True)
    latest_path = state_dir / "run_state_latest.json"
    latest_path.write_text('{"historical": true}', encoding="utf-8")

    def exit_during_pos_connect(**_kwargs):  # type: ignore[no-untyped-def]
        raise SystemExit("simulated dependency exit")

    runner = AutomationRunner(
        config,
        settings_path=tmp_path / "app.yaml",
        app_version="test",
        connect_pos_window_func=exit_during_pos_connect,
        save_as_handler_factory=lambda _config: MockSaveAsHandler(),
    )

    with latest_path.open("r", encoding="utf-8"):
        with pytest.raises(SystemExit, match="simulated dependency exit"):
            runner.run()

    recovery_paths = list(state_dir.glob("run_state_recovery_*.json"))
    assert len(recovery_paths) == 1
    recovery_payload = json.loads(recovery_paths[0].read_text(encoding="utf-8"))
    assert recovery_payload["status"] == "failed"
    assert latest_path.read_text(encoding="utf-8") == '{"historical": true}'
    journal_path = next((tmp_path / "logs").rglob("automation_runtime_*.jsonl"))
    run_started = json.loads(journal_path.read_text(encoding="utf-8").splitlines()[0])
    assert run_started["run_state_path"] == str(recovery_paths[0])
    assert run_started["run_state_primary_path"] == str(latest_path)
    assert run_started["run_state_recovery_reason"] == "LATEST_STATE_REPLACE_DENIED"


def _summary_actual_group_columns(  # type: ignore[no-untyped-def]
    sheet,
    month_label: str,
    *,
    allow_legacy: bool = False,
) -> list[int]:
    branches = (
        "站前4樓",
        "站前11樓",
        "忠孝7樓",
        "忠孝國際醫學3樓",
        "忠孝健康7樓",
        "忠孝預防醫學3樓",
    )
    for col in range(1, sheet.max_column + 1):
        if sheet.cell(2, col).value != month_label or sheet.cell(3, col).value != "Actual":
            continue
        columns = [col]
        for branch in branches:
            next_col = columns[-1] + 1
            if sheet.cell(2, next_col).value != branch or sheet.cell(3, next_col).value != "Actual":
                break
            columns.append(next_col)
        if len(columns) == 1 + len(branches):
            return columns
        if allow_legacy and len(columns) == len(branches):
            return columns
    raise AssertionError(f"Summary Actual group not found: {month_label}")


def _find_workbook_item_row(sheet, item_code: str) -> int:  # type: ignore[no-untyped-def]
    for row in range(4, sheet.max_row + 1):
        if str(sheet.cell(row, 2).value).strip() == item_code:
            return row
    raise AssertionError(f"item not found: {item_code}")


@pytest.mark.parametrize("error_code", ["REPORT_SCREEN_NOT_OPENED", "REPORT_ROOT_MENU_NOT_FOUND"])
def test_automation_runner_treats_report_surface_failure_as_recoverable(
    error_code: str,
    tmp_path: Path,
) -> None:
    config = _load_runner_config(tmp_path)
    config.pos_recovery.enabled = True
    config.pos_recovery.max_restarts_per_run = 1
    runner = AutomationRunner(config, settings_path=tmp_path / "app.yaml", app_version="test")

    assert runner._can_recover_pos(ReportAutomationError(error_code, "POS 報表畫面失效"), 0) is True
